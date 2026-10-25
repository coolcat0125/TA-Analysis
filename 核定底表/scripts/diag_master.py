# -*- coding: utf-8 -*-
"""诊断基线：不改数，先量化（Phase 0）。

D1 车长错配率：12 家×车系「同系车长极差>300mm」（补齐比亚迪41%/吉利6% 以外的未知数）
D2 三方对账：合并版 xlsx vs master_vehicles.json（派生库）vs 411/412 扫描底表
D3 W7 扫描：整备质量/车长 出界 [0.18, 0.65]（17 AGENTS 2026-09-13 固化判据）
D4 归一化残余变体：同组内疑似同车异名（前缀包含关系等）
PY 核定优先级清单：T1 错配嫌疑 / T2 行数≥3 / T3 2025-26 新车 / T4 长尾

产出：data/diag_report.md + data/diag_data.json
"""
import collections
import json
import os
import sys
from datetime import date

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import (SRC_XLSX, DERIVED_DB, SCAN_XLSX, DATA, DIAG_MD, DIAG_JSON,
                    num, year_of_batch, is_ptype_name)
from oem_groups import GROUPS, BOARDS, EXCL_ENT, SERIES_BOARD, board_of, attr_group
from series_norm import norm_series

SPREAD = 300          # 同系车长极差阈值(mm)
W7_LO, W7_HI = 0.18, 0.65
TOL = 0.5             # 对账容差（17 派生库核验口径）


def load_main():
    """主表全量行 → dict 列表（含归属/归一化）。"""
    wb = openpyxl.load_workbook(SRC_XLSX, read_only=True, data_only=True)
    ws = wb['NEV公告参数汇总']
    rows = ws.iter_rows(values_only=True)
    hdr = list(next(rows))
    I = {h: i for i, h in enumerate(hdr)}
    data = [r for r in rows if r[I['产品型号']]]
    wb.close()

    out = []
    for r in data:
        mark = str(r[I['产品商标']] or '')
        ent = str(r[I['企业名称']] or '')
        key, jv, why = attr_group(mark, ent)
        batch = int(num(r[I['批次']]) or 0)
        # 车系名（占位型产品类型 → 置空，C 问题单独统计）
        series_raw = ''
        for k in ('通用名称', '车型名称', '产品名称'):
            v = str(r[I[k]] or '').strip()
            if v and not is_ptype_name(v):
                series_raw = v
                break
        board, series = None, None
        if key:
            bd = None
            for sub, nm in SERIES_BOARD.get(key, []):
                if sub and sub in series_raw:
                    bd = nm
                    break
            if bd is None:
                bd = board_of(mark, ent, BOARDS[key])
            board = bd
            series = norm_series(series_raw, key, bd) if series_raw else None
        out.append({
            'batch': batch, 'model': str(r[I['产品型号']]).strip(),
            'mark': mark, 'ent': ent,
            'group': key, 'jv': jv, 'why': why,
            'common_raw': series_raw,
            'common_src': str(r[I['通用名称']] or ''),
            'board': board, 'series': series,
            'len': num(r[I['车长(mm)']]),
            'wb': num(r[I['轴距(mm)']]),
            'wt': num(r[I['整备质量(kg)']]),
            'cap': num(r[I['电池容量(kWh)']]),
        })
    return out


def d1_mismatch(rows):
    """D1：按集团统计同系车长极差>SPREAD 的行占比（合资行不计入，主口径外板块计入）。"""
    stats, bad = {}, []
    for key in GROUPS:
        rs = [r for r in rows if r['group'] == key and not r['jv']]
        by_series = collections.defaultdict(list)
        for r in rs:
            if r['series'] and r['len']:
                by_series[r['series']].append(r)
        tot = sum(len(v) for v in by_series.values())
        bad_rows, bad_series = 0, []
        for s, v in by_series.items():
            lens = collections.Counter(int(r['len']) for r in v)
            spread = max(lens) - min(lens)
            if spread > SPREAD:
                bad_rows += len(v)
                bad_series.append({'group': key, 'series': s, 'rows': len(v),
                                   'spread': spread,
                                   'lens': [{'len': L, 'n': n} for L, n in
                                            sorted(lens.items())]})
        bad += sorted(bad_series, key=lambda z: -z['rows'])
        stats[key] = {'cn': GROUPS[key][0], 'rows': len(rs),
                      'rows_with_len': tot,
                      'bad_series': len(bad_series), 'bad_rows': bad_rows,
                      'rate': round(bad_rows / tot, 3) if tot else None}
    return stats, bad


def d2_crosscheck(rows):
    """D2：xlsx vs master_vehicles.json 按 (批,型号) 逐格对账。"""
    if not os.path.exists(DERIVED_DB):
        return {'skip': 'no master_vehicles.json'}
    recs = json.load(open(DERIVED_DB, encoding='utf-8'))
    mv = {}
    for r in recs:
        mv[(str(r.get('batchNo') or ''), str(r.get('model') or ''))] = r
    n_pair = n_len = mm_len = mm_wb = mm_cap = 0
    mism = []
    for r in rows:
        k = (str(r['batch']), r['model'])
        o = mv.get(k)
        if not o:
            continue
        n_pair += 1
        if r['len'] is not None and o.get('lengthMm') not in (None, ''):
            n_len += 1
            if abs(r['len'] - num(o['lengthMm'])) > TOL:
                mm_len += 1
                if len(mism) < 400:
                    mism.append({'key': k, 'field': '车长',
                                 'xlsx': r['len'], 'db': num(o['lengthMm']),
                                 'series': r['series'], 'group': r['group']})
        if r['wb'] is not None and o.get('wheelbaseMm') not in (None, ''):
            if abs(r['wb'] - num(o['wheelbaseMm'])) > TOL:
                mm_wb += 1
        if r['cap'] is not None and o.get('batteryKwh') not in (None, ''):
            if abs(r['cap'] - num(o['batteryKwh'])) > TOL:
                mm_cap += 1
    # 扫描底表 411/412 vs 派生库 411
    scan = load_scan()
    s411 = {r['model']: r for r in scan if str(r['batch']).startswith('411')}
    n411 = sum(1 for m in s411 if any(k[1] == m and k[0] == '411' for k in mv))
    return {'pairs': n_pair, 'len_compared': n_len, 'len_mismatch': mm_len,
            'wb_mismatch': mm_wb, 'cap_mismatch': mm_cap,
            'mismatch_samples': mism,
            'scan411_models': len(s411), 'scan411_in_db': n411}


def d3_w7(rows):
    """D3：整备/车长比出界（0.18~0.65），17 W7 判据；同时按 15 口径 [200,700] kg/m 计数。"""
    hits = []
    for r in rows:
        if not (r['len'] and r['wt']):
            continue
        ratio = r['wt'] / r['len']
        if not (W7_LO <= ratio <= W7_HI):
            kgm = r['wt'] / (r['len'] / 1000)
            hits.append({'group': r['group'], 'series': r['series'],
                         'batch': r['batch'], 'model': r['model'],
                         'len': r['len'], 'wt': r['wt'],
                         'ratio': round(ratio, 3), 'kg_per_m': round(kgm, 1),
                         'in_15_bound': 200 <= kgm <= 700})
    return hits


def d4_variants(rows):
    """D4：同集团内疑似同车异名（一方是另一方前缀/去前缀形式）。"""
    out = []
    for key in GROUPS:
        names = collections.Counter(
            r['series'] for r in rows
            if r['group'] == key and not r['jv'] and r['series'])
        items = [(s, n) for s, n in names.items() if s and s != '未命名']
        for s, n in items:
            for s2, n2 in items:
                if s != s2 and s.endswith(s2) and len(s) - len(s2) <= 4 and n < n2:
                    out.append({'group': key, 'keep': s2, 'drop': s,
                                'n_keep': n2, 'n_drop': n})
    return sorted(out, key=lambda z: (z['group'], -z['n_drop']))


def load_scan():
    """411/412 扫描底表（表头 r3）。"""
    wb = openpyxl.load_workbook(SCAN_XLSX, read_only=True, data_only=True)
    out = []
    for sn in ('411正式车型', '412公示车型'):
        ws = wb[sn]
        it = ws.iter_rows(values_only=True)
        for _ in range(2):
            next(it)
        hdr = list(next(it))
        J = {h: i for i, h in enumerate(hdr)}
        for r in it:
            if not r or not r[J['产品型号']]:
                continue
            out.append({'sheet': sn, 'batch': r[J['批次']], 'status': r[J['状态']],
                        'model': str(r[J['产品型号']]),
                        'common': str(r[J['通用名称']] or ''),
                        'len': num(r[J['车长(mm)']]), 'wb': num(r[J['轴距(mm)']]),
                        'wt': num(r[J['整备质量(kg)']]),
                        'cap': num(r[J['电池容量(kWh)']]),
                        'url': str(r[J.get('来源URL')] or '') if '来源URL' in J else ''})
    wb.close()
    return out


def priority(rows):
    """核定优先级：T1 错配 / T2 行数≥3 / T3 2025-26 新增 / T4 长尾。"""
    agg = {}
    for r in rows:
        if not (r['group'] and not r['jv'] and r['series'] and r['series'] != '未命名'):
            continue
        k = (r['group'], r['series'])
        a = agg.setdefault(k, {'group': r['group'], 'series': r['series'],
                               'board': r['board'], 'rows': 0, 'lens': collections.Counter(),
                               'batches': [9999, 0]})
        a['rows'] += 1
        if r['len']:
            a['lens'][int(r['len'])] += 1
        a['batches'][0] = min(a['batches'][0], r['batch'])
        a['batches'][1] = max(a['batches'][1], r['batch'])
    out = []
    for a in agg.values():
        lens = list(a['lens'])
        spread = (max(lens) - min(lens)) if lens else None
        tier = 'T4'
        if a['rows'] >= 3:
            tier = 'T2'
        if spread and spread > SPREAD:
            tier = 'T1'
        elif a['batches'][0] >= 401:
            tier = 'T3' if tier == 'T4' else tier
        out.append({'tier': tier, 'group': a['group'], 'series': a['series'],
                    'board': a['board'], 'rows': a['rows'],
                    'spread': spread, 'batches': a['batches'],
                    'lens': sorted(a['lens'].items(), key=lambda kv: -kv[1])[:6]})
    order = {'T1': 0, 'T2': 1, 'T3': 2, 'T4': 3}
    return sorted(out, key=lambda z: (order[z['tier']], -z['rows'], z['group']))


def main():
    rows = load_main()
    print('主表行', len(rows), '| 12 家内(非合资)', sum(1 for r in rows if r['group'] and not r['jv']),
          '| 12 家内合资', sum(1 for r in rows if r['jv']), '| 未归属',
          sum(1 for r in rows if not r['group']))

    g_stats, bad_series = d1_mismatch(rows)
    cc = d2_crosscheck(rows)
    w7 = d3_w7(rows)
    var = d4_variants(rows)
    pri = priority(rows)
    scan = load_scan()

    tiers = collections.Counter(p['tier'] for p in pri)
    os.makedirs(DATA, exist_ok=True)
    json.dump({'run': str(date.today()), 'group_stats': g_stats,
               'bad_series': bad_series, 'crosscheck': cc, 'w7': w7,
               'variants': var, 'priority': pri,
               'scan': {'n411': sum(1 for s in scan if s['sheet'] == '411正式车型'),
                        'n412': sum(1 for s in scan if s['sheet'] == '412公示车型')}},
              open(DIAG_JSON, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

    L = ['# 核定底表 · 诊断基线报告', f'', f'运行：{date.today()} · 源：合并版 xlsx（341~410）', '']
    L.append('## D1 车长错配率（同系极差>300mm，合资不计入）')
    L.append('| 集团 | 总行 | 有车长行 | 错配车系数 | 错配行 | 错配率 |')
    L.append('|---|---|---|---|---|---|')
    for k, s in sorted(g_stats.items(), key=lambda kv: -(kv[1]['rate'] or 0)):
        L.append(f"| {s['cn']} | {s['rows']} | {s['rows_with_len']} | "
                 f"{s['bad_series']} | {s['bad_rows']} | "
                 f"{(s['rate'] or 0) * 100:.0f}% |")
    L.append('')
    L.append('## D2 三方对账（xlsx vs 派生库 master_vehicles.json）')
    L.append(f"- 配对 (批,型号)：{cc.get('pairs')}；车长可比 {cc.get('len_compared')}，"
             f"不一致 {cc.get('len_mismatch')}；轴距不一致 {cc.get('wb_mismatch')}；"
             f"容量不一致 {cc.get('cap_mismatch')}")
    L.append(f"- 扫描底表：411 正式 {cc.get('scan411_models')} 条（派生库有 411 "
             f"{cc.get('scan411_in_db')} 条）、412 公示另行单列")
    L.append('')
    L.append('## D3 W7 出界（整备/车长 比 ∉ [0.18,0.65]）')
    g_w7 = collections.Counter((h['group'] or '未归属') for h in w7)
    L.append('- ' + '；'.join(f'{k}:{v}' for k, v in g_w7.most_common()) or '- 无')
    L.append('')
    L.append('## D4 疑似同车异名（前缀包含，归一化残余）')
    for v in var[:40]:
        L.append(f"- [{v['group']}] 保留 `{v['keep']}`({v['n_keep']}) ← 吸收 `{v['drop']}`({v['n_drop']})")
    L.append('')
    L.append('## 核定优先级（T1 错配 / T2 ≥3行 / T3 2025-26 新 / T4 长尾）')
    L.append(f"- T1 {tiers['T1']} 系 / T2 {tiers['T2']} 系 / T3 {tiers['T3']} 系 / T4 {tiers['T4']} 系")
    L.append('')
    L.append('### T1 全量')
    for p in [z for z in pri if z['tier'] == 'T1']:
        lens = ' '.join(f'{L_}×{n}' for L_, n in p['lens'])
        L.append(f"- [{p['group']}] {p['series']} 行{p['rows']} 批{p['batches'][0]}~{p['batches'][1]} 极差{p['spread']}：{lens}")
    open(DIAG_MD, 'w', encoding='utf-8').write('\n'.join(L))
    print('SAVED', DIAG_MD, 'and', DIAG_JSON)
    print('T1/T2/T3/T4 =', dict(tiers), '| W7 hits =', len(w7), '| variants =', len(var))


if __name__ == '__main__':
    main()
