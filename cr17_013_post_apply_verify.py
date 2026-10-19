# -*- coding: utf-8 -*-
"""CR-17-013: 411/412 落盘后的全量核验（对照落盘前备份逐格比对 + 三门禁 + 并发列不变量）。

用法：python cr17_013_post_apply_verify.py [--backup <xlsx>] [--json <out>]
默认自动取 archive/ 下最新的 v4.9.32_pre-411412_backup_*.xlsx。
"""
import argparse, glob, hashlib, io, json, os, subprocess, sys
from collections import Counter
from datetime import datetime
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
WB = os.path.join(HERE, "NEV公告参数汇总表_合并版（341~410批）.xlsx")
ARCHIVE = os.path.join(HERE, "archive")
OUT = os.path.join(HERE, "audit-output", "cr17_013_post_apply_verify_20261019.json")
OUT_DIR = os.path.join(HERE, "audit-output")

ALLOWED_ENRICH_COLS = {"官方参数页URL", "车型名称", "电池容量(kWh)", "前电机功率/扭矩",
                       "电机总功率/扭矩", "后电机功率/扭矩", "是否增程(EREV)"}
PEER_SALES_COLS = ("月销量(辆)", "销量数据月份", "销量更新日期")
BATCHES_NEW = (411, 412)


def log(*a):
    print(*a, flush=True)


def sha256(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for c in iter(lambda: f.read(1 << 20), b""):
            h.update(c)
    return h.hexdigest()


def sheet_rows(path, name):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[name]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        hdr = rows[0] if rows else []
        return hdr, rows
    finally:
        wb.close()


def norm(v):
    if isinstance(v, str):
        return v.strip()
    return v


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--backup", default=None)
    ap.add_argument("--json", default=OUT)
    args = ap.parse_args()

    baks = sorted(glob.glob(os.path.join(ARCHIVE, "v4.9.32_pre-411412_backup_*.xlsx")))
    bak = args.backup or (baks[-1] if baks else None)
    if not bak or not os.path.exists(bak):
        sys.exit("[FATAL] 未找到落盘前备份（archive/v4.9.32_pre-411412_backup_*.xlsx）")
    rep = {"at": datetime.now().isoformat(timespec="seconds"), "backup": os.path.relpath(bak, HERE),
           "workbook": os.path.basename(WB), "sha_backup": sha256(bak), "sha_now": sha256(WB)}

    # ---------- 1) 主表逐格比对 ----------
    hdr_b, rows_b = sheet_rows(bak, "NEV公告参数汇总")
    hdr_n, rows_n = sheet_rows(WB, "NEV公告参数汇总")
    rep["main_hdr_backup"] = hdr_b
    rep["main_hdr_now"] = hdr_n
    rep["backup_data_rows"] = len(rows_b) - 1
    rep["now_data_rows"] = len(rows_n) - 1
    same_hdr = hdr_b == hdr_n
    rep["header_unchanged"] = same_hdr

    idx_b = {h: i for i, h in enumerate(hdr_b) if h}
    idx_n = {h: i for i, h in enumerate(hdr_n) if h}
    ncol = max(len(hdr_b), len(hdr_n))

    changes = Counter()
    outside = []
    enriched_cells = 0
    for r in range(1, min(len(rows_b), len(rows_n))):
        rb, rn = rows_b[r], rows_n[r]
        for c in range(ncol):
            vb = rb[c] if c < len(rb) else None
            vn = rn[c] if c < len(rn) else None
            if norm(vb) != norm(vn):
                cname = hdr_n[c] if c < len(hdr_n) and hdr_n[c] else f"col{c}"
                if vb is None and vn is not None:
                    changes[cname] += 1
                    enriched_cells += 1
                    if cname not in ALLOWED_ENRICH_COLS:
                        outside.append({"row": r + 1, "col": cname, "before": vb, "after": vn})
                else:
                    outside.append({"row": r + 1, "col": cname, "before": vb, "after": vn,
                                    "kind": "非空→改写/清空"})
    rep["enrich_cells"] = enriched_cells
    rep["enrich_by_col"] = dict(changes)
    rep["unexpected_changes"] = outside[:50]
    rep["unexpected_change_count"] = len(outside)

    # ---------- 2) 新增行不变量 ----------
    new_rows = rows_n[len(rows_b):]
    rep["new_rows"] = len(new_rows)
    viol = []
    batches = Counter()
    for k, r in enumerate(new_rows):
        g = lambda c: r[idx_n[c]] if c in idx_n and idx_n[c] < len(r) else None
        b = g("批次")
        batches[str(b)] += 1
        if b not in BATCHES_NEW:
            viol.append((k, "批次", b))
        if not g("产品型号"):
            viol.append((k, "产品型号", None))
        for c in PEER_SALES_COLS:
            if c in idx_n and g(c) is not None:
                viol.append((k, c, g(c)))
        if "车系ID" in idx_n and g("车系ID") is not None:
            viol.append((k, "车系ID", g("车系ID")))
        if "销量覆盖状态" in idx_n and not g("销量覆盖状态"):
            viol.append((k, "销量覆盖状态", None))
    rep["new_batches"] = dict(batches)
    rep["new_row_violations"] = viol[:30]
    rep["new_row_violation_count"] = len(viol)

    # ---------- 3) 台账 ----------
    hdr_lb, rows_lb = sheet_rows(bak, "变更记录")
    hdr_ln, rows_ln = sheet_rows(WB, "变更记录")
    rep["ledger_rows_backup"] = len(rows_lb) - 1
    rep["ledger_rows_now"] = len(rows_ln) - 1
    led_new = rows_ln[len(rows_lb):]
    seqs = [r[0] for r in led_new]
    rep["ledger_appended"] = len(led_new)
    rep["ledger_seq_range"] = [seqs[0], seqs[-1]] if seqs else None
    rep["ledger_seq_strictly_increasing"] = all(
        isinstance(seqs[i], (int, float)) and isinstance(seqs[i + 1], (int, float)) and seqs[i + 1] == seqs[i] + 1
        for i in range(len(seqs) - 1)) if len(seqs) > 1 else None
    rep["ledger_type_counts"] = dict(Counter(str(r[5]) for r in led_new))

    # ---------- 4) 时间表 / 字段映射 ----------
    for sh, key in (("批次时间表", "time"), ("字段映射说明", "map")):
        _, rb = sheet_rows(bak, sh)
        _, rn = sheet_rows(WB, sh)
        rep[f"{key}_rows_backup"] = len(rb) - 1
        rep[f"{key}_rows_now"] = len(rn) - 1
        rep[f"{key}_appended"] = [list(x) for x in rn[len(rb):]]

    # ---------- 5) 门禁：audit_data --strict ----------
    tmp = os.path.join(os.environ.get("TEMP", "/tmp"), "cr17_013_audit")
    os.makedirs(tmp, exist_ok=True)
    r1 = subprocess.run([sys.executable, os.path.join(HERE, "audit_data.py"), "--strict",
                         "--input", WB, "--output-dir", tmp],
                        capture_output=True, text=True, encoding="utf-8", errors="replace", cwd=HERE)
    rep["audit_rc"] = r1.returncode
    rep["audit_out"] = (r1.stdout or "").strip().splitlines()[-2:]
    jp = os.path.join(tmp, "data_audit_report.json")
    if os.path.exists(jp):
        j = json.load(open(jp, encoding="utf-8"))
        rep["audit_records"] = j["dataset"]["records"]
        rep["audit_columns"] = j["dataset"]["columns"]
        rep["audit_gate_passed"] = j["quality_gates"]["passed"]
        rep["audit_gate_failures"] = j["quality_gates"].get("failures")
        rep["audit_batch_max"] = max(int(k) for k in j["dataset"]["batches"] if str(k).isdigit())

    # ---------- 6) verify_consistency ----------
    cj = os.path.join(OUT_DIR, "cr17_013_consistency_after.json")
    r2 = subprocess.run([sys.executable, os.path.join(HERE, "verify_consistency.py"),
                         "--input", WB, "--json", cj],
                        capture_output=True, text=True, encoding="utf-8", errors="replace", cwd=HERE)
    rep["consistency_rc"] = r2.returncode
    rep["consistency_out"] = (r2.stdout or "").strip().splitlines()[-3:]
    if os.path.exists(cj):
        c = json.load(open(cj, encoding="utf-8"))
        rep["consistency_summary"] = c.get("summary") or {k: (len(v) if isinstance(v, list) else v)
                                                          for k, v in c.items() if k in ("critical", "warn", "stats")}

    with open(args.json, "w", encoding="utf-8") as f:
        json.dump(rep, f, ensure_ascii=False, indent=2)
    log(json.dumps({k: v for k, v in rep.items() if k not in ("main_hdr_backup", "main_hdr_now",
                                                              "time_appended", "map_appended",
                                                              "unexpected_changes")},
                   ensure_ascii=False, indent=2))
    log("\n[时间表新增]", rep.get("time_appended"))
    log("[字段映射新增]", rep.get("map_appended"))
    log("[非预期变更]", rep["unexpected_change_count"], rep["unexpected_changes"][:5])
    log("→ 报告", args.json)


main()
