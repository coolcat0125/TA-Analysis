# -*- coding: utf-8 -*-
"""capacity_evidence_v4925.py — 容量矛盾 188 行·汽车之家取证轮（方向 C 续作，可复跑）

纪律（Update/夜间分派_20260928.md §P1）：
  · 唯一外部证据才提名更正；多值/体系并存弃权留档（36≠37 教训）
  · 禁止电耗×续航反推选值（09-22 同源循环裁决）——本脚本只做外部取证
  · 代际窗口=公告月份 ≥7 允许 Y+1（S4 同款）；品牌守卫 media_fill.brand_conflict
用法：
  python capacity_evidence_v4925.py --fetch   # 抓取+提名（网络，限速 0.55s）
  python capacity_evidence_v4925.py --apply   # 提名落库（浅红+台账，先审 fetch 报告）
"""
import argparse
import json
import os
import re
import sys
import time
from collections import defaultdict

import openpyxl
from openpyxl.styles import PatternFill

import media_fill as mf

HERE = os.path.dirname(os.path.abspath(__file__))
WB = os.path.join(HERE, 'NEV公告参数汇总表_合并版（341~410批）.xlsx')
XF = os.path.join(HERE, 'audit-output', 'capacity_dcd_crossref_20260927.json')
OUT = os.path.join(HERE, 'audit-output', 'capacity_autohome_evidence_20260927.json')
CACHE = os.path.join(HERE, 'audit-output', '_cap_evidence_cache.json')
RED = PatternFill('solid', fgColor='FFC7CE')
PH = re.compile(r'^(纯电动|插电式|插电|增程式?|换电式).{0,6}(轿车|多用途乘用车|乘用车|SUV|越野车)$')
# 特殊名家族 → autohome 检索词（09-21 别名先例同源）
QUERY_ALIAS = {
    '极星2': 'Polestar 2', '极星': 'Polestar', 'EU5': '北汽EU5', 'ES7': '蔚来ES7',
    'kiwi ev': '宝骏KiWi EV', 'kiwi': '宝骏KiWi EV', '精灵1': 'smart精灵#1',
    '精灵#1': 'smart精灵#1', '#1': 'smart精灵#1', 'mifa 9': '上汽大通MIFA 9',
    'mifa 6': '上汽大通MIFA 6', 'mifa 5': '上汽大通MIFA 5', '迈腾gte': '迈腾GTE',
    'magotan gte': '迈腾GTE', 'e·10x': '思皓E10X', 'e10x': '思皓E10X',
}


def query_for(base):
    k = base.lower().strip()
    return QUERY_ALIAS.get(k, QUERY_ALIAS.get(base, base))


def batch_ym():
    wb = openpyxl.load_workbook(WB, read_only=True)
    ws = wb['批次时间表']
    hdr = [str(c.value or '').strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    bi, di = hdr.index('批次'), hdr.index('公告日期')
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        m = re.match(r'^(\d{4})-(\d{2})', str(r[di] or ''))
        if r[bi] is not None and m:
            try:
                out[int(float(str(r[bi])))] = (int(m.group(1)), int(m.group(2)))
            except Exception:
                pass
    wb.close()
    return out


def spec_year(name):
    m = re.search(r'(\d{4})\s*款', str(name))
    return int(m.group(1)) if m else None


def fetch_all_capacity(seriesid):
    """getParamConf 全款型矩阵：[(year, 容量kWh, spec名)]；listSpec+1 次参数请求/车系
    （specid=0 返回空，必须带真实 specid，datalist 才含全系款型矩阵）"""
    ls = mf.list_spec(seriesid)
    specs = ls.get('specs') or []
    if not specs:
        return None, 'no_specs'
    u = (f"https://car-web-api.autohome.com.cn/car/param/getParamConf"
         f"?mode=1&site=2&seriesid={seriesid}&specid={specs[0]['id']}")
    d = mf.get(u)
    r = d.get('result') or {}
    tm = {}
    for grp in r.get('titlelist') or []:
        for it in grp.get('items') or []:
            tm[it.get('titleid')] = it.get('itemname')
    cap_title = None
    for tid, nm in tm.items():
        if nm and ('电池能量' in str(nm) or '电池容量' in str(nm)):
            cap_title = tid
            break
    if cap_title is None:
        return None, 'no_capacity_param'
    out = []
    for dl in r.get('datalist') or []:
        sid = dl.get('specid')
        sn = dl.get('specname') or ''
        val = None
        for pc in dl.get('paramconflist') or []:
            if pc.get('titleid') == cap_title:
                v = pc.get('itemname')
                if pc.get('sublist') and pc['sublist']:
                    v = pc['sublist'][0].get('name')
                val = v
                break
        if val is None or str(val).strip() in ('-', ''):
            continue
        try:
            fv = float(str(val).split('/')[0].replace(',', ''))
        except Exception:
            continue
        if 5 <= fv <= 150:
            out.append((spec_year(sn), fv, str(sn)[:40]))
    return out, None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--fetch', action='store_true')
    ap.add_argument('--apply', action='store_true')
    a = ap.parse_args()
    xf = json.load(open(XF, encoding='utf-8'))
    props = xf['proposals']
    ym = batch_ym()
    BY = {b: (y + (1 if m >= 7 else 0)) for b, (y, m) in ym.items()}

    if a.apply:
        rep = json.load(open(OUT, encoding='utf-8'))
        fixes = [p for p in rep['proposals'] if p.get('verdict') == 'FIX']
        print(f"落库提名 {len(fixes)} 格")
        wb = openpyxl.load_workbook(WB)
        ws = wb['NEV公告参数汇总']
        hdr = [str(c.value or '').strip() for c in ws[1]]
        hi = {h: i for i, h in enumerate(hdr)}
        led = wb['变更记录']
        seq = max((r[0] for r in led.iter_rows(min_row=2, values_only=True)
                   if isinstance(r[0], int)), default=0)
        n = 0
        for p in fixes:
            tgt = None
            for row in ws.iter_rows(min_row=2):
                if (str(row[hi['批次']].value).strip() == str(p['batch'])
                        and str(row[hi['产品型号']].value or '').strip() == p['model']):
                    tgt = row
                    break
            if tgt is None:
                continue
            cell = tgt[hi['电池容量(kWh)']]
            old = cell.value
            cell.value = p['fix_value']
            cell.fill = RED
            seq += 1
            led.append((seq, p['batch'], p['model'], p.get('power_type', ''), '电池容量(kWh)',
                        '媒体容量核证更正', str(old), str(p['fix_value']),
                        f"汽车之家车系参数（{p['series_name']}，sid={p['sid']}，检索 2026-09-27；"
                        f"代际窗口≤{p['cap_year']} 唯一值 {p['fix_value']}kWh）"))
            n += 1
        wb.save(WB)
        print(f"已落库 {n} 格（台账至 seq {seq}）——三件套+看板+提交")
        return

    # ---------- fetch ----------
    cache = {}
    if os.path.exists(CACHE):
        cache = json.load(open(CACHE, encoding='utf-8'))
    done = 0
    for p in props:
        gen = str(p['generic'] or '').strip()
        base = re.split(r'[,，、;；]', gen)[0].strip()
        if PH.search(base):
            base = str(p.get('model') or '')
        if not base or len(base) < 2:
            p['verdict'] = 'SKIP_NO_NAME'
            continue
        skey = base
        if skey not in cache:
            q = query_for(base)
            try:
                sug = mf.suggest(q)
            except Exception as e:
                cache[skey] = {'err': f'suggest:{e}'}
                time.sleep(2)
                continue
            time.sleep(0.55)
            cand = mf.find_series(sug, p.get('power_type', ''), name=base, brand=str(p.get('enterprise') or ''))
            if cand is None:
                cache[skey] = {'err': 'no_series_brand_guard'}
                continue
            try:
                sid = int(str(cand['wordid']))
            except Exception:
                cache[skey] = {'err': 'wordid_not_int'}
                continue
            time.sleep(0.55)
            try:
                caps, err = fetch_all_capacity(sid)
            except Exception as e:
                caps, err = None, f'param:{e}'
            cache[skey] = {'sid': sid, 'series_name': cand['key'], 'caps': caps, 'err': err}
            time.sleep(0.55)
        done += 1
        if done % 20 == 0:
            json.dump(cache, open(CACHE, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
            print(f'  …{done} 车系已取')
    json.dump(cache, open(CACHE, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

    # ---------- 提名判定 ----------
    nfix = nuniq = nmulti = nnosrc = 0
    for p in props:
        gen = str(p['generic'] or '').strip()
        base = re.split(r'[,，、;；]', gen)[0].strip()
        if PH.search(base):
            base = str(p.get('model') or '')
        c = cache.get(base) or {}
        p['series_name'] = c.get('series_name')
        p['sid'] = c.get('sid')
        caps = c.get('caps')
        if not caps:
            p['verdict'] = 'ABSTAIN_NO_SRC' if not c.get('err') else f"ABSTAIN_{c.get('err')}"
            nnosrc += 1
            continue
        b = int(p['batch']) if str(p['batch']).isdigit() else None
        cap_year = BY.get(b)
        if cap_year is None:
            p['verdict'] = 'ABSTAIN_NO_BATCH_YEAR'
            continue
        elig = {(y2, v) for (y2, v, _n) in caps if y2 is not None and y2 <= cap_year}
        if not elig:
            p['verdict'] = 'ABSTAIN_FUTURE_ONLY'
            continue
        vals = sorted({v for _y, v in elig})
        p['evidence'] = sorted({(y2, v) for y2, v in elig})[:12]
        if len(vals) != 1:
            p['verdict'] = 'ABSTAIN_MULTI'
            p['eligible_values'] = vals
            nmulti += 1
            continue
        v = vals[0]
        try:
            cur = float(str(p['cur_capacity']).split('/')[0])
        except Exception:
            cur = None
        if cur is not None and abs(cur - v) < 0.5:
            p['verdict'] = 'KEEP_SAME'
            nuniq += 1
            continue
        p['verdict'] = 'FIX'
        p['fix_value'] = v
        p['cap_year'] = cap_year
        nfix += 1
    rep = {'created': '2026-09-27', '纪律': '唯一外部证据才提名；代际窗口=公告月≥7允Y+1；禁止反推',
           'stats': {'fix': nfix, 'keep_same': nuniq, 'abstain_multi': nmulti, 'abstain_nosrc': nnosrc},
           'proposals': props}
    json.dump(rep, open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print(f"提名 FIX {nfix} | 同值 KEEP {nuniq} | 多值弃权 {nmulti} | 无源弃权 {nnosrc} → {OUT}")


if __name__ == '__main__':
    main()
