#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""cr17_004_005_feasibility.py — CR-17-004/005 可得性取证（只读，不写库）

回答：为什么 电池包供应商 2.04% / 官方参数页URL 2.06%？是"没做"还是"源里就没有"？
取证链：
  A. 变更记录扫描：历史是否有过这两字段的补全批次？（若历史 0 次 → 从未专项攻坚）
  B. 媒体校验来源侧信道：公告 PDF 里往往同页给出 电池包供应商/整车参数页，
     统计 媒体校验来源 中出现的可抽取模式（字段名/链接域），判断是否存在可批量抽取的富化源。
  C. 电芯供应商(30.23%) 作为对照：证明"公告侧确实带供应商字段"，不是字段本身不存在。
输出：audit-output/cr17_004_005_feasibility_20261005.json
"""
import json, os, re, collections
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
WB = os.path.join(HERE, 'NEV公告参数汇总表_合并版（341~410批）.xlsx')
OUT = os.path.join(HERE, 'audit-output', 'cr17_004_005_feasibility_20261005.json')

I_CELL, I_PACK, I_URL, I_MEDIA, I_BATCH = 31, 32, 33, 34, 0
SUP_HINT = re.compile(r'(电池包|pack|供应商|电池系统|电芯|BMS)', re.I)
PARAM_HINT = re.compile(r'(参数|配置|规格|官网|车型|product|car/spec|/car/|/models?/)', re.I)


def s(v):
    return '' if v is None else str(v).strip()


def main():
    wb = openpyxl.load_workbook(WB, read_only=True, data_only=True)

    # A. 变更记录溯源
    log = wb['变更记录']
    hits_pack, hits_url, per_type = collections.Counter(), collections.Counter(), collections.Counter()
    n_log = 0
    for r in log.iter_rows(min_row=2, values_only=True):
        if not r or r[0] is None:
            continue
        n_log += 1
        per_type[str(r[5] or '')] += 1
        fld = str(r[4] or '')
        if '电池包' in fld:
            hits_pack[fld] += 1
        if 'URL' in fld or '参数页' in fld:
            hits_url[fld] += 1

    # B/C. 底表侧信道
    ws = wb['NEV公告参数汇总']
    n = media = sup_hint = param_hint = 0
    media_batches = collections.Counter()
    cell_with_pack = 0
    cell_n = 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        n += 1
        m = s(r[I_MEDIA])
        if m:
            media += 1
            media_batches[str(r[I_BATCH])] += 1
            if SUP_HINT.search(m):
                sup_hint += 1
            if PARAM_HINT.search(m):
                param_hint += 1
        cs, ps = s(r[I_CELL]), s(r[I_PACK])
        if cs:
            cell_n += 1
            if ps:
                cell_with_pack += 1

    out = {
        'generated': '2026-10-05', 'records': n,
        'A_change_log': {'rows': n_log,
                         'field_histogram_top': dict(per_type.most_common(10)),
                         'battery_pack_related_records': dict(hits_pack),
                         'official_url_related_records': dict(hits_url)},
        'B_media_sidechannel': {
            'media_filled_rows': media, 'media_coverage': round(media / n * 100, 2),
            'rows_with_supplier_hint': sup_hint,
            'rows_with_param_page_hint': param_hint,
            'batches_with_media': len(media_batches),
            'recent_batches_media': dict(sorted(media_batches.items())[-12:]),
        },
        'C_control': {'cellSupplier_nonempty': cell_n,
                      'cell_coverage': round(cell_n / n * 100, 2),
                      'rows_with_both_cell_and_pack': cell_with_pack},
    }
    json.dump(out, open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print('变更记录总行', n_log)
    print('  电池包相关', dict(hits_pack))
    print('  官方URL相关', dict(hits_url))
    print('  变更类型Top', dict(per_type.most_common(6)))
    print(f"媒体校验来源非空 {media} ({out['B_media_sidechannel']['media_coverage']}%)"
          f" 供应商线索 {sup_hint}  参数页线索 {param_hint}")
    print('  近批媒体覆盖', out['B_media_sidechannel']['recent_batches_media'])
    print('电芯', cell_n, '其中也有电池包', cell_with_pack)
    print('->', OUT)


if __name__ == '__main__':
    main()