#!/usr/bin/env python3
"""Evidence-led NEV parameter enrichment. Copyright © 2026 David YE."""
from __future__ import annotations
import argparse, csv, datetime as dt, json, re, sys
from pathlib import Path
from openpyxl import load_workbook

COPYRIGHT = "Copyright © 2026 David YE"
ALIASES = {
    "model": ["产品型号", "车型型号"], "capacity": ["电池容量", "电池组总能量"],
    "range": ["纯电续航"], "consumption": ["百公里电耗"],
    "energy": ["电池组总能量", "动力蓄电池组总能量"], "mass": ["电池组总质量", "动力蓄电池组总质量"],
    "density": ["电池能量密度"],
}
RANGES = {"百公里电耗": (3, 40), "电池能量密度": (30, 400)}

def missing(v): return v is None or str(v).strip() in {"", "/", "-", "--"}
def number(v):
    m = re.search(r"-?\d+(?:\.\d+)?", str(v).replace(",", "")) if not missing(v) else None
    return float(m.group()) if m else None
def col(headers, aliases):
    return next((i for i, h in enumerate(headers) if any(a in h for a in aliases)), None)
def valid(field, value):
    lo_hi = RANGES.get(field)
    return value is not None and (not lo_hi or lo_hi[0] <= value <= lo_hi[1])

def main():
    p = argparse.ArgumentParser(description="TA-Analysis 证据台账驱动的数据补全")
    p.add_argument("--input", required=True); p.add_argument("--ledger", required=True); p.add_argument("--output", required=True)
    p.add_argument("--include-c-inference", action="store_true", help="允许写入 C 级推断候选；默认关闭")
    a = p.parse_args()
    wb = load_workbook(a.input); ws = wb["NEV公告参数汇总"] if "NEV公告参数汇总" in wb.sheetnames else wb.active
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    ix = {k: col(headers, v) for k, v in ALIASES.items()}
    if ix["model"] is None: raise ValueError("未找到产品型号列")
    rows = {str(ws.cell(r, ix["model"] + 1).value).strip(): r for r in range(2, ws.max_row + 1) if not missing(ws.cell(r, ix["model"] + 1).value)}
    if "【证据补充审计】" in wb.sheetnames: del wb["【证据补充审计】"]
    audit = wb.create_sheet("【证据补充审计】"); audit.append(["产品型号", "字段", "原值", "新值", "证据等级", "方法", "来源ID", "来源URL", "置信度", "状态", "时间", COPYRIGHT])
    applied = skipped = 0
    with open(a.ledger, encoding="utf-8-sig", newline="") as f:
        for e in csv.DictReader(f):
            model, field, tier = e.get("product_model", "").strip(), e.get("field", "").strip(), e.get("evidence_tier", "").upper()
            row = rows.get(model); target = col(headers, [field]); value = number(e.get("value"))
            approved = tier in {"A", "B"} or (tier == "C" and a.include_c_inference)
            ocr_ok = e.get("evidence_type", "").upper() != "OCR" or e.get("review_status", "").lower() == "verified"
            if not (row and target is not None and approved and ocr_ok and valid(field, value) and missing(ws.cell(row, target + 1).value)):
                skipped += 1; continue
            old = ws.cell(row, target + 1).value; ws.cell(row, target + 1).value = value
            audit.append([model, field, old, value, tier, e.get("method"), e.get("source_id"), e.get("source_url"), e.get("confidence"), "applied", dt.date.today().isoformat(), COPYRIGHT]); applied += 1
    # C-level deterministic candidates are auditable and opt-in; never overwrite declared data.
    for r in range(2, ws.max_row + 1):
        cap = number(ws.cell(r, ix["capacity"] + 1).value) if ix["capacity"] is not None else None; ran = number(ws.cell(r, ix["range"] + 1).value) if ix["range"] is not None else None
        if ix["consumption"] is not None and missing(ws.cell(r, ix["consumption"] + 1).value) and cap and ran:
            v = cap * 100 / ran
            audit.append([ws.cell(r, ix["model"] + 1).value, "百公里电耗", None, round(v, 3), "C", "容量×100÷纯电续航", "DERIVED-CONSUMPTION", "", "medium", "candidate" if not a.include_c_inference else "applied", dt.date.today().isoformat(), COPYRIGHT])
            if a.include_c_inference and valid("百公里电耗", v): ws.cell(r, ix["consumption"] + 1).value = round(v, 3); applied += 1
        energy = number(ws.cell(r, ix["energy"] + 1).value) if ix["energy"] is not None else None; mass = number(ws.cell(r, ix["mass"] + 1).value) if ix["mass"] is not None else None
        if ix["density"] is not None and missing(ws.cell(r, ix["density"] + 1).value) and energy and mass:
            v = energy * 1000 / mass
            audit.append([ws.cell(r, ix["model"] + 1).value, "电池能量密度", None, round(v, 3), "C", "电池组总能量×1000÷电池组总质量", "DERIVED-DENSITY", "", "high", "candidate" if not a.include_c_inference else "applied", dt.date.today().isoformat(), COPYRIGHT])
            if a.include_c_inference and valid("电池能量密度", v): ws.cell(r, ix["density"] + 1).value = round(v, 3); applied += 1
    audit.freeze_panes = "A2"; wb.save(a.output)
    print(json.dumps({"applied": applied, "skipped": skipped, "audit_sheet": audit.title, "copyright": COPYRIGHT}, ensure_ascii=False))

if __name__ == "__main__":
    try: main()
    except Exception as e: print(f"ERROR: {e}", file=sys.stderr); sys.exit(2)
