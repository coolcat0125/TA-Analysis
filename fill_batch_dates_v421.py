# -*- coding: utf-8 -*-
"""v4.2.1：用装备中心「公告发布」栏目官方日期回填批次时间表的推断行（只补空，不覆盖已有官方日期）。"""
import json
import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
SRC = "NEV公告参数汇总表_合并版（341~410批）.xlsx"
DATES = json.load(open("_scripts/batch_dates_official.json", encoding="utf-8"))["official"]

wb = load_workbook(SRC)
ws = wb["批次时间表"]
rows = list(ws.iter_rows(min_row=2, values_only=False))
hdr = [c.value for c in ws[1]]
print("表头:", hdr)
ci = {name: i for i, name in enumerate(hdr)}

updated, skipped_have, no_data = [], [], []
for row in rows:
    batch = row[ci["批次"]].value
    if batch is None:
        continue
    key = str(int(batch)) if not isinstance(batch, str) else batch.strip()
    precision = row[ci["精度"]].value
    cur_date = row[ci["公告日期"]].value
    if key not in DATES:
        no_data.append(key)
        continue
    if cur_date not in (None, "") or (precision and "官方" in str(precision)):
        skipped_have.append(key)
        continue
    date, url, _ = DATES[key]
    row[ci["公告年月"]].value = date[:7]
    row[ci["公告日期"]].value = date
    row[ci["精度"]].value = "官方精确日期"
    row[ci["说明/来源"]].value = f"工信部装备中心公告发布栏目挂网日期 {date}（{url}）"
    updated.append((key, date))

print(f"回填 {len(updated)} 行:", updated)
print(f"已有官方日期跳过 {len(skipped_have)} 行:", skipped_have)
print(f"无官方数据 {len(no_data)} 行:", no_data)
wb.save(SRC)
print("已保存")
