# -*- coding: utf-8 -*-
"""
底表通篇刷新 v3.7.0
====================
口径原则：保留免征购置税维度；所有补全均为可溯源的确定值，缺失留空不编造。

动作清单：
  A. 新增列「车长(mm)」(第31列，尾部追加，既有列索引不变) —— 占位待公告原件补充，本版不填值。
  B. 百公里电耗公式补全（全表范围重算）：ec = 电池容量(kWh) / 纯电续航(km) × 100，
     物理守卫 5~40；仅填空值，不覆盖已有值。填充沿用『公式补全』蓝色底纹(BDD7EE)，逐格记入变更记录。
  C. 数值类型归一化：文本型数字统一转为数值（不改数值本身，不入变更记录）。
用法：
  python refresh_workbook_v37.py            # 刷新 NEV公告参数汇总表_合并版（341~410批）.xlsx（就地）
Copyright © 2026 David YE
"""
import openpyxl
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, 'NEV公告参数汇总表_合并版（341~410批）.xlsx')
SHEET = 'NEV公告参数汇总'
BLUE = openpyxl.styles.PatternFill('solid', fgColor='BDD7EE')   # 公式补全蓝色底纹（与既有台账一致）

HDR = ['批次', '产品型号', '产品商标', '企业名称', '车型名称', '产品名称', '通用名称', '产品类型',
       '细分市场', '动力类型', '是否减免购置税', '整备质量(kg)', '轴距(mm)', '纯电续航里程(km)',
       '电池容量(kWh)', '电池类型', '电池能量密度(Wh/kg)', '百公里电耗(kWh/100km)', '电机峰值功率(kW)',
       '电机总功率(kW)', '电机生产企业', '电机功率/扭矩', '电机型号', '综合油耗(L/100km)',
       'B状态油耗(L/100km)', '发动机排量(mL)', '发动机功率(kW)', '发动机生产企业', '发动机型号', '数据来源']
NEW_COL_TITLE = '车长(mm)'


def main():
    print(f'[1/5] 加载工作簿: {os.path.basename(SRC)}（全量模式，较慢请稍候）')
    wb = openpyxl.load_workbook(SRC)

    ws = wb[SHEET]
    hdr = [c.value for c in ws[1]]
    assert hdr[:30] == HDR, f'表头口径不符，中止以保护数据: {hdr}'
    n_rows = ws.max_row
    idx = {h: i for i, h in enumerate(hdr)}
    c_batch, c_model = idx['批次'], idx['产品型号']
    c_cap, c_rng, c_ec = idx['电池容量(kWh)'] , idx['纯电续航里程(km)'], idx['百公里电耗(kWh/100km)']

    # ---------- A. 新增 车长(mm) ----------
    print('[2/5] 新增占位列「车长(mm)」…')
    if NEW_COL_TITLE not in hdr:
        col_new = ws.max_column + 1
        hc = ws.cell(1, col_new, NEW_COL_TITLE)
        hc.font = ws.cell(1, idx['轴距(mm)'] + 1).font.copy()
        ec_header_fill = ws.cell(1, idx['轴距(mm)'] + 1).fill
        if ec_header_fill and ec_header_fill.patternType:
            hc.fill = ec_header_fill.copy()
        print(f'      写入第 {col_new} 列表头（记录区留空待补充）')
    else:
        col_new = hdr.index(NEW_COL_TITLE) + 1
        print('      已存在，跳过')

    # 变更记录定位
    cr = wb['变更记录']
    last_seq = cr.max_row - 1          # 表头占用1行
    # 找真正最大序号（防中间有空行）
    for r in range(cr.max_row, 1, -1):
        v = cr.cell(r, 1).value
        if isinstance(v, int):
            last_seq = v
            break

    def ledger(batch, model, ptype, field, ctype, old, new, srcs):
        nonlocal last_seq
        last_seq += 1
        cr.append([last_seq, batch, model, ptype, field, ctype,
                   old if old is not None else '(空)', new, srcs])

    # ---------- B/C. 通篇刷新主循环 ----------
    print('[3/5] 通篇校验与公式补全…')
    n_ec_new = 0
    n_num_cast = 0
    num_fields = ['整备质量(kg)', '轴距(mm)', '纯电续航里程(km)', '电池容量(kWh)',
                  '电池能量密度(Wh/kg)', '百公里电耗(kWh/100km)', '电机峰值功率(kW)',
                  '电机总功率(kW)', '综合油耗(L/100km)', 'B状态油耗(L/100km)',
                  '发动机排量(mL)', '发动机功率(kW)']
    for ri in range(2, n_rows + 1):
        batch = ws.cell(ri, c_batch + 1).value
        model = ws.cell(ri, c_model + 1).value or ''
        ptype = str(ws.cell(ri, idx['动力类型'] + 1).value or '')

        # B. 电耗公式补全（BEV/PHEV 同口径，守卫 5~40 kWh/100km）
        cap, rng, ec = (ws.cell(ri, c_cap + 1).value,
                        ws.cell(ri, c_rng + 1).value,
                        ws.cell(ri, c_ec + 1).value)
        cap_f = cap if isinstance(cap, (int, float)) else None
        rng_f = rng if isinstance(rng, (int, float)) else None
        has_valid_ec = isinstance(ec, (int, float))
        if (ec is None or str(ec).strip() == '') and cap_f and rng_f and rng_f > 0:
            v = round(cap_f / rng_f * 100, 2)
            if 5 <= v <= 40:
                cell = ws.cell(ri, c_ec + 1, v)
                cell.fill = BLUE
                ledger(int(batch), model, ptype, '百公里电耗(kWh/100km)', '补充',
                       None, v, '公式补全·v3.7通篇刷新（容量/续航×100）')
                n_ec_new += 1

        # C. 文本数字 → 数值
        for fname in num_fields:
            cell = ws.cell(ri, idx[fname] + 1)
            v = cell.value
            if isinstance(v, str):
                s = v.strip()
                if s.replace('.', '', 1).isdigit():
                    fv = float(s) if '.' in s else int(s)
                    cell.value = fv
                    n_num_cast += 1

    # ---------- D. 汇总打印 ----------
    tx_col = ws.iter_rows(min_row=2, max_row=n_rows, min_col=idx['是否减免购置税']+1,
                          max_col=idx['是否减免购置税']+1, values_only=True)
    vals = [r[0] for r in tx_col]
    tx_yes = sum(1 for v in vals if str(v).strip() == '是')
    tx_no = sum(1 for v in vals if str(v).strip() == '否')
    print(f'      免征购置税维度保留：是 {tx_yes} 条 / 否 {tx_no} 条 / 未填报 {n_rows-1-tx_yes-tx_no} 条')
    print(f'      电耗公式新补全 {n_ec_new} 条（5~40 守卫内），均已记入变更记录')
    print(f'      数值类型归一化 {n_num_cast} 处')

    # ---------- E. 保存并复检 ----------
    print('[4/5] 保存工作簿…')
    final_name = os.path.basename(SRC)
    saved_as = None
    try:
        wb.save(SRC)
        saved_as = SRC
        print(f'      已就地覆盖 {final_name}')
    except PermissionError:
        tmp = SRC.replace('（341~410批）', '（341~410批）v37tmp')
        wb.save(tmp)
        print(f'      原文件被 Excel 占用，已先写入 {os.path.basename(tmp)}')
        import time
        for i in range(5):
            time.sleep(20)
            try:
                os.replace(tmp, SRC)
                saved_as = SRC
                print(f'      重试成功：已替换为 {final_name}')
                break
            except PermissionError:
                print(f'      替换重试 {i+1}/5：原文件仍被占用…')
        if not saved_as:
            kept = SRC.replace('（341~410批）', '（341~410批·刷新版）')
            os.replace(tmp, kept)
            saved_as = kept
            print(f'      原件持续被占用，刷新版保留为 {os.path.basename(kept)}（关闭Excel后可手动改名回主名）')
    print('[5/5] 复检…')
    wb2 = openpyxl.load_workbook(saved_as, read_only=True)
    ws2 = wb2[SHEET]
    hdr2 = [c.value for c in next(ws2.iter_rows(min_row=1, max_row=1))]
    ok_dim = (ws2.max_row == n_rows) and (ws2.max_column >= 31) and (hdr2[-1] == NEW_COL_TITLE)
    print(f'      维度 {ws2.max_row}行 × {ws2.max_column}列（{"OK" if ok_dim else "异常"}）; 尾列表头={hdr2[-1]}')
    cr2 = wb2['变更记录']
    print(f'      变更记录现有 {cr2.max_row - 1} 行（本版新增 {n_ec_new} 条）')
    wb2.close()
    assert ok_dim, '复检失败'
    print('✓ 底表刷新完成')


if __name__ == '__main__':
    sys.exit(main())
