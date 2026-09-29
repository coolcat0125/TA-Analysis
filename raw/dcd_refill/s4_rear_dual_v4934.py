# -*- coding: utf-8 -*-
"""s4_rear_dual_v4934.py — S4 扩展：双电机签名定向回填后电机（Wave 4，可复跑）

原理：行签名「前电机有值 + 后电机空 + 总功率>前功率 ⇒ 双电机车」——tied 款型中
只保留带 rear_electric_max_power 的款型，唯一值才补（配置身份已由驱动形式锚定）。
覆盖对象：dual_gap 64 行 + rear_only 中总>前 行 + S4 残池 359 中的后电机格。
铁律：只补空 / 唯一值 / 代际窗口(月≥7允Y+1) / 台账绿底 / W1 门禁复核。
用法：--scan 量化；--apply 落库（先 dry 审样例）。
"""
import argparse
import json
import os
import re
import time
from collections import defaultdict

import openpyxl
from openpyxl.styles import PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
WB = os.path.join(ROOT, 'NEV公告参数汇总表_合并版（341~410批）.xlsx')
REFILL = HERE
STATE = os.path.join(REFILL, 'batch_state.json')
FUEL = os.path.join(REFILL, 'dcd_fuel_form.json')
GREEN = PatternFill('solid', fgColor='C6EFCE')


def fnum(v):
    try:
        return float(str(v).split('/')[0].replace(',', '').replace('kW', '').strip())
    except Exception:
        return None


def year_of(c):
    try:
        return int(str(c.get('year')).strip()[:4])
    except Exception:
        return None


def load_batch_ym():
    wb = openpyxl.load_workbook(WB, read_only=True)
    ws = wb['批次时间表']
    hdr = [str(c.value or '').strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    bi, di = hdr.index('批次'), hdr.index('公告日期')
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        m = re.match(r'^(\d{4})-(\d{2})', str(r[di] or ''))
        if r[bi] is not None and m:
            try:
                out[int(float(str(r[bi])))] = (int(m.group(1)), int(m.group(2)))
            except Exception:
                pass
    wb.close()
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--scan', action='store_true')
    ap.add_argument('--apply', action='store_true')
    ap.add_argument('--verbose', action='store_true')
    a = ap.parse_args()
    if not a.scan and not a.apply:
        ap.print_help()
        return
    BYM = load_batch_ym()
    fuel_ff = {}
    try:
        fuel = json.load(open(FUEL, encoding='utf-8'))
        for sid_f, lst in fuel.items():
            for sp in (lst or []):
                if isinstance(sp, dict) and sp.get('id'):
                    fuel_ff[(str(sid_f), str(sp['id']))] = str(sp.get('ff') or '')
    except Exception:
        pass
    st = json.load(open(STATE, encoding='utf-8'))
    targets = {}
    for bid, b in sorted(st['batches'].items()):
        for s in b.get('series', []):
            if isinstance(s, dict) and s.get('sid'):
                kw = str(s['kw']).lower().replace(' ', '')
                if kw and kw not in targets:
                    targets[kw] = str(s['sid'])

    wb = openpyxl.load_workbook(WB)
    ws = wb['NEV公告参数汇总']
    hdr = [str(c.value or '').strip() for c in ws[1]]
    hi = {h: i for i, h in enumerate(hdr)}
    rows = list(ws.iter_rows(min_row=2))
    stats = defaultdict(int)
    cands = []
    samples = []
    for kw, sid in targets.items():
        fp = os.path.join(REFILL, f'dcd_s{sid}.json')
        if not os.path.exists(fp):
            continue
        dcd = json.load(open(fp, encoding='utf-8'))
        all_cars = dcd.get('cars', [])
        for row in rows:
            g = (str(row[hi['通用名称']].value or '') + str(row[hi['车型名称']].value or '')).lower().replace(' ', '')
            if kw not in g:
                continue
            pt = str(row[hi['动力类型']].value or '')
            if not ('纯电' in pt or '插电' in pt or '增程' in pt or 'BEV' in pt or 'PHEV' in pt):
                continue
            is_bev = '纯电' in pt or 'BEV' in pt
            code = str(row[hi['产品型号']].value or '').strip()
            rear_cell = row[hi['后电机功率/扭矩']]
            if rear_cell.value is not None and str(rear_cell.value).strip() != '':
                continue  # 只补空
            front = fnum(row[hi['前电机功率/扭矩']].value)
            total = fnum(row[hi['电机总功率/扭矩']].value)
            if front is None or total is None or total <= front:
                continue  # 双电机签名不成立（总>前 才是双电机）
            stats['dual_rows'] += 1
            try:
                bno = int(float(str(row[hi['批次']].value or '').strip()))
            except Exception:
                continue
            binfo = BYM.get(bno)
            if not binfo:
                continue
            cap_year = binfo[0] + (1 if binfo[1] >= 7 else 0)
            trims = []
            for c in all_cars:
                v = c.get('vals', {})
                if fuel_ff.get((sid, str(c.get('id'))), '') == '汽油':
                    continue
                if is_bev and not v.get('battery_capacity'):
                    continue
                if not v.get('rear_electric_max_power'):
                    continue  # 双电机款型必须有后电机值
                trims.append(c)
            if not trims:
                stats['no_dual_trim'] += 1
                continue
            r_rng = fnum(row[hi['纯电续航里程(km)']].value)
            r_cw = fnum(row[hi['整备质量(kg)']].value)
            scored = []
            for c in trims:
                v = c['vals']
                nm = re.search(r'(\d{3,4})\s*km', str(c.get('name')))
                cltc = fnum(v.get('cltc_recharge_mileage')) or (float(nm.group(1)) if nm else None)
                cw = fnum(v.get('curb_weight'))
                sc = 0
                if r_rng and cltc:
                    d0 = abs(r_rng - cltc)
                    sc += 0 if d0 <= 15 else (1 if d0 <= 30 else 3)
                if r_cw and cw:
                    d0 = abs(r_cw - cw)
                    sc += 0 if d0 <= 40 else (1 if d0 <= 90 else 3)
                scored.append((sc, c))
            scored.sort(key=lambda t: t[0])
            best = scored[0][0]
            tied = [c for sc, c in scored if sc == best]
            elig = [c for c in tied if year_of(c) is not None and year_of(c) <= cap_year]
            if elig:
                ty = max(year_of(c) for c in elig)
                tied2 = [c for c in elig if year_of(c) == ty]
            else:
                tied2 = tied
            vals = set()
            for c in tied2:
                m = re.match(r'^\s*(\d+(?:\.\d+)?)(?:\s*kW)?', str(c['vals'].get('rear_electric_max_power')), re.I)
                if m:
                    vals.add(float(m.group(1)))
            if len(vals) == 1:
                val = next(iter(vals))
                # 双电机自洽：总 = 前 + 后（文本首值口径，±2 容差）
                if abs(total - (front + val)) <= 2.0:
                    cands.append({'row_idx': row[0].row, 'batch': bno, 'model': code, 'pt': pt,
                                  'val': val, 'sid': sid, 'kw': kw,
                                  'front': front, 'total': total})
                    stats['fillable'] += 1
                    if a.verbose and len(samples) < 20:
                        samples.append(f'{bno}批 {code} [{kw}] 后电机→{val}（前{front}+后{val}=总{total}）')
                else:
                    stats['sum_mismatch'] += 1
            elif not vals:
                stats['no_val'] += 1
            else:
                stats['still_ambig'] += 1
    print(f"== S4-rear-dual scan ==  双电机行 {stats['dual_rows']} | 可填 {stats['fillable']} | "
          f"无双电机款型 {stats['no_dual_trim']} | 求和不符 {stats['sum_mismatch']} | 仍歧义 {stats['still_ambig']}")
    for s in samples[:12]:
        print(' ', s)
    if not a.apply:
        print('DRY-RUN（未写入）')
        return
    led = wb['变更记录']
    seq = max((r[0] for r in led.iter_rows(min_row=2, values_only=True)
               if isinstance(r[0], int)), default=0)
    n = 0
    seen = set()
    for cd in cands:
        key = (cd['row_idx'],)
        if key in seen:
            continue
        seen.add(key)
        row = ws[cd['row_idx']]
        cell = row[hi['后电机功率/扭矩']]
        if cell.value is not None and str(cell.value).strip() != '':
            continue
        cell.value = f"{cd['val']:g}"
        cell.fill = GREEN
        seq += 1
        led.append((seq, cd['batch'], cd['model'], cd['pt'], '后电机功率/扭矩', '媒体参数补全',
                    '(空)', f"{cd['val']:g}",
                    f"懂车帝参数页(sid={cd['sid']})（懂车帝参数页，双电机签名定向 检索 {time.strftime('%Y-%m-%d')}；"
                    f"前{cd['front']:g}+后{cd['val']:g}=总{cd['total']:g} 自洽）"))
        n += 1
    if not n:
        print('无可写格')
        return
    if a.apply and os.environ.get('DRY1') == '1':
        print(f'DRY 预演：应写 {n} 格（未写盘）')
        return
    wb.save(WB)
    print(f"已落库 {n} 格（台账 seq {seq-n+1}~{seq}）——三件套+看板+提交")


if __name__ == '__main__':
    main()
