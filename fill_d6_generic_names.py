# -*- coding: utf-8 -*-
"""fill_d6_generic_names.py — D6：占位符通用名称定向回填（官方能耗源 20260915）

源：工信部新能源车型能耗数据20260915.xlsx（新版+旧版双表，车辆型号→通用名称）
靶：底表中 通用名称=动力类型泄漏值（PLACEHOLDER_GENERIC_RE 同款判定）的行
规则：型号精确匹配到官方通用名称且与占位符不同→回填（绿底+台账）；无源维持
"""
import json
import os
import re
import sys

import openpyxl
from openpyxl.styles import PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
WB = os.path.join(HERE, 'NEV公告参数汇总表_合并版（341~410批）.xlsx')
SRC = os.path.join(HERE, '工信部新能源车型能耗数据20260915.xlsx')
OUT = os.path.join(HERE, 'audit-output', 'd6_generic_fill_report.json')

GREEN = PatternFill('solid', fgColor='C6EFCE')

PLACEHOLDER_GENERIC_RE = re.compile(
    r'^(纯电动|插电式|插电混动|增程式|增程|混合动力|燃料电池|混动|纯|双层)?'
    r'(轿车|SUV|SUV车|MPV|MPV车|多用途乘用车|乘用车|运动型乘用车|客车|货车|卡车|底盘|用车)$')


def is_ph(g):
    s = str(g or '').strip()
    return bool(s) and bool(PLACEHOLDER_GENERIC_RE.match(s))


def norm_code(s):
    return re.sub(r'[\s\-]+', '', str(s or '').strip()).upper()


def main():
    apply = '--apply' in sys.argv
    # 1) 官方源 型号→通用名称
    src = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    code2name = {}
    for sn in src.sheetnames:
        ws = src[sn]
        it = ws.iter_rows(min_row=1, values_only=True)
        hdr = [str(h or '').strip() for h in next(it)]
        try:
            ic, im, ig = hdr.index('车辆型号'), hdr.index('生产企业'), hdr.index('通用名称')
        except ValueError:
            continue
        for r in it:
            code = norm_code(r[ic])
            name = str(r[ig] or '').strip()
            if code and name and not is_ph(name):
                code2name.setdefault(code, name)
    src.close()
    print(f'official code->name map: {len(code2name)}')

    # 2) 底表占位行
    wb = openpyxl.load_workbook(WB)
    ws = wb['NEV公告参数汇总']
    hdr = [str(c.value or '').strip() for c in ws[1]]
    hi = {h: i for i, h in enumerate(hdr)}
    ledger = wb['变更记录']
    seq = max((r[0] for r in ledger.iter_rows(min_row=2, values_only=True)
               if isinstance(r[0], int)), default=0)

    filled = noplay = 0
    samples = []
    for row in ws.iter_rows(min_row=2):
        g = row[hi['通用名称']].value
        if not is_ph(g):
            continue
        code = norm_code(row[hi['产品型号']].value)
        name = code2name.get(code)
        if not name:
            noplay += 1
            continue
        if name == str(g).strip():
            noplay += 1
            continue
        row[hi['通用名称']].value = name
        row[hi['通用名称']].fill = GREEN
        seq += 1
        ledger.append((seq, row[hi['批次']].value, row[hi['产品型号']].value or '',
                       row[hi['动力类型']].value or '', '通用名称', '官方源定向回填',
                       str(g), name, 'D6：官方能耗源20260915 通用名称（型号精确匹配）'))
        filled += 1
        if len(samples) < 10:
            samples.append((row[hi['批次']].value, row[hi['产品型号']].value, str(g), name))

    print(f'filled={filled} unmatched/kept={noplay} ledger_seq->{seq}')
    json.dump({'filled': filled, 'kept': noplay, 'samples': samples},
              open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    for s in samples:
        print(' ', s)
    if not apply:
        print('DRY-RUN（未写入）')
        return
    wb.save(WB)
    print('saved')


if __name__ == '__main__':
    main()
