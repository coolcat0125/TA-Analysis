# -*- coding: utf-8 -*-
"""build_deliverables_v4923.py — 下游交付①·45表分析底表重建（v4.9.23 口径，可复跑）

源：NEV公告参数汇总表_合并版（341~410批）.xlsx（权威底表 5,392 行 × 39 列）
旧版基线快照（v4.3.6 口径 4,613 行）：audit-output/old_deliverable_45sheets_snapshot_20260927.json
输出：NEV公告车型行业分析_数据底表.xlsx（45 表，全部由当前底表重算）

口径要点：
  · 动力分组 BEV=纯电/换电；PHEV=插电/增程（与看板一致）
  · 年度=批次公告日期年份（批次时间表·官方精确日期）
  · 电耗/续航=工信部能耗数据20260915 官方对齐值（v4.9.22）；工况 BEV=CLTC、PHEV/EREV=WLTC
  · 电机功率取「电机总功率/扭矩」首值；仅乘用车纯净口径（GB/T 3730.1）
  · 各表样本=该表所需字段均有值的行（有值子集），与旧版一致
"""
import datetime
import math
import os
import re
import shutil
import statistics as st
from collections import Counter, defaultdict

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = '.'
SRC = 'NEV公告参数汇总表_合并版（341~410批）.xlsx'
OUT = 'NEV公告车型行业分析_数据底表.xlsx'
VER = 'v4.9.23'

THIN = Side(style='thin', color='B0B7C3')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HFILL = PatternFill('solid', fgColor='D9E1F2')
TFONT = Font(bold=True, size=13, color='1F3864')
HFONT = Font(bold=True, size=10)
YELLOW = PatternFill('solid', fgColor='FFF2CC')


def fnum(v):
    try:
        return float(str(v).split('/')[0].replace(',', '').strip())
    except Exception:
        return None


def load_records():
    wb = openpyxl.load_workbook(SRC, read_only=True)
    ws = wb['NEV公告参数汇总']
    hdr = [str(c.value or '').strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    hi = {h: i for i, h in enumerate(hdr)}
    # 批次→公告年
    bt = wb['批次时间表']
    bth = [str(c.value or '').strip() for c in next(bt.iter_rows(min_row=1, max_row=1))]
    bi, di = bth.index('批次'), bth.index('公告日期')
    batch_year = {}
    for r in bt.iter_rows(min_row=2, values_only=True):
        m = re.match(r'^(\d{4})', str(r[di] or ''))
        if r[bi] is not None and m:
            try:
                batch_year[int(float(str(r[bi])))] = int(m.group(1))
            except Exception:
                pass
    recs = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        def g(name):
            return r[hi[name]] if name in hi else None
        pt = str(g('动力类型') or '').strip()
        if '纯电' in pt or '换电' in pt or pt == 'BEV':
            grp = 'BEV'
        elif '插电' in pt or '增程' in pt or pt == 'PHEV':
            grp = 'PHEV'
        else:
            continue
        bno = g('批次')
        try:
            bno = int(float(str(bno)))
        except Exception:
            continue
        recs.append({
            '批次': bno, '年': batch_year.get(bno), '组': grp,
            '商标': str(g('产品商标') or '').strip(),
            '企业': str(g('企业名称') or '').strip(),
            '通用': str(g('通用名称') or '').strip(),
            '细分': str(g('细分市场') or '').strip(),
            '续航': fnum(g('纯电续航里程(km)')),
            '容量': fnum(g('电池容量(kWh)')),
            '密度': fnum(g('电池能量密度(Wh/kg)')),
            '电耗': fnum(g('百公里电耗(kWh/100km)')),
            '油耗': fnum(g('综合油耗(L/100km)')),
            '整备': fnum(g('整备质量(kg)')),
            '功率': fnum(g('电机总功率/扭矩')),
            '排量': fnum(g('发动机排量(mL)')),
            '电池类型': str(g('电池类型') or '').strip(),
            '电芯': str(g('电芯供应商') or '').strip(),
            '电机企': str(g('电机生产企业') or '').strip(),
            '月销': fnum(g('月销量(辆)')),
        })
    wb.close()
    return recs


def pct(vals, p):
    if not vals:
        return None
    s = sorted(vals)
    k = (len(s) - 1) * p / 100.0
    f = math.floor(k)
    c = min(f + 1, len(s) - 1)
    return s[f] + (s[c] - s[f]) * (k - f)


def year_rows(recs, grp=None):
    d = defaultdict(list)
    for r in recs:
        if r['年'] and (grp is None or r['组'] == grp):
            d[r['年']].append(r)
    return d


def agg_by_year(recs, field, grp=None, filt=None, scale=1.0):
    """年 → (均值, 中位数, 样本数)；filt=行谓词；scale 换算"""
    out = {}
    for y, rs in sorted(year_rows(recs, grp).items()):
        vals = [r[field] * scale for r in rs
                if r[field] is not None and (filt is None or filt(r))]
        if vals:
            out[y] = (sum(vals) / len(vals), st.median(vals), len(vals))
    return out


def per_row_ratio(recs, a, b, grp):
    """per-row a/b×1000 均值（分子分母均有值的子集）"""
    out = {}
    for y, rs in sorted(year_rows(recs, grp).items()):
        vals = [r[a] / r[b] * 1000 for r in rs
                if r[a] is not None and r[b] not in (None, 0)]
        if vals:
            out[y] = sum(vals) / len(vals)
    return out


def pearson(xs, ys):
    n = len(xs)
    if n < 3:
        return None
    mx, my = sum(xs) / n, sum(ys) / n
    sx = math.sqrt(sum((x - mx) ** 2 for x in xs))
    sy = math.sqrt(sum((y - my) ** 2 for y in ys))
    if sx == 0 or sy == 0:
        return None
    return sum((x - mx) * (y - my) for x, y in zip(xs, ys)) / (sx * sy)


class Book:
    def __init__(self):
        self.wb = openpyxl.Workbook()
        self.wb.remove(self.wb.active)

    def sheet(self, name, title=None):
        ws = self.wb.create_sheet(name)
        if title:
            ws.cell(1, 1, title).font = TFONT
        return ws

    def table(self, ws, start_row, header, rows, num_fmt=None, widths=None):
        """写一张带表头样式的小表；rows=二维数组（值或 (值,fmt)）"""
        for j, h in enumerate(header, 1):
            c = ws.cell(start_row, j, h)
            c.font = HFONT
            c.fill = HFILL
            c.border = BORDER
            c.alignment = Alignment(horizontal='center')
        for i, row in enumerate(rows, start_row + 1):
            for j, v in enumerate(row, 1):
                c = ws.cell(i, j, v)
                c.border = BORDER
                if isinstance(v, float):
                    c.number_format = num_fmt or '0.0'
                    if abs(v) >= 1000:
                        c.number_format = '#,##0'
        if widths:
            for j, w in enumerate(widths, 1):
                ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w


def build(recs):
    b = Book()
    total = len(recs)
    n_bev = sum(1 for r in recs if r['组'] == 'BEV')
    n_phev = total - n_bev
    today = datetime.date.today().isoformat()

    # ---------- 1 封面 ----------
    ws = b.sheet('封面')
    ws['B2'] = 'NEV 公告车型行业分析 · 数据底表'
    ws['B2'].font = Font(bold=True, size=18, color='1F3864')
    ws['B4'] = f'口径版本 {VER} · {total:,} 款 × 39 列（第 341~411 批 · 纯乘用车）'
    ws['B5'] = f'生成日期 {today} · 权威源 NEV公告参数汇总表_合并版（341~410批）.xlsx'
    ws['B6'] = '电耗/续航已与工信部能耗数据20260915 官方对齐 · 销量参考=易车零售月度口径'
    ws['B7'] = '分析表均由权威底表程序化重算（生成器 build_deliverables_v4923.py，可复跑）'
    for r in (2, 4, 5, 6, 7):
        ws.cell(r, 2).alignment = Alignment(vertical='center')
    ws.column_dimensions['B'].width = 96

    # ---------- 2 目录 ----------
    ws = b.sheet('目录', '目录')
    names = ['封面', '目录', '数据补充说明', '执行摘要', '数据总览', '年度分布', 'BEV续航分析',
             'PHEV续航分析', '续航对比分析', '电池容量分析', 'BEV容量分布', 'PHEV容量分布',
             '电池类型趋势-BEV', '电池类型趋势-PHEV', '电池能量密度', '续航分布-BEV',
             '续航分布-PHEV', '电机功率分析', '整备质量分析', '功率质量比', '整车能量密度',
             '能效分析', 'PHEV油耗分析', 'PHEV油耗-排量年份', '技术路线矩阵', '供应商集中度',
             '电池供应商CR', '电机供应商CR', '电池供应商排名', '电机供应商排名', '供应商组合',
             '品牌-电池供应商', '品牌-电机供应商', '细分市场-BEV', '细分市场-PHEV',
             '细分市场续航对比', '细分市场容量对比', '质量按细分市场', '相关性-BEV',
             '相关性-PHEV', '同比增长-BEV', '同比增长-PHEV', '关键指标汇总-BEV',
             '关键指标汇总-PHEV', '年度综合指标', '月销量分析']
    b.table(ws, 3, ['序', '工作表'], list(enumerate(names, 1)), widths=[6, 26])

    # ---------- 3 数据补充说明 ----------
    ws = b.sheet('数据补充说明', '数据补充说明')
    notes = [
        '① 口径双原则：仅乘用车（GB/T 3730.1 M1 类）；仅主流车企品牌。',
        '② 动力分组：BEV=纯电动/换电式纯电动；PHEV=插电式混合动力/插电式增程（与看板一致）。',
        '③ 年度=批次《公告》发布日期年份（批次时间表·官方精确日期 341~410 批；411 批为公示暂计入 2026）。',
        '④ 电耗/续航口径：以工信部新能源车型能耗数据20260915 为准（v4.9.22 全量对齐）；',
        '    工况按动力类型推定：BEV=CLTC，PHEV/EREV=WLTC；单元格不再保留工况标注。',
        '⑤ 电机功率取「电机总功率/扭矩」首值（kW）；功率质量比/整车能量密度为逐行比值均值（有值子集）。',
        '⑥ 各表样本=该表所需字段均有值的行；样本数随底表覆盖率滚动变化。',
        '⑦ 销量参考=易车零售量月度口径（2026-08，160 车系，仅参考列）。',
        '⑧ 供应商排名/CR/HHI 基于电芯供应商、电机生产企业有值子集，覆盖率有限仅作趋势参考。',
        '⑨ 旧版（v4.3.6 口径 4,613 款）快照存 audit-output/old_deliverable_45sheets_snapshot_20260927.json。',
    ]
    for i, t in enumerate(notes, 3):
        ws.cell(i, 1, t)
    ws.column_dimensions['A'].width = 100

    # ---------- 4 执行摘要 ----------
    yrs = sorted({r['年'] for r in recs if r['年']})
    yy = yrs[-1]
    bev = [r for r in recs if r['组'] == 'BEV']
    phev = [r for r in recs if r['组'] == 'PHEV']
    bev_y = [r for r in bev if r['年'] == yy]
    phev_y = [r for r in phev if r['年'] == yy]
    mean = lambda rs, f: (lambda v: round(sum(v) / len(v), 2) if v else None)(
        [r[f] for r in rs if r[f] is not None])
    ws = b.sheet('执行摘要', '执行摘要 - 关键指标')
    ws.cell(3, 1, '（结论详见演示文稿）')
    kpis = [
        ['指标', '数值', '单位', '备注'],
        ['公告车型总量', total, '款', f'BEV {n_bev:,} / PHEV {n_phev:,}'],
        [f'{yy}年公告车型', len(bev_y) + len(phev_y), '款',
         f'BEV {len(bev_y)} / PHEV {len(phev_y)}'],
        [f'{yy}年PHEV占比', round(len(phev_y) / max(1, len(bev_y) + len(phev_y)) * 100, 1), '%',
         '插混/增程持续放量'],
        [f'BEV平均续航({yy})', mean(bev_y, '续航'), 'km',
         f'较{yrs[0]}年显著提升'],
        [f'BEV平均电池容量({yy})', mean(bev_y, '容量'), 'kWh', '大电量化持续'],
        [f'PHEV平均纯电续航({yy})', mean(phev_y, '续航'), 'km', '插混长途化'],
        [f'BEV平均百公里电耗({yy})', mean(bev_y, '电耗'), 'kWh/100km', '能效持续优化'],
        ['月销量覆盖(2026-08)', sum(1 for r in recs if r['月销']), '行', '易车零售口径·参考'],
    ]
    b.table(ws, 4, kpis[0], kpis[1:], widths=[26, 12, 12, 34])

    # ---------- 5 数据总览 ----------
    ws = b.sheet('数据总览', '数据总览')
    ws.cell(3, 1, '按动力类型').font = HFONT
    b.table(ws, 4, ['动力类型', '车型数', '占比%'],
            [['BEV', n_bev, round(n_bev / total * 100, 1)],
             ['PHEV', n_phev, round(n_phev / total * 100, 1)]], widths=[12, 10, 10])
    ws.cell(8, 1, '按年度').font = HFONT
    yr_rows = []
    for y in yrs:
        nb = sum(1 for r in recs if r['年'] == y and r['组'] == 'BEV')
        np_ = sum(1 for r in recs if r['年'] == y and r['组'] == 'PHEV')
        yr_rows.append([y, nb, np_, nb + np_, round(np_ / max(1, nb + np_) * 100, 1)])
    b.table(ws, 9, ['年份', 'BEV', 'PHEV', '合计', 'PHEV占比%'], yr_rows)

    # ---------- 6 年度分布 ----------
    ws = b.sheet('年度分布', '年度车型数量分布')
    ws.cell(3, 1, '（含411批公示行计入2026）')
    b.table(ws, 4, ['年份', 'BEV(款)', 'PHEV(款)', '合计(款)', 'PHEV占比%'], yr_rows,
            widths=[8, 10, 10, 10, 12])

    # ---------- 7/8 续航分析 ----------
    for name, grp in (('BEV续航分析', 'BEV'), ('PHEV续航分析', 'PHEV')):
        ws = b.sheet(name, f'{grp} 续航里程年度分析')
        rows = []
        for y in yrs:
            vals = sorted(r['续航'] for r in recs
                          if r['年'] == y and r['组'] == grp and r['续航'] is not None)
            if vals:
                rows.append([y, round(sum(vals) / len(vals), 2), round(st.median(vals), 1),
                             round(pct(vals, 25), 0), round(pct(vals, 75), 0),
                             round(pct(vals, 90), 0), len(vals)])
        b.table(ws, 3, ['年份', '均值(km)', '中位数(km)', 'P25(km)', 'P75(km)', 'P90(km)', '样本数'],
                rows, widths=[8, 10, 11, 9, 9, 9, 8])

    # ---------- 9 续航对比 ----------
    ws = b.sheet('续航对比分析', 'BEV vs PHEV 续航年均值')
    mb, mp_ = agg_by_year(recs, '续航', 'BEV'), agg_by_year(recs, '续航', 'PHEV')
    b.table(ws, 3, ['年份', 'BEV均值(km)', 'PHEV均值(km)', '差值(km)'],
            [[y, round(mb[y][0], 2), round(mp_[y][0], 2), round(mb[y][0] - mp_[y][0], 1)]
             for y in yrs if y in mb and y in mp_])

    # ---------- 10 电池容量 ----------
    ws = b.sheet('电池容量分析', '电池容量年度分析')
    cb, cp = agg_by_year(recs, '容量', 'BEV'), agg_by_year(recs, '容量', 'PHEV')
    b.table(ws, 3, ['年份', 'BEV均值(kWh)', 'PHEV均值(kWh)'],
            [[y, round(cb[y][0], 2), round(cp[y][0], 2)] for y in yrs if y in cb and y in cp])

    # ---------- 11/12 容量分布 ----------
    CAP_BINS = [('<20', 0, 20), ('20-30', 20, 30), ('30-40', 30, 40), ('40-50', 40, 50),
                ('50-60', 50, 60), ('60-70', 60, 70), ('70-80', 70, 80), ('≥80', 80, 1e9)]
    CAP_BINS_P = [('<20', 0, 20), ('20-30', 20, 30), ('30-40', 30, 40),
                  ('40-50', 40, 50), ('≥50', 50, 1e9)]
    for name, grp, bins in (('BEV容量分布', 'BEV', CAP_BINS), ('PHEV容量分布', 'PHEV', CAP_BINS_P)):
        ws = b.sheet(name, f'{grp} 电池容量分布（按年度）')
        ws.cell(3, 1, '按容量区间计数（kWh，区间左闭右开）')
        hdr = ['年份'] + [b0 for b0, _, _ in bins]
        body = []
        for y in yrs:
            vals = [r['容量'] for r in recs
                    if r['年'] == y and r['组'] == grp and r['容量'] is not None]
            if vals:
                body.append([y] + [sum(1 for v in vals if lo <= v < hi) for _, lo, hi in bins])
        b.table(ws, 4, hdr, body)

    # ---------- 13/14 电池类型趋势 ----------
    def btype(v):
        if '磷酸铁锂' in v:
            return 'LFP'
        if '三元' in v:
            return 'NCM'
        return '其他'
    for name, grp in (('电池类型趋势-BEV', 'BEV'), ('电池类型趋势-PHEV', 'PHEV')):
        ws = b.sheet(name, f'{grp} 电池类型趋势')
        ws.cell(3, 1, '数量与占比')
        body = []
        for y in yrs:
            rs = [r for r in recs if r['年'] == y and r['组'] == grp and r['电池类型']]
            if not rs:
                continue
            c = Counter(btype(r['电池类型']) for r in rs)
            n = len(rs)
            body.append([y, c.get('LFP', 0), c.get('NCM', 0), c.get('其他', 0),
                         round(c.get('LFP', 0) / n * 100, 1),
                         round(c.get('NCM', 0) / n * 100, 1),
                         round(c.get('其他', 0) / n * 100, 1)])
        b.table(ws, 4, ['年份', 'LFP(款)', 'NCM(款)', '其他(款)', 'LFP占比%', 'NCM占比%', '其他占比%'],
                body)

    # ---------- 15 电池能量密度 ----------
    ws = b.sheet('电池能量密度', '电池能量密度年度趋势')
    db, dp = agg_by_year(recs, '密度', 'BEV'), agg_by_year(recs, '密度', 'PHEV')
    b.table(ws, 3, ['年份', 'BEV均值(Wh/kg)', 'PHEV均值(Wh/kg)'],
            [[y, round(db[y][0], 1), round(dp[y][0], 1)] for y in yrs if y in db and y in dp])

    # ---------- 16/17 续航分布 ----------
    RAN_B = [('<250', 0, 250), ('250-350', 250, 350), ('350-450', 350, 450),
             ('450-550', 450, 550), ('550-650', 550, 650), ('650-750', 650, 750), ('≥750', 750, 1e9)]
    RAN_P = [('<100', 0, 100), ('100-150', 100, 150), ('150-200', 150, 200),
             ('200-250', 200, 250), ('≥250', 250, 1e9)]
    for name, grp, bins in (('续航分布-BEV', 'BEV', RAN_B), ('续航分布-PHEV', 'PHEV', RAN_P)):
        ws = b.sheet(name, f'{grp} 续航里程分布（按年度）')
        ws.cell(3, 1, '按续航区间计数（km，区间左闭右开）')
        hdr = ['年份'] + [b0 for b0, _, _ in bins]
        body = []
        for y in yrs:
            vals = [r['续航'] for r in recs
                    if r['年'] == y and r['组'] == grp and r['续航'] is not None]
            if vals:
                body.append([y] + [sum(1 for v in vals if lo <= v < hi) for _, lo, hi in bins])
        b.table(ws, 4, hdr, body)

    # ---------- 18 电机功率 ----------
    ws = b.sheet('电机功率分析', '电机功率年度分析')
    pb, pp = agg_by_year(recs, '功率', 'BEV'), agg_by_year(recs, '功率', 'PHEV')
    b.table(ws, 3, ['年份', 'BEV均值(kW)', 'PHEV均值(kW)'],
            [[y, round(pb[y][0], 2), round(pp[y][0], 2)] for y in yrs if y in pb and y in pp])

    # ---------- 19 整备质量 ----------
    ws = b.sheet('整备质量分析', '整备质量年度分析')
    wb_, wp_ = agg_by_year(recs, '整备', 'BEV'), agg_by_year(recs, '整备', 'PHEV')
    b.table(ws, 3, ['年份', 'BEV均值(kg)', 'PHEV均值(kg)'],
            [[y, round(wb_[y][0], 1), round(wp_[y][0], 1)] for y in yrs if y in wb_ and y in wp_])

    # ---------- 20 功率质量比 ----------
    ws = b.sheet('功率质量比', '功率质量比年度趋势')
    rb, rp = per_row_ratio(recs, '功率', '整备', 'BEV'), per_row_ratio(recs, '功率', '整备', 'PHEV')
    b.table(ws, 3, ['年份', 'BEV(kW/kg×1000)', 'PHEV(kW/kg×1000)'],
            [[y, round(rb[y], 1), round(rp[y], 1)] for y in yrs if y in rb and y in rp])

    # ---------- 21 整车能量密度 ----------
    ws = b.sheet('整车能量密度', '整车能量密度年度趋势')
    eb, ep = per_row_ratio(recs, '容量', '整备', 'BEV'), per_row_ratio(recs, '容量', '整备', 'PHEV')
    b.table(ws, 3, ['年份', 'BEV(Wh/kg)', 'PHEV(Wh/kg)'],
            [[y, round(eb[y], 1), round(ep[y], 1)] for y in yrs if y in eb and y in ep])

    # ---------- 22 能效分析 ----------
    ws = b.sheet('能效分析', 'BEV 百公里电耗年度分析')
    ee = agg_by_year(recs, '电耗', 'BEV')
    b.table(ws, 3, ['年份', '均值(kWh/100km)', '中位数', '样本数'],
            [[y, round(ee[y][0], 2), round(ee[y][1], 2), ee[y][2]] for y in yrs if y in ee])

    # ---------- 23 PHEV 油耗 ----------
    ws = b.sheet('PHEV油耗分析', 'PHEV 油耗与纯电续航年度分析')
    rows = []
    for y in yrs:
        oil = [r['油耗'] for r in recs if r['年'] == y and r['组'] == 'PHEV'
               and r['油耗'] is not None and 0.1 <= r['油耗'] <= 12]
        rng = [r['续航'] for r in recs if r['年'] == y and r['组'] == 'PHEV'
               and r['续航'] is not None]
        if oil and rng:
            rows.append([y, round(sum(oil) / len(oil), 2),
                         round(sum(rng) / len(rng), 1), len(oil)])
    b.table(ws, 3, ['年份', '综合油耗均值(L/100km)', '纯电续航均值(km)', '样本数'], rows)

    # ---------- 24 PHEV 油耗×排量 ----------
    ws = b.sheet('PHEV油耗-排量年份', 'PHEV 油耗按排量区间年度均值')
    DISP_BINS = [('<1.0L', 0, 1000), ('1.0-1.5L', 1000, 1500), ('1.5-2.0L', 1500, 2000),
                 ('2.0-2.5L', 2000, 2500), ('≥2.5L', 2500, 1e9)]
    body = []
    for lab, lo, hi_ in DISP_BINS:
        row = [lab]
        for y in yrs:
            vals = [r['油耗'] for r in recs if r['年'] == y and r['组'] == 'PHEV'
                    and r['油耗'] is not None and r['排量'] is not None
                    and lo <= r['排量'] < hi_ and 0.1 <= r['油耗'] <= 12]
            row.append(round(sum(vals) / len(vals), 2) if vals else '')
        body.append(row)
    b.table(ws, 3, ['排量\\年份'] + [str(y) for y in yrs], body)

    # ---------- 25 技术路线矩阵 ----------
    ws = b.sheet('技术路线矩阵', 'BEV 技术路线四象限')
    ws.cell(3, 1, '阈值=BEV 续航/容量中位数（全体 BEV 有值子集）')
    bvr = [r['续航'] for r in bev if r['续航'] is not None]
    bvc = [r['容量'] for r in bev if r['容量'] is not None]
    thr_r = st.median(bvr) if bvr else 0
    thr_c = st.median(bvc) if bvc else 0
    quad = Counter()
    for r in bev:
        if r['续航'] is None or r['容量'] is None:
            continue
        quad[('低' if r['续航'] < thr_r else '高') + '续航' +
             ('低' if r['容量'] < thr_c else '高') + '容量'] += 1
    qn = sum(quad.values())
    b.table(ws, 4, ['象限', '车型数', '占比%'],
            [[k, quad.get(k, 0), round(quad.get(k, 0) / qn * 100, 1)]
             for k in ('低续航低容量', '高续航高容量', '低续航高容量', '高续航低容量') if k in quad],
            widths=[16, 10, 10])
    ws.cell(9, 1, f'阈值：续航 {thr_r:.0f} km · 容量 {thr_c:.0f} kWh')

    # ---------- 26 供应商集中度 HHI ----------
    def hhi(vals):
        c = Counter(vals)
        n = sum(c.values())
        return round(sum((v / n * 100) ** 2 for v in c.values()), 1) if n >= 5 else None
    ws = b.sheet('供应商集中度', '供应商 HHI 年度趋势')
    ws.cell(3, 1, f'电池(电芯)口径样本 {sum(1 for r in recs if r["电芯"])} · 电机口径样本 {sum(1 for r in recs if r["电机企"])}（有值子集）')
    body = []
    for y in yrs:
        hb = hhi([r['电芯'] for r in recs if r['年'] == y and r['电芯']])
        hm = hhi([r['电机企'] for r in recs if r['年'] == y and r['电机企']])
        if hb or hm:
            body.append([y, hb, hm])
    b.table(ws, 4, ['年份', '电池供应商HHI', '电机供应商HHI'], body)

    # ---------- 27/28 供应商 CR ----------
    def cr(vals, k):
        c = Counter(vals)
        n = sum(c.values())
        if n < 5:
            return None
        return round(sum(v for _, v in c.most_common(k)) / n * 100, 1)
    for name, fld in (('电池供应商CR', '电芯'), ('电机供应商CR', '电机企')):
        ws = b.sheet(name, f'{fld} 供应商集中度 CR 年度趋势')
        body = []
        for y in yrs:
            vals = [r[fld] for r in recs if r['年'] == y and r[fld]]
            if len(vals) >= 5:
                body.append([y, cr(vals, 1), cr(vals, 3), cr(vals, 5), len(vals)])
        b.table(ws, 3, ['年份', 'CR1(%)', 'CR3(%)', 'CR5(%)', '样本数'], body)

    # ---------- 29/30 供应商排名 ----------
    for name, fld, label in (('电池供应商排名', '电芯', '电芯供应商'),
                             ('电机供应商排名', '电机企', '电机生产企业')):
        ws = b.sheet(name, f'{label}排名 Top10（有值子集）')
        c = Counter(r[fld] for r in recs if r[fld])
        b.table(ws, 3, ['排名', label, '车型数'],
                [[i, k, v] for i, (k, v) in enumerate(c.most_common(10), 1)],
                widths=[6, 30, 10])

    # ---------- 31 供应商组合 ----------
    ws = b.sheet('供应商组合', '电池-电机供应商组合 Top10')
    c = Counter((r['电芯'], r['电机企']) for r in recs if r['电芯'] and r['电机企'])
    b.table(ws, 3, ['组合（电芯 × 电机企业）', '车型数'],
            [[f'{k[0]} × {k[1]}', v] for k, v in c.most_common(10)], widths=[52, 10])

    # ---------- 32/33 品牌×供应商 ----------
    def brand_sup(name, bfld, sfld, slabel, topb, tops):
        ws = b.sheet(name, f'Top{topb}品牌 × Top{tops}{slabel}（车型数）')
        brands = [k for k, _ in Counter(r[bfld] for r in recs if r[bfld]).most_common(topb)]
        sups = [k for k, _ in Counter(r[sfld] for r in recs if r[sfld]).most_common(tops)]
        body = []
        for bd in brands:
            row = [bd]
            for sp in sups:
                row.append(sum(1 for r in recs if r[bfld] == bd and r[sfld] == sp))
            body.append(row)
        b.table(ws, 3, ['品牌'] + sups, body, widths=[14] + [16] * tops)
    brand_sup('品牌-电池供应商', '商标', '电芯', '电芯供应商', 8, 5)
    brand_sup('品牌-电机供应商', '商标', '电机企', '电机供应商', 10, 5)

    # ---------- 34/35 细分市场 ----------
    for name, grp in (('细分市场-BEV', 'BEV'), ('细分市场-PHEV', 'PHEV')):
        ws = b.sheet(name, f'{grp} 细分市场汇总（车型数）')
        segs = sorted({r['细分'] for r in recs if r['细分'] and r['组'] == grp})
        body = []
        for sg in segs:
            row = [sg]
            for y in yrs:
                row.append(sum(1 for r in recs if r['组'] == grp and r['细分'] == sg and r['年'] == y))
            row.append(sum(row[1:]))
            body.append(row)
        b.table(ws, 3, ['细分市场'] + [str(y) for y in yrs] + ['合计'], body, widths=[14] + [8] * (len(yrs) + 1))

    # ---------- 36/37 细分市场对比 ----------
    def seg_cmp(name, fld, unit, label):
        ws = b.sheet(name, f'BEV 细分市场{label} 2023 vs 2025 均值对比')
        body = []
        segs = sorted({r['细分'] for r in recs if r['细分'] and r['组'] == 'BEV'})
        for sg in segs:
            v23 = [r[fld] for r in recs if r['组'] == 'BEV' and r['细分'] == sg
                   and r['年'] == 2023 and r[fld] is not None]
            v25 = [r[fld] for r in recs if r['组'] == 'BEV' and r['细分'] == sg
                   and r['年'] == 2025 and r[fld] is not None]
            if v23 and v25:
                m23, m25 = sum(v23) / len(v23), sum(v25) / len(v25)
                body.append([sg, round(m23, 1), round(m25, 1),
                             round((m25 - m23) / m23 * 100, 1)])
        b.table(ws, 3, ['细分市场', f'2023均值({unit})', f'2025均值({unit})', '变化%'], body,
                widths=[14, 14, 14, 10])
    seg_cmp('细分市场续航对比', '续航', 'km', '续航')
    seg_cmp('细分市场容量对比', '容量', 'kWh', '容量')

    # ---------- 38 质量按细分市场 ----------
    ws = b.sheet('质量按细分市场', '各细分市场整备质量均值（有值子集）')
    seg_m = {}
    for r in recs:
        if r['细分'] and r['整备'] is not None:
            seg_m.setdefault(r['细分'], []).append(r['整备'])
    rows = sorted(((k, sum(v) / len(v)) for k, v in seg_m.items() if len(v) >= 5),
                  key=lambda x: -x[1])[:12]
    b.table(ws, 3, ['细分市场', '整备均值(kg)'],
            [[k, round(v, 1)] for k, v in rows], widths=[16, 14])

    # ---------- 39/40 相关性 ----------
    CORR_B = [('续航', '续航'), ('电池容量', '容量'), ('电机功率', '功率'),
              ('整备质量', '整备'), ('能量密度', '密度'), ('电耗', '电耗')]
    CORR_P = CORR_B + [('综合油耗', '油耗')]
    for name, grp, spec in (('相关性-BEV', 'BEV', CORR_B), ('相关性-PHEV', 'PHEV', CORR_P)):
        ws = b.sheet(name, f'{grp} 相关性矩阵（Pearson，有值成对子集）')
        body = []
        for lab_a, f_a in spec:
            row = [lab_a]
            for lab_b, f_b in spec:
                xs, ys = [], []
                for r in recs:
                    if r['组'] == grp and r[f_a] is not None and r[f_b] is not None:
                        xs.append(r[f_a])
                        ys.append(r[f_b])
                p = pearson(xs, ys)
                row.append(round(p, 3) if p is not None else '')
            body.append(row)
        b.table(ws, 3, [''] + [lab for lab, _ in spec], body)

    # ---------- 41/42 同比增长 ----------
    def yoy_rows(grp, spec):
        body = []
        for lab, f in spec:
            row = [lab]
            prev = None
            for y in yrs:
                vals = [r[f] for r in recs if r['年'] == y and r['组'] == grp and r[f] is not None]
                m = sum(vals) / len(vals) if vals else None
                row.append(round((m - prev) / prev * 100, 1)
                           if (m is not None and prev) else '')
                prev = m if m is not None else prev
            body.append(row)
        return body
    WSPEC = [('平均续航(km)', '续航'), ('平均容量(kWh)', '容量'), ('平均功率(kW)', '功率'),
             ('平均整备(kg)', '整备'), ('平均能量密度(Wh/kg)', '密度'), ('平均电耗(kWh/100km)', '电耗')]
    PSPEC = WSPEC + [('平均油耗(L/100km)', '油耗')]
    for name, grp, spec in (('同比增长-BEV', 'BEV', WSPEC), ('同比增长-PHEV', 'PHEV', PSPEC)):
        ws = b.sheet(name, f'{grp} 关键指标同比增长率(%)')
        b.table(ws, 3, ['指标\\年份'] + [str(y) for y in yrs], yoy_rows(grp, spec))

    # ---------- 43/44 关键指标汇总 ----------
    for name, grp, spec in (('关键指标汇总-BEV', 'BEV', WSPEC), ('关键指标汇总-PHEV', 'PHEV', PSPEC)):
        ws = b.sheet(name, f'{grp} 关键指标（全周期描述统计）')
        body = []
        for lab, f in spec:
            vals = [r[f] for r in recs if r['组'] == grp and r[f] is not None]
            if vals:
                body.append([lab, round(sum(vals) / len(vals), 2),
                             round(st.median(vals), 1), len(vals)])
        b.table(ws, 3, ['指标', '均值', '中位数', '样本数'], body)

    # ---------- 45 年度综合指标 ----------
    ws = b.sheet('年度综合指标', '年度综合技术指标（有值子集均值）')
    hdr = ['年份', 'BEV续航', 'BEV容量', 'BEV功率', 'BEV质量', 'BEV整车密度',
           'PHEV续航', 'PHEV容量', 'PHEV功率', 'PHEV质量', 'PHEV整车密度']
    rows = []
    for y in yrs:
        row = [y]
        for grp in ('BEV', 'PHEV'):
            for f in ('续航', '容量', '功率', '整备'):
                m = agg_by_year(recs, f, grp).get(y)
                row.append(round(m[0], 2) if m else '')
            m = per_row_ratio(recs, '容量', '整备', grp).get(y)
            row.append(round(m, 1) if m else '')
        rows.append(row)
    b.table(ws, 3, hdr, rows, widths=[8] + [11] * 10)

    # ---------- 46 月销量分析（易车零售月度口径，参考） ----------
    ws = b.sheet('月销量分析', '月销量参考（易车零售 · 2026-08 · 车系级数值）')
    ws.cell(3, 1, '按通用名称聚合取最大值（同系多行共享车系级销量）；806 行有值中的 Top20')
    _s = {}
    for r in recs:
        if r['月销']:
            k = r['通用'] or r['商标'] or r['企业'][:10]
            if r['月销'] > _s.get(k, (0, ''))[0]:
                _s[k] = (r['月销'], r['商标'] or r['企业'][:8])
    top_sales = sorted(((k, v[0], v[1]) for k, v in _s.items()), key=lambda x: -x[1])[:20]
    b.table(ws, 4, ['排名', '车系（通用名称）', '月销量(辆)', '品牌'],
            [[i, _html_e(k), int(v), ent] for i, (k, v, ent) in enumerate(top_sales, 1)],
            widths=[6, 30, 12, 14])

    return b, names


def _html_e(s):
    import html as _h
    return _h.escape(str(s))


def main():
    recs = load_records()
    book, names = build(recs)
    # 覆盖检查：目录与实际 sheet 一致（45 表 + 月销量分析）
    actual = book.wb.sheetnames
    missing = [n for n in names if n not in actual]
    assert not missing, f'缺表: {missing}'
    assert '月销量分析' in actual, '缺月销量分析表'
    _bak = f"archive/行业分析底表_pre-v4923_backup_{datetime.date.today():%Y%m%d}.xlsx"
    if not os.path.exists(_bak):
        shutil.copy2(OUT, _bak)  # 09-30 教训：重复运行不得覆盖原始备份
    book.wb.save(OUT)
    # 复开验证
    chk = openpyxl.load_workbook(OUT, read_only=True)
    assert len(chk.sheetnames) == 46, len(chk.sheetnames)
    print(f'生成完成：46 表（45+月销量分析） → {OUT}（底表 {len(recs)} 行口径 {VER}）')
    print('sheets:', '、'.join(chk.sheetnames[:8]), '…')


if __name__ == '__main__':
    main()
