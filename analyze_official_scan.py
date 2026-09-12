#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""analyze_official_scan.py — 官方库扫描结果 × 底表 对账分析
产出：
  1. 型号→官方批次全映射；底表完全缺失的型号清单（按最早官方批次=批准批次归组）
  2. 各批官方乘用车数 vs 底表行数 对账表
  3. 底表行归属反核：行批次官方库不含该型号 → 潜在误植清单
用法：python analyze_official_scan.py
"""
import os, re, json, glob
from collections import defaultdict
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SCAN = r'D:\#AI\00 Project\17 TA analysis\02-数据底表\data\nev-announcements\raw\official_nev_passenger_scan.json'
scan = json.load(open(SCAN, encoding='utf-8'))

# 官方映射
model_batches = defaultdict(set)
batch_models = defaultdict(set)
for b, lst in scan.items():
    for x in lst:
        m = (x.get('clxh') or '').strip()
        if m:
            model_batches[m].add(b)
            batch_models[b].add(m)

# 底表
wb = openpyxl.load_workbook(os.path.join(HERE, 'NEV公告参数汇总表_合并版（341~410批）.xlsx'), read_only=True)
ws = wb['NEV公告参数汇总']
headers = [str(c.value).strip() if c.value is not None else '' for c in next(ws.iter_rows(min_row=1, max_row=1))]
hi = {h: i for i, h in enumerate(headers)}
wb_model_batches = defaultdict(set)
wb_batch_rows = defaultdict(int)
wb_rows = []
for row in ws.iter_rows(min_row=2, values_only=True):
    if not any(v not in (None, '') for v in row):
        continue
    b = str(row[hi['批次']]).strip(); m = str(row[hi['产品型号']] or '').strip()
    wb_batch_rows[b] += 1
    wb_rows.append((b, m))
    if m:
        wb_model_batches[m].add(b)

# 1) 底表完全缺失的官方型号
missing = {}
for m, bs in model_batches.items():
    if m not in wb_model_batches:
        missing[m] = sorted(bs, key=int)
print(f'官方乘用车型号总数: {len(model_batches)} | 底表缺失: {len(missing)}')
by_first = defaultdict(list)
for m, bs in missing.items():
    by_first[bs[0]].append(m)
print('\n按最早官方批次（=建议补录批次）:')
for b in sorted(by_first, key=int):
    print(f'  {b}批: {len(by_first[b])} 款')

# 2) 各批对账
print('\n批次对账（官方乘用车款数 vs 底表行数）:')
print(f"{'批':4s} {'官方':>4s} {'底表':>4s} {'差':>4s}")
gaps = {}
for b in sorted(set(list(batch_models) + list(wb_batch_rows)), key=int):
    o = len(batch_models.get(b, set())); w = wb_batch_rows.get(b, 0)
    gaps[b] = o - w
    print(f'{b:4s} {o:4d} {w:4d} {o-w:4d}')

# 3) 底表行归属反核（行批次官方库不含该型号）
mismatch = []
for b, m in wb_rows:
    if m and m in model_batches and b not in model_batches[m]:
        mismatch.append({'批次': b, '型号': m, '官方批次': sorted(model_batches[m], key=int)})
print(f'\n底表行归属反核（行批次官方库不含该型号）: {len(mismatch)} 行')
for x in mismatch[:20]:
    print(f"  {x['批次']}/{x['型号']:24s} 官方批次: {x['官方批次']}")

json.dump({'missing_by_batch': {k: v for k, v in by_first.items()},
           'mismatch': mismatch}, open(os.path.join(HERE, 'audit-output', 'official_scan_analysis.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print('\n-> audit-output/official_scan_analysis.json')
