# -*- coding: utf-8 -*-
"""v4.4.0 电机列对账：强制 总=前+后（用户口径），不变量违例以台账修正。
规则：前 存在 → 期望总 = 前 + (后 or 0)；若 当前总 ≠ 期望总 → 重写总列（浅红+台账）。
      前 为空 → 不动（仅总未拆行维持数值总）。
"""
import sys

import openpyxl
from openpyxl.styles import PatternFill

sys.stdout.reconfigure(encoding="utf-8")
SRC = "NEV公告参数汇总表_合并版（341~410批）.xlsx"
RED = PatternFill("solid", fgColor="FFC7CE")


def parse_pt(v):
    if v is None:
        return None, None, ""
    s = str(v).strip()
    if not s:
        return None, None, ""
    parts = s.split("/")
    try:
        p = float(parts[0]) if parts[0] else None
    except ValueError:
        p = None
    try:
        t = float(parts[1]) if len(parts) > 1 and parts[1] else None
    except ValueError:
        t = None
    return p, t, s


def pt_cell(p, t):
    if p is None:
        return ""
    a = str(int(p)) if float(p) == int(p) else str(p)
    if t is None:
        return a
    b = str(int(t)) if float(t) == int(t) else str(t)
    return f"{a}/{b}"


def main():
    wb = openpyxl.load_workbook(SRC)
    ws = wb["NEV公告参数汇总"]
    hdr = [str(c.value).strip() if c.value else "" for c in ws[1]]
    fi = {h: i for i, h in enumerate(hdr)}
    iF, iT, iR = fi["前电机功率/扭矩"], fi["电机总功率/扭矩"], fi["后电机功率/扭矩"]
    cr = wb["变更记录"]
    seq = 0
    for rr in range(cr.max_row, 1, -1):
        v = cr.cell(rr, 1).value
        if isinstance(v, int):
            seq = v
            break
    fixed = kept = 0
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value is None:
            continue
        fp, ft, _ = parse_pt(ws.cell(r, iF + 1).value)
        rp, rt, _ = parse_pt(ws.cell(r, iR + 1).value)
        tp, tt, ts = parse_pt(ws.cell(r, iT + 1).value)
        if fp is None:
            continue
        exp_p = fp + (rp or 0)
        exp_t = None
        if ft is not None and (rt is not None or rp is None):
            exp_t = ft + (rt or 0)
        if tp is not None and abs(tp - exp_p) <= 2 and (
                (exp_t is None) or (tt is not None and abs(tt - exp_t) <= 2)):
            kept += 1
            continue
        val = pt_cell(exp_p, exp_t)
        cell = ws.cell(r, iT + 1)
        old = cell.value
        cell.value = val
        cell.fill = RED
        fixed += 1
        seq += 1
        cr.append([seq, ws.cell(r, 1).value, ws.cell(r, fi["产品型号"] + 1).value,
                   ws.cell(r, fi["动力类型"] + 1).value, "电机总功率/扭矩", "数据清洗更正",
                   "" if old is None else str(old)[:40], val, "对账修正：总=前+后（共识独立补全致不一致）"])
    wb.save(SRC)
    print(f"对账完成：一致 {kept} 行，重写总列 {fixed} 行（台账末序号 {seq}）")


if __name__ == "__main__":
    main()
