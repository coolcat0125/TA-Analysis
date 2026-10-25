# -*- coding: utf-8 -*-
"""Normalize batch 411/412 scan data and compute delivery metrics."""

from __future__ import annotations

import json
import re
import statistics
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
REPO = ROOT.parent / "01-主仓库" / "TA-Analysis"
MASTER = REPO / "NEV公告参数汇总表_合并版（341~410批）.xlsx"
B411_ARCHIVE = ROOT.parent / "02-数据底表" / "data" / "nev-announcements" / "raw" / "b411"


def clean(value: Any) -> str | None:
    if value is None:
        return None
    text = " ".join(str(value).replace("\xa0", " ").split()).strip()
    return text or None


def first_number(value: Any) -> float | None:
    match = re.search(r"-?\d+(?:\.\d+)?", str(value or ""))
    return float(match.group()) if match else None


def numbers(value: Any) -> list[float]:
    return [float(item) for item in re.findall(r"-?\d+(?:\.\d+)?", str(value or ""))]


def mean_or_none(values: list[float]) -> float | None:
    return round(statistics.mean(values), 1) if values else None


def classify_route(record: dict[str, Any]) -> str:
    text = "".join(
        str(record.get(key) or "")
        for key in ("产品名称", "产品型号", "公示详情原始其它")
    )
    if "增程" in text or "REEV" in text or "DREEV" in text:
        return "增程(EREV)"
    if "插电" in text or "PHEV" in text or "汽油/电" in text:
        return "插电混动(PHEV)"
    return "纯电(BEV)"


def is_erev(record: dict[str, Any]) -> bool:
    return classify_route(record) == "增程(EREV)"


def component_engine_power(record: dict[str, Any]) -> float | None:
    values: list[float] = []
    for row in record.get("发动机组件_公示详情") or []:
        value = first_number(row.get("功率"))
        if value is not None:
            values.append(value)
    if values:
        return max(values)
    text = str(record.get("公示详情原始其它") or "")
    matches = re.findall(r"发动机(?:最大净功率|净功率值|净功率)[:：为]?\s*([\d.]+)\s*kW", text)
    return first_number(matches[0]) if matches else None


def component_engine_displacement(record: dict[str, Any]) -> float | None:
    values: list[float] = []
    for row in record.get("发动机组件_公示详情") or []:
        value = first_number(row.get("排量"))
        if value is not None and value > 0:
            values.append(value)
    if values:
        return max(values)
    return first_number(record.get("发动机排量(mL)"))


def merge_archive(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    extracted = json.loads((B411_ARCHIVE / "details_extracted.json").read_text(encoding="utf-8"))
    cell_fill = json.loads(
        (B411_ARCHIVE / "cell_fill_411_20260918.json").read_text(encoding="utf-8")
    )
    for record in records:
        model = str(record.get("产品型号") or "").strip().upper()
        archive = extracted.get(model, {})
        mappings = {
            "车长(mm)": archive.get("车长(mm)"),
            "整备质量(kg)": archive.get("整备质量(kg)"),
            "轴距(mm)": archive.get("轴距(mm)"),
            "电池类型_正式": archive.get("电池类型"),
            "电芯供应商_正式": archive.get("电芯供应商"),
            "电池包供应商_正式": archive.get("电池包供应商"),
            "发动机型号": archive.get("发动机型号"),
            "发动机生产企业": archive.get("发动机生产企业"),
            "发动机排量(mL)": archive.get("发动机排量(mL)"),
            "发动机功率(kW)": archive.get("发动机功率(kW)"),
        }
        for key, value in mappings.items():
            if value not in (None, "") and not record.get(key):
                record[key] = value
        fill = cell_fill.get(model, {})
        if fill.get("cell"):
            record["电芯供应商_正式"] = fill["cell"]
        if fill.get("pack"):
            record["电池包供应商_正式"] = fill["pack"]
    return records


def load_history() -> tuple[set[str], dict[str, list[float]], dict[str, float | int | None]]:
    workbook = load_workbook(MASTER, read_only=True, data_only=True)
    sheet = workbook["NEV公告参数汇总"]
    header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    index = {name: pos for pos, name in enumerate(header)}
    models: set[str] = set()
    baseline_groups: dict[str, list[float]] = defaultdict(list)
    baseline_rows: list[Any] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        model = clean(row[index["产品型号"]])
        if model:
            models.add(model.upper())
        if str(row[index["批次"]]) != "410":
            continue
        baseline_rows.append(row)
        product_name = str(row[index["产品名称"]] or "")
        model_name = str(row[index["产品型号"]] or "")
        power_type = str(row[index["动力类型"]] or "")
        text = product_name + model_name + power_type
        if "增程" in text or "REEV" in text:
            route = "增程(EREV)"
        elif "插电" in text or "PHEV" in text:
            route = "插电混动(PHEV)"
        else:
            route = "纯电(BEV)"
        baseline_groups[route].append(row)
    workbook.close()
    def baseline_metric(field: str) -> float | None:
        position = index[field]
        values = []
        for row in baseline_rows:
            value = first_number(row[position]) if position < len(row) else None
            if value is not None:
                values.append(value)
        return mean_or_none(values)

    baseline_overall = {
        "数量": len(baseline_rows),
        "平均车长(mm)": baseline_metric("车长(mm)"),
        "平均电机功率(kW)": baseline_metric("电机总功率/扭矩"),
        "平均续航(km)": baseline_metric("纯电续航里程(km)"),
        "平均电池容量(kWh)": baseline_metric("电池容量(kWh)"),
    }
    return models, baseline_groups, baseline_overall


def normalize_record(record: dict[str, Any], history_models: set[str]) -> dict[str, Any]:
    route = classify_route(record)
    model = clean(record.get("产品型号")) or ""
    battery_type = (
        clean(record.get("电池类型_公示其它"))
        or clean(record.get("电池类型_正式"))
        or clean(record.get("电池类型"))
    )
    cell_supplier = (
        clean(record.get("电芯供应商_公示其它"))
        or clean(record.get("电芯供应商_正式"))
        or clean(record.get("电芯供应商"))
    )
    pack_supplier = (
        clean(record.get("电池包供应商_公示其它"))
        or clean(record.get("电池包供应商_正式"))
        or clean(record.get("电池包供应商"))
    )
    engine_power = component_engine_power(record)
    engine_displacement = component_engine_displacement(record)
    missing = [
        label
        for label, value in (
            ("车长", record.get("车长(mm)")),
            ("续航", record.get("纯电续航里程(km)")),
            ("电池容量", record.get("电池容量(kWh)")),
            ("电池类型", battery_type),
            ("电机功率", record.get("电机总功率_kW")),
            ("发动机功率", engine_power),
        )
        if value in (None, "", [])
    ]
    return {
        "批次": int(record.get("批次") or 0),
        "状态": clean(record.get("批次状态")),
        "技术路线": route,
        "是否增程": "是" if is_erev(record) else "否",
        "企业名称": clean(record.get("企业名称")),
        "产品商标": clean(record.get("产品商标")),
        "通用名称": clean(record.get("通用名称")),
        "产品型号": model,
        "产品名称": clean(record.get("产品名称")),
        "车长(mm)": first_number(record.get("车长(mm)")),
        "车宽(mm)": first_number(record.get("车宽(mm)")),
        "车高(mm)": first_number(record.get("车高(mm)")),
        "轴距(mm)": first_number(record.get("轴距(mm)")),
        "整备质量(kg)": first_number(record.get("整备质量(kg)")),
        "总质量(kg)": first_number(record.get("总质量(kg)")),
        "最高车速(km/h)": first_number(record.get("最高车速(km/h)")),
        "额定载客(人)": clean(record.get("额定载客(人)")),
        "纯电续航(km)": first_number(record.get("纯电续航里程(km)")),
        "电池容量(kWh)": first_number(record.get("电池容量(kWh)")),
        "电池类型": battery_type,
        "电芯供应商": cell_supplier,
        "电池包供应商": pack_supplier,
        "电机峰值功率数列(kW)": clean(record.get("电机峰值功率数列_kW")),
        "电机总功率(kW)": first_number(record.get("电机总功率_kW")),
        "驱动形式": clean(record.get("驱动形式_电机数量推断")),
        "发动机排量(mL)": engine_displacement,
        "发动机功率(kW)": engine_power,
        "综合油耗(L/100km)": first_number(record.get("综合油耗(L/100km)")),
        "型号级新增": "历史批次已有同型号" if model.upper() in history_models else "型号级新增",
        "来源URL": clean(record.get("来源URL") or record.get("详情页URL")),
        "缺失关键字段": "、".join(missing) if missing else "",
    }


def group_stats(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    output = []
    for route in ("纯电(BEV)", "插电混动(PHEV)", "增程(EREV)"):
        group = [record for record in records if record["技术路线"] == route]
        if not group:
            continue
        power_known = [record for record in group if record["电机总功率(kW)"] is not None]
        drive_known = [record for record in group if record["驱动形式"]]
        output.append(
            {
                "技术路线": route,
                "数量": len(group),
                "平均续航(km)": mean_or_none(
                    [float(record["纯电续航(km)"]) for record in group if record["纯电续航(km)"] is not None]
                ),
                "平均电池容量(kWh)": mean_or_none(
                    [
                        float(record["电池容量(kWh)"])
                        for record in group
                        if record["电池容量(kWh)"] is not None
                    ]
                ),
                "平均电机功率(kW)": mean_or_none(
                    [float(record["电机总功率(kW)"]) for record in power_known]
                ),
                "平均车长(mm)": mean_or_none(
                    [float(record["车长(mm)"]) for record in group if record["车长(mm)"] is not None]
                ),
                "平均整备质量(kg)": mean_or_none(
                    [
                        float(record["整备质量(kg)"])
                        for record in group
                        if record["整备质量(kg)"] is not None
                    ]
                ),
                "功率有效样本": len(power_known),
                "200kW及以上占比": (
                    round(
                        sum(1 for record in power_known if float(record["电机总功率(kW)"]) >= 200)
                        / len(power_known),
                        3,
                    )
                    if power_known
                    else None
                ),
                "多电机占比": (
                    round(
                        sum(1 for record in drive_known if record["驱动形式"] in ("双电机", "三电机", "四电机"))
                        / len(drive_known),
                        3,
                    )
                    if drive_known
                    else None
                ),
                "多电机有效样本": len(drive_known),
            }
        )
    return output


def baseline_stats(groups: dict[str, list[Any]]) -> list[dict[str, Any]]:
    header_index = {
        "纯电续航里程(km)": 13,
        "电池容量(kWh)": 14,
        "电机总功率/扭矩": 19,
        "车长(mm)": 30,
        "整备质量(kg)": 11,
    }
    output = []
    for route, rows in groups.items():
        def values(field: str) -> list[float]:
            position = header_index[field]
            out = []
            for row in rows:
                if position < len(row):
                    value = first_number(row[position])
                    if value is not None:
                        out.append(value)
            return out

        output.append(
            {
                "技术路线": route,
                "数量": len(rows),
                "平均续航(km)": mean_or_none(values("纯电续航里程(km)")),
                "平均电池容量(kWh)": mean_or_none(values("电池容量(kWh)")),
                "平均电机功率(kW)": mean_or_none(values("电机总功率/扭矩")),
                "平均车长(mm)": mean_or_none(values("车长(mm)")),
                "平均整备质量(kg)": mean_or_none(values("整备质量(kg)")),
            }
        )
    return output


def coverage(records: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    fields = (
        "车长(mm)",
        "轴距(mm)",
        "整备质量(kg)",
        "纯电续航(km)",
        "电池容量(kWh)",
        "电池类型",
        "电芯供应商",
        "电池包供应商",
        "电机总功率(kW)",
        "驱动形式",
        "综合油耗(L/100km)",
    )
    return {
        field: {
            "covered": sum(1 for record in records if record.get(field) not in (None, "", [])),
            "total": len(records),
        }
        for field in fields
    }


def main() -> None:
    history_models, baseline_groups, baseline_overall = load_history()
    records_411 = [
        normalize_record(record, history_models)
        for record in merge_archive(json.loads((DATA / "details_411.json").read_text(encoding="utf-8")))
    ]
    records_412 = [
        normalize_record(record, history_models)
        for record in json.loads((DATA / "details_412.json").read_text(encoding="utf-8"))
    ]

    payload = {
        "as_of": "2026-10-17",
        "records": {"411": records_411, "412": records_412},
        "batch_summary": {
            "overall_410": baseline_overall,
            "410": baseline_stats(baseline_groups),
            "411": group_stats(records_411),
            "412": group_stats(records_412),
        },
        "coverage": {"411": coverage(records_411), "412": coverage(records_412)},
        "new_model_counts": {
            "411": Counter(record["型号级新增"] for record in records_411),
            "412": Counter(record["型号级新增"] for record in records_412),
        },
        "sources": [
            {
                "name": "第411批正式公告",
                "url": "https://www.miit.gov.cn/jgsj/zbys/wjfb/art/2026/art_ef051926ce924a4da93e1677c7e18b34.html",
            },
            {
                "name": "第412批拟发布内容公示",
                "url": "https://www.miit.gov.cn/jgsj/zbys/qcgy/art/2026/art_ec785d818a604ccd9618b73d4f13f062.html",
            },
            {
                "name": "第412批车辆新产品公示清单",
                "url": "https://www.miit.gov.cn/datainfo/cpgg/art/2026/art_43a4b4be0de2483d85f28e9423201494.html",
            },
        ],
    }
    target = DATA / "delivery_payload.json"
    target.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8",
    )
    print(f"Wrote {target}")
    print(json.dumps(payload["batch_summary"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
