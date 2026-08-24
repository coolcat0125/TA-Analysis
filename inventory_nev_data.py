#!/usr/bin/env python3
"""Inventory the TA-Analysis source workbook for auditable enrichment.

Copyright © 2026 David YE
SPDX-License-Identifier: MIT
"""
from __future__ import annotations
import csv, datetime as dt, glob, hashlib, json, re
from pathlib import Path
from openpyxl import load_workbook

COPYRIGHT = "Copyright © 2026 David YE"

def missing(v): return v is None or str(v).strip() in {"", "/", "-", "--", "N/A"}
def num(v):
    if missing(v): return None
    m = re.search(r"-?\d+(?:\.\d+)?", str(v).replace(",", ""))
    return float(m.group()) if m else None
def sha256(path):
    h=hashlib.sha256()
    with open(path,"rb") as f:
        for chunk in iter(lambda:f.read(1024*1024),b""): h.update(chunk)
    return h.hexdigest()
def find(headers,*terms):
    return next((i for i,h in enumerate(headers) if any(t in h for t in terms)),None)

files=sorted(glob.glob("NEV公告参数汇总表_合并版*.xlsx"))
if not files: raise SystemExit("source workbook not found")
source=Path(files[-1]); wb=load_workbook(source,read_only=True,data_only=True)
ws=wb["NEV公告参数汇总"] if "NEV公告参数汇总" in wb.sheetnames else wb[wb.sheetnames[0]]
it=ws.iter_rows(values_only=True); headers=[str(x).strip() if x is not None else "" for x in next(it)]
rows=[list(r) for r in it if any(not missing(x) for x in r)]
stats=[]
for i,h in enumerate(headers):
    filled=sum(1 for r in rows if i<len(r) and not missing(r[i])); stats.append({"index":i+1,"header":h,"filled":filled,"missing":len(rows)-filled,"coverage":round(filled/len(rows),6)})
ix={"batch":find(headers,"批次"),"model":find(headers,"产品型号","车型型号"),"torque":find(headers,"电机最大扭矩","电机总扭矩","电机扭矩"),"capacity":find(headers,"电池容量","电池组总能量"),"range":find(headers,"纯电续航"),"consumption":find(headers,"百公里电耗"),"energy":find(headers,"动力蓄电池组总能量","电池组总能量"),"mass":find(headers,"动力蓄电池组总质量","电池组总质量"),"density":find(headers,"电池能量密度"),"supplier":find(headers,"电池生产企业","电池供应商"),"chemistry":find(headers,"电池类型","电池材料"),"voltage":find(headers,"电压平台","系统电压"),"length":find(headers,"电池包长","电池长度"),"width":find(headers,"电池包宽","电池宽度"),"height":find(headers,"电池包高","电池高度")}
def get(r,k): return r[ix[k]] if ix.get(k) is not None and ix[k]<len(r) else None
infer_consumption=sum(1 for r in rows if missing(get(r,"consumption")) and num(get(r,"capacity")) and num(get(r,"range")))
infer_density=sum(1 for r in rows if missing(get(r,"density")) and num(get(r,"energy")) and num(get(r,"mass")))
torque_missing=[]; packs80=[]
for n,r in enumerate(rows,start=2):
    model=str(get(r,"model") or "").strip(); batch=str(get(r,"batch") or "").strip(); cap=num(get(r,"capacity"))
    if ix["torque"] is not None and missing(get(r,"torque")): torque_missing.append({"excel_row":n,"batch":batch,"product_model":model})
    if cap is not None and 75<=cap<=85:
        packs80.append({"excel_row":n,"batch":batch,"product_model":model,"capacity_kwh":cap,"supplier":get(r,"supplier"),"chemistry":get(r,"chemistry"),"voltage":get(r,"voltage"),"pack_length":get(r,"length"),"pack_width":get(r,"width"),"pack_height":get(r,"height"),"pack_mass":get(r,"mass"),"system_density":get(r,"density")})
out=Path("audit-output"); out.mkdir(exist_ok=True)
report={"schema_version":"1.0.0","copyright":COPYRIGHT,"generated_at_utc":dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat(),"source":{"file":source.name,"sha256":sha256(source),"sheet":ws.title},"records":len(rows),"headers":headers,"column_statistics":stats,"resolved_columns":ix,"inferable":{"consumption":infer_consumption,"energy_density":infer_density},"motor_torque_missing_count":len(torque_missing),"pack_80kwh_count":len(packs80),"constraints":["pack dimensions require direct L×W×H evidence","800V is not treated as exact voltage","derived pack mass only uses same-row energy/system density and must be marked 非称重实测","source values are never overwritten"]}
(out/"data_inventory.json").write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
for name,data,fields in [("motor_torque_missing.csv",torque_missing,["excel_row","batch","product_model"]),("battery_80kwh_priority.csv",packs80,["excel_row","batch","product_model","capacity_kwh","supplier","chemistry","voltage","pack_length","pack_width","pack_height","pack_mass","system_density"])]:
    with open(out/name,"w",encoding="utf-8-sig",newline="") as f:
        f.write(f"# {COPYRIGHT}\n")
        w=csv.DictWriter(f,fieldnames=fields); w.writeheader(); w.writerows(data)
with open(out/"column_coverage.csv","w",encoding="utf-8-sig",newline="") as f:
    f.write(f"# {COPYRIGHT}\n")
    w=csv.DictWriter(f,fieldnames=["index","header","filled","missing","coverage"]); w.writeheader(); w.writerows(stats)
print(json.dumps({"records":len(rows),"torque_missing":len(torque_missing),"infer_consumption":infer_consumption,"infer_density":infer_density,"packs80":len(packs80)},ensure_ascii=False))
