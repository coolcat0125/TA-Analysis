# -*- coding: utf-8 -*-
"""
v4.9.35 · 续航工况四列尾插落库（CR-17-031）
依据：Muse《TA公告数据问题清单与执行建议_20261025》§1/§3 结论，映射到本地权威底表。
铁律：只尾插 4 列新注记，不改动任何既有单元格（cols 1-41 零差异）；
     证据分级明标（逐项核实/同族同值/佐证级/法规默认/证书值/未决/存疑/不可用）；
     插混 ≥350 批批量 WLTC 默认提案（Muse B2）未经用户批准，不执行。
列设计（本地化命名，语义=Muse DICT 工况四列）：
  col42 测试循环   CLTC / WLTC / NEDC / 未区分
  col43 CLTC续航(km)  池值的 CLTC 镜像（或双值并列行/对照值）
  col44 WLTC续航(km)  池值的 WLTC 镜像（或双值并列行/对照值）
  col45 续航工况证据  分级前缀 + 依据（Muse 章节引用）
用法：python cycle_annotation_v4935.py --dry-run | --apply
"""
import sys, os, json, shutil, datetime
from copy import copy
import openpyxl
from openpyxl.styles import PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
XL = os.path.join(HERE, 'NEV公告参数汇总表_合并版（341~410批）.xlsx')
SHEET = 'NEV公告参数汇总'
BACKUP = os.path.join(HERE, 'archive', 'v4.9.34_pre-cycle-annotation_backup_20261026.xlsx')
FILL = PatternFill('solid', start_color='FFD9F2F2', end_color='FFD9F2F2')  # 浅青底纹
NEW_HEADERS = ['测试循环', 'CLTC续航(km)', 'WLTC续航(km)', '续航工况证据']
SRC = 'Muse工况核实2026-10-25'

EV_LAW = ('法规默认·CLTC：GB/T 18386.1-2021（2021-03-09发布/2021-10-01实施），'
          '350批起纯电乘用车按CLTC申报；非逐项核实，个别变更/扩展行可能沿用旧测试值（Muse §1.3）')
EV_TRANS = ('过渡期·未区分：341~349批 CLTC/NEDC 并存（同批 ModelY=NEDC、元PLUS=CLTC 实证），'
            '不批量默认，待逐项核实（Muse §1.3）')

def _b(r):
    try: return int(str(r[0]).strip())
    except Exception: return None

def _code(r): return str(r[1] or '').strip()

# ---------- 逐项/同族/佐证级 WLTC（WLTC续航=池值；个别双值行列另给 CLTC 对照） ----------
WLTC_DUAL = {  # code -> (batch, CLTC对照值, ev)
    'FZ6530REEVWD19C': (403, 500, '逐项核实·WLTC：零跑D19 双值并列 CLTC 500 / WLTC 370（Muse §1.2）'),
    'FZ6530REEVWD19F': (407, 400, '逐项核实·WLTC：零跑D19 双值并列 CLTC 400 / WLTC 285（Muse §1.2）'),
}
WLTC_POOL = {  # code -> (batch, ev)  WLTC续航=池值，CLTC列留空
    'CC6511DQ22CPHEV': (411, '逐项核实·WLTC：魏牌蓝山 220/270（Muse §1.2）'),
    'CC6531DQ22APHEV': (411, '逐项核实·WLTC：魏牌蓝山 220/270（Muse §1.2）'),
    'CC6511DQ22APHEV': (411, '逐项核实·WLTC：魏牌蓝山 220/270（Muse §1.2）'),
    'CC6531DQ22CPHEV': (411, '逐项核实·WLTC：魏牌蓝山 220/270（Muse §1.2）'),
    'CC2040BV23BPHEV': (398, '同族同值·WLTC：坦克700 190km 值族（Muse 逐项核实190；本地398/411两行同值同配）'),
    'CC2040BV23CPHEV': (411, '同族同值·WLTC：坦克700 190km 值族（Muse 逐项核实190；本地398/411两行同值同配）'),
    'SC6530AAAB6HEV': (411, '逐项核实·WLTC：阿维塔T09 300/320/330 为工信部WLTC申报值（媒体逐字证据）；v1.3误标CLTC已v1.4纠正，CLTC列留空（Muse §1.6）'),
    'SC6530AAAA6HEV': (411, '逐项核实·WLTC：阿维塔T09 300/320/330 为工信部WLTC申报值（媒体逐字证据）；v1.3误标CLTC已v1.4纠正，CLTC列留空（Muse §1.6）'),
    'MR6521DCHEV116': (411, '逐项核实·WLTC：银河M8 两驱231/四驱251 为WLTC申报值（媒体逐字证据）；v1.3误标CLTC已v1.4纠正（Muse §1.6）'),
    'MR6521DCHEV117': (411, '逐项核实·WLTC：银河M8 两驱231/四驱251 为WLTC申报值（媒体逐字证据）；v1.3误标CLTC已v1.4纠正（Muse §1.6）'),
    'MR6482DCHEV03': (372, '逐项核实·WLTC：领克08 185（Muse §1.2）'),
    'MR6482DCHEV02': (372, '佐证级·WLTC：领克08 205（佐证级证据；Muse §1.2）'),
    'MR6481DCHEV100': (384, '佐证级·WLTC：领克08 205（佐证级证据；Muse §1.2）'),
    'SQR6480CHEVT28': (380, '同族同值·WLTC：风云T9 100/160 值族（Muse 逐项核实100/160；402批32.659kWh版另注未决）'),
    'SQR6481CHEVT28': (383, '同族同值·WLTC：风云T9 100/160 值族（Muse 逐项核实100/160；402批32.659kWh版另注未决）'),
    'SQR6480CHEVT285': (398, '同族同值·WLTC：风云T9 100/160 值族（Muse 逐项核实100/160；402批32.659kWh版另注未决）'),
    'SQR6480CHEVT281': (385, '同族同值·WLTC：风云T9 100/160 值族（Muse 逐项核实100/160；402批32.659kWh版另注未决）'),
    'LXA6511SHEVX1': (362, '同族同值·WLTC：理想L8 168km（Muse 逐项核实7行；本地同族同值共8行，176km行另注未决）'),
    'LXA6511SHEVX4': (367, '同族同值·WLTC：理想L8 168km（Muse 逐项核实7行；本地同族同值共8行，176km行另注未决）'),
    'LXA6511SHEVX3': (368, '同族同值·WLTC：理想L8 168km（Muse 逐项核实7行；本地同族同值共8行，176km行另注未决）'),
    'LXA6511SHEVX5': (372, '同族同值·WLTC：理想L8 168km（Muse 逐项核实7行；本地同族同值共8行，176km行另注未决）'),
    'LXA6511SHEVX01': (373, '同族同值·WLTC：理想L8 168km（Muse 逐项核实7行；本地同族同值共8行，176km行另注未决）'),
    'LXA6511SHEVX03': (373, '同族同值·WLTC：理想L8 168km（Muse 逐项核实7行；本地同族同值共8行，176km行另注未决）'),
    'LXA6511SHEVX04': (373, '同族同值·WLTC：理想L8 168km（Muse 逐项核实7行；本地同族同值共8行，176km行另注未决）'),
    'LXA6511SHEVX05': (373, '同族同值·WLTC：理想L8 168km（Muse 逐项核实7行；本地同族同值共8行，176km行另注未决）'),
    'CC6530DC21BPHEV': (404, '佐证级·WLTC：V9X远山 363（本次未独立复核公告原文，佐证级；Muse §1.2）'),
}
EV_BYD_WLTC = ('同族同值·WLTC：18.316kWh 比亚迪DM-i WLTC值族 秦101/宋85/汉102（Muse 逐项核实该值族；'
               '本行同电池同值；NEDC值族120/110/100/121 见345批行）')
WLTC_VALUE_SET = {101, 85, 102}

# ---------- NEDC（无 NEDC续航 列设计，值列留空，循环=NEDC） ----------
NEDC_MY_CODES = {'TSL6480BEVBA3', 'TSL6480BEVBA4', 'TSL6480BEVBA5'}
EV_NEDC_MY = '逐项核实·NEDC：346批 ModelY 594/640/566（过渡期批次 NEDC 实证；Muse §1.2/§1.4）'
EV_NEDC_BYD = ('同族同值·NEDC：345批 18.316kWh DM-i NEDC值族 秦120/宋100·110/汉121（Muse 逐项核实；'
               '宋Pro 110/宋MAX 105 为同族同值扩展）')

# ---------- CLTC 逐项核实（CLTC续航=池值） ----------
CLTC_POOL = {  # code -> (batch, ev)
    'LXA6503BEVW05': (411, '逐项核实·CLTC：411批 理想i6 780/770（Muse §1.2）'),
    'LXA6503BEVW04': (411, '逐项核实·CLTC：411批 理想i6 720/710（Muse §1.2）'),
    'LXA6503BEVW06': (411, '逐项核实·CLTC：411批 理想i6 720/710（Muse §1.2）'),
    'BMW7000BABEV': (411, '逐项核实·CLTC：411批 802km（Muse 记「理想i3(802)」，本地该值行为新世代BMW i3长轴距版，按值锚定，命名差异留痕；Muse §1.2）'),
    'SVW646193AEV': (411, '逐项核实·CLTC：411批 ID.ERA 5X 415/520（Muse §1.2）'),
    'SVW646194AEV': (411, '逐项核实·CLTC：411批 ID.ERA 5X 415/520（Muse §1.2）'),
    'BYD7000CFBEV3': (411, '逐项核实·CLTC：411批 海鸥 320（Muse §1.2）'),
    'TSL6480BEVBA7': (391, '逐项核实·CLTC：ModelY 719（Muse v1.4 新增12行之一）'),
    'TSL6480BEVBA6': (409, '逐项核实·CLTC：ModelY 659（Muse v1.4 新增12行之一；本地该值存为字符串，原样镜像）'),
    'BYD7003BEVA': (346, '逐项核实·CLTC：346批 元PLUS 430/510（过渡期批次；同批 ModelY=NEDC、元PLUS=CLTC，工况不按批次切换的实证；Muse §1.3/§1.4）'),
    'BYD7003BEVA1': (346, '逐项核实·CLTC：346批 元PLUS 430/510（过渡期批次；同批 ModelY=NEDC、元PLUS=CLTC，工况不按批次切换的实证；Muse §1.3/§1.4）'),
}

# ---------- M7 证书值家族（373~389 批 SKE6500SSHEV*，规则：容量×驱动功率分型） ----------
EV_M7_BASE = ('证书值·未区分：问界M7 池值175/162/170/156 既非CLTC（后驱240/四驱210）亦不能一概标WLTC'
              '（后驱WLTC 200；四驱WLTC 162与175两源冲突不裁决）（Muse §1.5，B类确认前现行处理）')

# ---------- P1 存疑（原值不动，只加注） ----------
P1_ROWS = {  # (batch, code) -> ev
    (411, 'HQ6472DCHEV111'): '存疑·P1：星舰7 19.09kWh 配 230km，约2倍于同电池家族最高115km，疑沿用未改；原值不动，待新申报/官网配置表（Muse P1）',
    (411, 'HQ6472DCHEV110'): '存疑·P1关联：星舰7 同批兄弟行 215km（Muse池该行电池缺失，本地容量40.06）；续航值待核，原值不动（Muse P1）',
    (411, 'MR7151DCHEV131'): '存疑·P1：银河A7 18.4kWh 配 202km，官方同电池版 CLTC 120~135km、同电池家族续航85~120km；原值不动（Muse P1）',
    (402, 'HQ7151DCHEV112'): '存疑·P1：星耀7 仅402批行 18.4kWh 配 165km（400/401批28.3kWh行同续航正常，兄弟行110/111/113/118/119改18.4kWh后续航正常，反衬165km疑沿用未改）；原值不动（Muse P1）',
    (404, 'HFC7150DREEV11'): '存疑·P1：江淮行 28.4kWh/315km 车型身份公开渠道无法确认；与SC7150AAAZ6HEV同批同电池同油耗4.97/B4.79 疑似复制/错位，两行并标；原值不动（Muse P1）',
    (404, 'SC7150AAAZ6HEV'): '存疑·P1：与HFC7150DREEV11同批同电池同油耗，对方行疑复制/错位，两行并标；本行原值不动（Muse P1）',
}

# ---------- P3 未决（不作归属） ----------
P3_ROWS = {  # (batch, code) -> ev
    (352, 'BYD7150YT6HEV10'): '未决·P3：汉双电机版(345/666kW) 101km 数值未决（单电机102km已证实WLTC；本行101值族归属待核）（Muse P3）',
    (357, 'BYD7150YT6HEV7'): '未决·P3：汉双电机版(345/666kW) 101km 数值未决（单电机102km已证实WLTC；本行101值族归属待核）（Muse P3）',
    (397, 'BYD7150YT6HEV14'): "未决·P3：汉双电机版 110.0km 数值未决（本地值为字符串'110.0'原样保留）（Muse P3）",
    (398, 'SKE6510SSHEVA2S'): '未决·P3：398批 问界M7新款行 242/165/175/251 工况与数值待核（Muse P3）',
    (398, 'SKE6510SSHEVA4S'): '未决·P3：398批 问界M7新款行 242/165/175/251 工况与数值待核（Muse P3）',
    (398, 'SKE6510SSHEVA5S'): '未决·P3：398批 问界M7新款行 242/165/175/251 工况与数值待核（Muse P3）',
    (398, 'SKE6510SSHEVAS'): '未决·P3：398批 问界M7新款行 242/165/175/251 工况与数值待核（Muse P3）',
    (398, 'SKE6510SSHEVR'): '未决·P3：398批 问界M7新款行 242/165/175/251 工况与数值待核（Muse P3）',
    (398, 'SKE6510SSHEVR2'): '未决·P3：398批 问界M7新款行 242/165/175/251 工况与数值待核（Muse P3）',
    (379, 'LXA6511SHEVX07'): '未决·P3：理想L8 176km 值未决（168km已证实WLTC；Muse P3/§1.5）',
    (379, 'LXA6511SHEVX08'): '未决·P3：理想L8 176km 值未决（168km已证实WLTC；Muse P3/§1.5）',
    (379, 'LXA6511SHEVX09'): '未决·P3：理想L8 176km 值未决（168km已证实WLTC；Muse P3/§1.5）',
    (405, 'FZ6530REEVWD19D'): '未决·P3：D19D 352km 疑似WLTC（仅比值线索，不作归属）（Muse P3）',
    (395, 'CSA6492LFSHEV1'): '未决·P3：智己LS6增程 290km 疑似WLTC（仅比值线索，无对照，不作归属）（Muse P3）',
    (395, 'CSA6492LFSHEV2'): '未决·P3：智己LS6增程 360km 疑似WLTC（对上市CLTC 450说法；仅比值线索，不作归属）（Muse P3）',
    (407, 'CSA6492LFSHEV3'): '未决·P3：智己LS6增程 290km 疑似WLTC（仅比值线索，无对照，不作归属）（Muse P3）',
    (407, 'CSA6492LFSHEV4'): '未决·P3：智己LS6增程 360km 疑似WLTC（对上市CLTC 450说法；仅比值线索，不作归属）（Muse P3）',
    (409, 'XMA6500KREEVA1'): '未决·P3：澎程N70 380km 疑似WLTC（对上市CLTC 505说法；仅比值线索，不作归属）（Muse P3）',
    (409, 'XMA6530KREEVA1'): '未决·P3：澎程N90 370km 疑似WLTC（对464说法；仅比值线索，不作归属）（Muse P3）',
    (411, 'LXA6530REEVX12'): "未决·P3：理想L9 340/332km 疑似WLTC（仅比值线索，不作归属；Muse池该行电池缺失、本地72.7）（Muse P3）",
    (342, 'TSL6480BEVBA2'): '未决·P3：ModelY 342批 540km 待核（Muse P3；过渡期批次不批量默认）',
    (402, 'SQR6481CHEVT286'): '未决·P3：风云T9 32.659kWh版 160km 待核（100/160值族核实行不含此电池版本；T9L同）（Muse P3）',
}

EV_P2_17 = ('不可用·P2：容量1.7kWh 比值出界簇（跨越星光8行+SC6471EAPBEV，续航144~276km/kWh物理不可能），'
            '电池/续航疑似污染，续航值不可用于分析；原值不动（Muse P2）')
EV_P2_442 = ('不可用·P2：44.2kWh 配 44~161km 低续航比值出界簇（对照：家越07 408/409批 44.2kWh/240km/5.43 证明44.2为真实电池值，'
             '问题在续航值）；原值不动（Muse P2）')


def classify(rows):
    """返回 per-row (cycle, cltc, wltc, ev, grade) 列表与统计。"""
    plan = [None] * len(rows)
    stat = {}
    def setp(i, cycle, cltc=None, wltc=None, ev=None, grade='blank'):
        assert plan[i] is None, f'行{i+2}重复分类'
        plan[i] = (cycle, cltc, wltc, ev, grade)
        stat[grade] = stat.get(grade, 0) + 1

    p1_keys = set(P1_ROWS); p3_keys = set(P3_ROWS)
    for i, r in enumerate(rows):
        code, b = _code(r), _b(r)
        cap = r[14]; rng = r[13]; ptype = str(r[9] or '')
        # --- 1. 逐项/同族 WLTC ---
        if code in WLTC_DUAL and b == WLTC_DUAL[code][0]:
            _, cltc_v, ev = WLTC_DUAL[code]
            setp(i, 'WLTC', cltc=cltc_v, wltc=rng, ev=ev, grade='WLTC-逐项')
            continue
        if code in WLTC_POOL and b == WLTC_POOL[code][0]:
            _, ev = WLTC_POOL[code]
            setp(i, 'WLTC', wltc=rng, ev=ev, grade='WLTC-逐项/同族')
            continue
        # 18.316 WLTC 值族（排除汉双电机 P3 行）
        if str(cap) == '18.316' and (b or 0) >= 346 and isinstance(rng, (int, float)) and rng in WLTC_VALUE_SET \
                and (b, code) not in p3_keys:
            setp(i, 'WLTC', wltc=rng, ev=EV_BYD_WLTC, grade='WLTC-同族值族')
            continue
        # --- 2. NEDC ---
        if b == 346 and code in NEDC_MY_CODES:
            setp(i, 'NEDC', ev=EV_NEDC_MY, grade='NEDC-逐项')
            continue
        if b == 345 and str(cap) == '18.316':
            setp(i, 'NEDC', ev=EV_NEDC_BYD, grade='NEDC-同族值族')
            continue
        # --- 3. CLTC 逐项 ---
        if code in CLTC_POOL and b == CLTC_POOL[code][0]:
            _, ev = CLTC_POOL[code]
            setp(i, 'CLTC', cltc=rng, ev=ev, grade='CLTC-逐项')
            continue
        # --- 4. M7 证书值家族 ---
        if code.startswith('SKE6500SSHEV') and 373 <= (b or 0) <= 389:
            try:
                pw = int(str(r[19]).split('/')[0])
            except Exception:
                pw = None
            if cap == 37.389:
                setp(i, '未区分', ev=EV_M7_BASE + '；本行37.389kWh版对照未决，两值列留空', grade='M7-证书值')
            elif cap == 40.06 and pw == 200:
                setp(i, '未区分', cltc=240, wltc=200,
                     ev=EV_M7_BASE + '；本行后驱40.06kWh→对照 CLTC 240 / WLTC 200', grade='M7-证书值')
            elif cap == 40.06 and pw in (330, 331, 400):
                setp(i, '未区分', cltc=210,
                     ev=EV_M7_BASE + '；本行四驱40.06kWh→对照 CLTC 210；WLTC 162与175冲突留空', grade='M7-证书值')
            else:
                raise AssertionError(f'M7 行未分型: {b} {code} cap={cap} pw={pw}')
            continue
        # --- 5. P1 / P3 ---
        if (b, code) in p1_keys:
            setp(i, '未区分', ev=P1_ROWS[(b, code)], grade='P1-存疑')
            continue
        if (b, code) in p3_keys:
            setp(i, '未区分', ev=P3_ROWS[(b, code)], grade='P3-未决')
            continue
        # --- 6. P2 两簇（BEV 1.7 簇叠加法规默认；44.2 簇 PHEV 未区分） ---
        if str(cap).strip() == '1.7' and ptype == 'BEV':
            setp(i, 'CLTC', cltc=rng, ev=EV_LAW + '；' + EV_P2_17, grade='P2-1.7kWh簇')
            continue
        if cap == 44.2 and ptype == 'PHEV' and isinstance(rng, (int, float)) and 0 < rng < 200:
            setp(i, '未区分', ev=EV_P2_442, grade='P2-44.2kWh簇')
            continue
        # --- 7. 法规默认 CLTC（BEV ≥350 有值有型号） ---
        if ptype == 'BEV' and (b or 0) >= 350 and rng is not None and code:
            setp(i, 'CLTC', cltc=rng, ev=EV_LAW, grade='法规默认CLTC')
            continue
        # --- 8. 过渡期 BEV（341~349 未逐项） ---
        if ptype == 'BEV' and 341 <= (b or 0) <= 349 and rng is not None:
            setp(i, '未区分', ev=EV_TRANS, grade='过渡期BEV')
            continue
        # --- 9. 默认未区分（PHEV 全部未标注行 / BEV 无值或无型号 / 412 公示等） ---
        setp(i, '未区分', grade='未区分-默认')
    return plan, stat


def main():
    apply = '--apply' in sys.argv
    wb = openpyxl.load_workbook(XL)
    ws = wb[SHEET]
    assert ws.max_row == 5476 and ws.max_column == 41, f'前置断言失败: {ws.max_row}x{ws.max_column}'
    rows = list(ws.iter_rows(min_row=2, max_row=5476, values_only=True))
    plan, stat = classify(rows)

    # 颜色查重
    lg = wb['颜色说明']
    legend_fills = {c.fill.start_color.rgb for row in lg.iter_rows(min_row=2) for c in row if c.fill and c.fill.fill_type == 'solid'}
    assert 'FFD9F2F2' not in legend_fills, '浅青底纹与既有图例颜色冲突'

    # 分类统计与分布
    from collections import Counter
    cyc = Counter(p[0] for p in plan)
    print('=== 分类统计 ===')
    for k, v in sorted(stat.items(), key=lambda x: -x[1]): print(f'  {k}: {v}')
    print('=== 循环分布 ===', dict(cyc))

    # 逐名单核验（应全部命中）
    named = list(WLTC_DUAL) + list(WLTC_POOL) + list(CLTC_POOL)
    hit = {_code(r) for r in rows}
    miss = [c for c in named if c not in hit]
    assert not miss, f'名单行未命中: {miss}'
    m7n = sum(1 for r in rows if _code(r).startswith('SKE6500SSHEV') and 373 <= (_b(r) or 0) <= 389)
    assert m7n == 15, f'M7 行数 {m7n} != 15'
    p2_442 = sum(1 for r in rows if r[14] == 44.2 and str(r[9]) == 'PHEV' and isinstance(r[13], (int, float)) and 0 < r[13] < 200)
    print(f'M7={m7n}  P2-44.2簇={p2_442}  P2-1.7簇={stat.get("P2-1.7kWh簇",0)}')

    if not apply:
        print('DRY-RUN 通过（未写盘）。--apply 执行落库。')
        return

    # 备份（存在性检查，防毒化槽位）
    if os.path.exists(BACKUP):
        print(f'备份已存在，跳过：{BACKUP}')
    else:
        shutil.copy2(XL, BACKUP)
        assert os.path.getsize(BACKUP) == os.path.getsize(XL), '备份大小校验失败'
        print(f'已备份 {BACKUP}')

    # 写表头
    hdr_style_src = ws.cell(row=1, column=41)
    for j, h in enumerate(NEW_HEADERS):
        c = ws.cell(row=1, column=42 + j, value=h)
        c.font = copy(hdr_style_src.font); c.alignment = copy(hdr_style_src.alignment)
        c.border = copy(hdr_style_src.border); c.fill = FILL

    # 台账序号
    led = wb['变更记录']
    seq = int(str(led.cell(row=led.max_row, column=1).value)) + 1
    led_start = seq
    now = '2026-10-26'
    def ledrow(seq_, b, code, pt, field, newv, src):
        led.append([seq_, b, code, pt, field, '新增列-工况标注', '(空)', newv, src])
        return seq_ + 1

    def src_of(grade):
        if grade.startswith(('WLTC', 'NEDC', 'CLTC-逐项')): return f'{SRC} §1.2/§1.6 逐项核实结论'
        if grade == 'M7-证书值': return f'{SRC} §1.5 证书值现行处理'
        if grade.startswith('P1'): return f'{SRC} P1 存疑清单'
        if grade.startswith('P2'): return f'{SRC} P2 比值出界簇'
        if grade.startswith('P3'): return f'{SRC} P3 未决清单'
        return SRC

    agg = Counter()
    for i, (cycle, cltc, wltc, ev, grade) in enumerate(plan):
        rn = i + 2
        r = rows[i]
        b, code, pt = _b(r), _code(r), str(r[9] or '')
        vals = [cycle, cltc, wltc, ev]
        wrote = []
        for j, v in enumerate(vals):
            if v is None: continue
            c = ws.cell(row=rn, column=42 + j, value=v)
            c.fill = FILL
            wrote.append(NEW_HEADERS[j])
        agg[grade] += 1
        if grade != '未区分-默认' and grade != '法规默认CLTC' and grade != '过渡期BEV':
            for f in wrote:
                v = dict(zip(NEW_HEADERS, vals))[f]
                seq = ledrow(seq, b, code, pt, f, str(v), src_of(grade))
    # 聚合登记（法规默认/过渡期/默认未区分）
    n_law = stat.get('法规默认CLTC', 0) + stat.get('CLTC-逐项', 0) + stat.get('P2-1.7kWh簇', 0)
    seq = ledrow(seq, '（整表）', '（整表）', 'BEV', '测试循环', f'新增列初始填充：CLTC {n_law} 行（法规默认{stat.get("法规默认CLTC",0)}+逐项{stat.get("CLTC-逐项",0)}+P2簇{stat.get("P2-1.7kWh簇",0)}）', 'GB/T 18386.1-2021；' + SRC)
    seq = ledrow(seq, '（整表）', '（整表）', 'BEV', 'CLTC续航(km)', f'新增列初始填充：{n_law} 行（=测试循环CLTC行的池值镜像，逐字复制不改写）', SRC)
    seq = ledrow(seq, '（整表）', '（整表）', 'BEV', '续航工况证据', f'新增列初始填充：{n_law} 行（法规默认/逐项核实/P2加注三类证据文本）', SRC)
    n_tr = stat.get('过渡期BEV', 0)
    seq = ledrow(seq, '（整表）', '（整表）', 'BEV', '测试循环/续航工况证据', f'新增列初始填充：过渡期未区分 {n_tr} 行（341~349批，证据列注明不批量默认）', SRC + ' §1.3')
    n_dflt = stat.get('未区分-默认', 0)
    seq = ledrow(seq, '（整表）', '（整表）', '', '测试循环', f'新增列初始填充：未区分(默认态) {n_dflt} 行（PHEV未标注行/BEV无续航值行/412公示等）', '默认态')
    print(f'台账：{led_start}..{seq-1}（+{seq-led_start} 行）')

    # 颜色说明 + 字段映射说明
    lg.append(['浅青底纹', '续航工况标注（新增列，v4.9.35）',
               '测试循环/CLTC续航/WLTC续航/续航工况证据四列；不改动既有值；证据分级：逐项核实/同族同值/佐证级/法规默认/证书值/存疑/未决/不可用'])
    fm = wb['字段映射说明']
    fm_notes = [
        ('新增列(v4.9.35)', '测试循环', 'Muse工况核实2026-10-25/法规默认规则', 'CLTC/WLTC/NEDC/未区分四值；插混≥350批批量WLTC默认提案(Muse B2,实测1449行)未经用户批准未执行'),
        ('新增列(v4.9.35)', 'CLTC续航(km)', '池值CLTC镜像', '测试循环=CLTC行的纯电续航池值逐字镜像；双值并列行(零跑D19)填官方CLTC对照值；M7证书值行填已证实对照值(240/210)'),
        ('新增列(v4.9.35)', 'WLTC续航(km)', '池值WLTC镜像', '测试循环=WLTC行的池值镜像；D19行填官方WLTC=池值；M7后驱40.06行填对照200'),
        ('新增列(v4.9.35)', '续航工况证据', 'Muse工况核实/法规条文', '证据分级前缀：逐项核实/同族同值/佐证级/法规默认/证书值/存疑/未决/不可用；法规默认=GB/T 18386.1-2021(2021-10-01实施)350批起BEV'),
    ]
    for row in fm_notes: fm.append(list(row))

    wb.save(XL)
    print(f'已写盘：{XL}')
    # 机器可读台账
    out = {'date': now, 'version': 'v4.9.35', 'source': SRC, 'columns': NEW_HEADERS,
           'stats': stat, 'cycle_dist': dict(cyc), 'ledger_seq': [led_start, seq - 1],
           'rules': {'law_default': EV_LAW, 'transition': EV_TRANS,
                     'byd_wltc_value_family': sorted(WLTC_VALUE_SET),
                     'm7_mapping': '40.06+200kW→240/200; 40.06+330/400kW→210/-; 37.389→-/-'}, }
    with open(os.path.join(HERE, 'audit-output', 'cycle_annotation_20261026.json'), 'w', encoding='utf-8') as f:
        json.dump(out, f, ensure_ascii=False, indent=2, default=str)
    print('audit-output/cycle_annotation_20261026.json 已生成')

if __name__ == '__main__':
    main()
