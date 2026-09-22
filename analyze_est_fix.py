# -*- coding: utf-8 -*-
"""analyze_est_fix.py — 枚举 integrate_dcd_refill.py 的 est 红修候选(逐格明细)

复刻 integrate 的 est_keys 逻辑,但只报告不写入。重点核验:
  1. 每个候选的 (批次,型号,字段) 现值 → DCD候选值
  2. 车长/轴距 是否成对变化(W5 铁律)
  3. 同型号同行内 车长/轴距 的最终状态是否自洽
"""
import json
import os
import re
from collections import defaultdict

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, 'raw', 'dcd_refill')
WB = os.path.join(HERE, 'NEV公告参数汇总表_合并版（341~410批）.xlsx')
FUEL = os.path.join(SRC, 'dcd_fuel_form.json')

FIELD_MAP = {
    '车长(mm)': 'length', '轴距(mm)': 'wheelbase', '整备质量(kg)': 'curb_weight',
    '纯电续航里程(km)': 'cltc_recharge_mileage', '百公里电耗(kWh/100km)': 'power_consumption',
    '电池容量(kWh)': 'battery_capacity', '电池类型': 'battery_type',
    '电池能量密度(Wh/kg)': 'battery_energy_density',
}


def fnum(v):
    try:
        return float(str(v).split('/')[0].replace(',', '').replace('km', '').strip())
    except Exception:
        return None


def main():
    mapping = json.load(open(os.path.join(SRC, 'series_mapping.json'), encoding='utf-8'))
    accepted = [m for m in mapping['reviewed'] if m.get('ok')]
    fuel = json.load(open(FUEL, encoding='utf-8')) if os.path.exists(FUEL) else {}
    fuel_ff = {}
    for sid_f, lst in fuel.items():
        for sp in (lst or []):
            if isinstance(sp, dict) and sp.get('id'):
                fuel_ff[(str(sid_f), str(sp['id']))] = str(sp.get('ff') or '')

    wb = openpyxl.load_workbook(WB, read_only=True)
    ws = wb['NEV公告参数汇总']
    hdr = [str(c.value or '').strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    hi = {h: i for i, h in enumerate(hdr)}

    led = wb['变更记录']
    est_keys = set()
    skip_keys = set()
    for r in led.iter_rows(min_row=2, values_only=True):
        if isinstance(r[0], int):
            k = (str(r[1] or '').strip(), str(r[2] or '').strip(), str(r[4] or '').strip())
            if '估算' in str(r[5] or '') or '估算' in str(r[8] or ''):
                est_keys.add(k)
            if r[5] == '错配回退':
                skip_keys.add(k)

    series_files = {}
    for f in os.listdir(SRC):
        m = re.match(r'dcd_s(\d+)\.json$', f)
        if m:
            series_files[m.group(1)] = json.load(open(os.path.join(SRC, f), encoding='utf-8'))

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    cands = []  # (sid, kw, batch, code, pt, field, cur, new)
    for m in accepted:
        sid = str(m['sid'])
        if sid not in series_files:
            continue
        dcd = series_files[sid]
        kw = m['kw'].lower().replace(' ', '')
        kw2 = kw.replace('plus', '')
        ALIAS = {'polestar2': ['极星2', 'polestar2'], '北汽eu5': ['eu5', '北京eu5']}
        extra = ALIAS.get(kw, [])
        for row in rows:
            g = (str(row[hi['通用名称']] or '') + str(row[hi['车型名称']] or '')).lower().replace(' ', '')
            pt = str(row[hi['动力类型']] or '')
            code = str(row[hi['产品型号']] or '').strip()
            batch = str(row[hi['批次']] or '').strip()
            if kw not in g and kw2 not in g and not any(a in g for a in extra):
                continue
            if not ('纯电' in pt or '插电' in pt or '增程' in pt or 'BEV' in pt or 'PHEV' in pt):
                continue
            is_bev = '纯电' in pt or 'BEV' in pt
            trims = []
            for c in dcd['cars']:
                v = c['vals']
                ff = fuel_ff.get((sid, str(c.get('id'))), '')
                if ff == '汽油':
                    continue
                if is_bev and not v.get('battery_capacity'):
                    continue
                if (not is_bev) and not v.get('battery_capacity') and 'DM' not in c['name'] and 'EV' not in c['name']:
                    continue
                trims.append(c)
            if not trims:
                continue
            r_rng, r_cw = fnum(row[hi['纯电续航里程(km)']]), fnum(row[hi['整备质量(kg)']])
            scored = []
            for c in trims:
                v = c['vals']
                nm_km = re.search(r'(\d{3,4})\s*km', c['name'])
                cltc = fnum(v.get('cltc_recharge_mileage')) or (float(nm_km.group(1)) if nm_km else None)
                cw = fnum(v.get('curb_weight'))
                score = 0
                if r_rng and cltc:
                    d = abs(r_rng - cltc)
                    score += (0 if d <= 15 else (1 if d <= 30 else 3))
                elif r_rng and not cltc:
                    score += 1
                if r_cw and cw:
                    d = abs(r_cw - cw)
                    score += (0 if d <= 40 else (1 if d <= 90 else 3))
                scored.append((score, c))
            scored.sort(key=lambda t: t[0])
            best = scored[0][0]
            tied = [c for s, c in scored if s == best]
            for fld, key in FIELD_MAP.items():
                if (batch, code, fld) in skip_keys:
                    continue
                cur = row[hi[fld]]
                if cur in (None, ''):
                    continue  # 空格走 green 路径,不在本分析范围
                key_est = (batch, code, fld)
                if key_est not in est_keys:
                    continue
                cand_vals = {fnum(c['vals'].get(key)) for c in tied}
                cand_vals.discard(None)
                if len(cand_vals) == 1 and fnum(cur) is not None and fnum(cur) != next(iter(cand_vals)):
                    cands.append({'sid': sid, 'kw': m['kw'], 'batch': batch, 'code': code, 'pt': pt,
                                  'field': fld, 'cur': fnum(cur), 'new': next(iter(cand_vals)),
                                  'tied_n': len(tied), 'best_score': best,
                                  'cur_len': fnum(row[hi['车长(mm)']]), 'cur_wb': fnum(row[hi['轴距(mm)']]),
                                  'new_len': next(iter({fnum(c['vals'].get('length')) for c in tied} - {None}), None),
                                  'new_wb': next(iter({fnum(c['vals'].get('wheelbase')) for c in tied} - {None}), None)})

    print(f'est 红修候选共 {len(cands)} 格:')
    by_field = defaultdict(int)
    for c in cands:
        by_field[c['field']] += 1
    print('按字段:', dict(by_field))
    print()
    # 按 (批次,型号) 分组看车长/轴距配对
    by_row = defaultdict(list)
    for c in cands:
        by_row[(c['batch'], c['code'])].append(c)
    unparied = 0
    for (b, code), cs in sorted(by_row.items()):
        fields = {c['field'] for c in cs}
        pair_tag = ''
        if '车长(mm)' in fields or '轴距(mm)' in fields:
            has_l = '车长(mm)' in fields
            has_w = '轴距(mm)' in fields
            if has_l != has_w:
                pair_tag = '  ⚠️W5单字段!'
                unparied += 1
        for c in cs:
            print(f"  {c['batch']} {c['code']} {c['field']}: {c['cur']} → {c['new']}  "
                  f"(kw={c['kw']}, tied={c['tied_n']}, score={c['best_score']}, "
                  f"cur_L/W={c['cur_len']}/{c['cur_wb']}, new_L/W={c['new_len']}/{c['new_wb']}){pair_tag}")
    print(f'\n⚠️ W5 单字段行数: {unpaired}')


if __name__ == '__main__':
    main()
