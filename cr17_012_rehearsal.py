# -*- coding: utf-8 -*-
"""CR-17-012: 落盘彩排（rehearsal）——在临时副本上模拟并发会话的 41 列扩展与销量列状态，
再跑 cr17_011 的 --check/--apply，逐格比对"只应改我说的那些格"。

用法：python cr17_012_rehearsal.py
输出：audit-output/cr17_012_rehearsal_20261019.json
"""
import io, json, os, shutil, subprocess, sys, hashlib
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
WB = os.path.join(HERE, "NEV公告参数汇总表_合并版（341~410批）.xlsx")
TMP = os.path.join(os.environ.get("TEMP", r"C:\Windows\Temp"), "probe17", "rehearsal")
SIM = os.path.join(TMP, "NEV公告参数汇总表_合并版（341~410批）.xlsx")
OUT = os.path.join(HERE, "audit-output", "cr17_012_rehearsal_20261019.json")

os.makedirs(TMP, exist_ok=True)
sys.stdout = io.TextIOWrapper(open(os.path.join(TMP, "rehearsal.log"), "wb"), encoding="utf-8", errors="replace")


def sha256(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for c in iter(lambda: f.read(1 << 20), b""):
            h.update(c)
    return h.hexdigest()


def norm(v):
    return None if v is None else (v if not isinstance(v, str) else v.strip())


def main():
    os.makedirs(TMP, exist_ok=True)
    if os.path.exists(SIM):
        os.remove(SIM)
    shutil.copy2(WB, SIM)
    shutil.rmtree(os.path.join(TMP, "_test_archive"), ignore_errors=True)
    sha0 = sha256(SIM)
    print("[1] 副本就绪", SIM, len(open(SIM, "rb").read()), "B sha=", sha0[:12])

    # ---- 模拟并发会话的写入：尾插 2 列 + 状态/车系ID 填充 ----
    wb = load_workbook(SIM)
    ws = wb["NEV公告参数汇总"]
    hdr = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    assert hdr.index("月销量(辆)") + 1 == 36, hdr.index("月销量(辆)")
    c_sales, c_mon, c_upd = hdr.index("月销量(辆)") + 1, hdr.index("销量数据月份") + 1, hdr.index("销量更新日期") + 1
    c_id, c_st = ws.max_column + 1, ws.max_column + 2
    ws.cell(row=1, column=c_id).value = "车系ID"
    ws.cell(row=1, column=c_st).value = "销量覆盖状态"
    filled = 0
    for r in range(2, ws.max_row + 1):
        if norm(ws.cell(row=r, column=c_sales).value) is not None:
            ws.cell(row=r, column=c_id).value = f"TEST-{r}"
            ws.cell(row=r, column=c_mon).value = "2026-09"
            ws.cell(row=r, column=c_upd).value = "2026-10-19"
            ws.cell(row=r, column=c_st).value = "在售-已匹配"
            filled += 1
        else:
            ws.cell(row=r, column=c_st).value = "未匹配-已停产或未上量"
    wb.save(SIM)
    sha_sim = sha256(SIM)
    print(f"[2] 模拟并发写入完成：列 {ws.max_column}、有销量行 {filled}、sha={sha_sim[:12]}")

    # 快照（逐格）用于比对
    def snap(path):
        w2 = load_workbook(path, read_only=True, data_only=True)
        s = w2["NEV公告参数汇总"]
        rows = [[norm(v) for v in r] for r in s.iter_rows(values_only=True)]
        led = w2["变更记录"].max_row
        t = w2["批次时间表"].max_row
        m = w2["字段映射说明"].max_row
        w2.close()
        return rows, led, t, m

    pre, pre_led, pre_t, pre_m = snap(SIM)

    # ---- 跑 cr17_011 --check ----
    r1 = subprocess.run([sys.executable, os.path.join(HERE, "cr17_011_apply_rows.py"), "--check", "--wb", SIM],
                        capture_output=True, text=True, encoding="utf-8", errors="replace", cwd=HERE)
    print("[3] --check rc=", r1.returncode)
    print(r1.stdout[-1800:])
    if r1.returncode != 0:
        print("STDERR:", r1.stderr[-1200:])
        return

    # ---- 跑 --apply ----
    r2 = subprocess.run([sys.executable, os.path.join(HERE, "cr17_011_apply_rows.py"), "--apply",
                         "--expect-sha", sha_sim, "--wb", SIM, "--no-backup"],
                        capture_output=True, text=True, encoding="utf-8", errors="replace", cwd=HERE)
    print("[4] --apply rc=", r2.returncode)
    print(r2.stdout[-2500:])
    if r2.returncode != 0:
        print("STDERR:", r2.stderr[-1500:])
        return

    # ---- 比对 ----
    post, post_led, post_t, post_m = snap(SIM)
    n_old = len(pre)
    diff = []
    for i in range(n_old):
        a, b = pre[i], post[i]
        if a == b:
            continue
        for j in range(max(len(a), len(b))):
            va = a[j] if j < len(a) else None
            vb = b[j] if j < len(b) else None
            if va != vb:
                diff.append({"row": i + 1, "col0": j, "before": va, "after": vb})
    print(f"[5] 原有 {n_old} 行 × {len(pre[0])} 列 逐格差异 = {len(diff)}（期望 = 富化格数 66）")
    from collections import Counter
    print("    差异列分布:", Counter(d["col0"] for d in diff).most_common())
    for d in diff[:8]:
        print("     ", d)

    new_rows = post[n_old:]
    print(f"[6] 新增行 = {len(new_rows)}（期望 83）；新增行非空列统计:")
    cnt = Counter()
    for r in new_rows:
        for j, v in enumerate(r):
            if v is not None:
                cnt[j] += 1
    print("    ", sorted(cnt.items()))
    print(f"[7] 台账 {pre_led} → {post_led}（+{post_led - pre_led}，期望 +91）")
    print(f"[8] 批次时间表 {pre_t} → {post_t}（+{post_t - pre_t}，期望 +2）")
    print(f"[9] 字段映射说明 {pre_m} → {post_m}（+{post_m - pre_m}，期望 +2）")

    # 新增行的并发列应严格留空/占位
    c_id0, c_st0, c_sales0 = c_id - 1, c_st - 1, c_sales - 1
    bad = []
    for k, r in enumerate(new_rows):
        for j, label in ((c_sales0, "月销量"), (c_sales0 + 1, "销量数据月份"), (c_sales0 + 2, "销量更新日期")):
            if r[j] is not None:
                bad.append((k, label, r[j]))
        if r[c_id0] is not None:
            bad.append((k, "车系ID", r[c_id0]))
        if r[c_st0] != "未匹配-命名待锚定":
            bad.append((k, "销量覆盖状态", r[c_st0]))
    print(f"[10] 新增行并发列传参违规 = {len(bad)}（期望 0）", bad[:5])

    # 幂等性：再跑一次 --check 应 0 新增
    r3 = subprocess.run([sys.executable, os.path.join(HERE, "cr17_011_apply_rows.py"), "--check", "--wb", SIM],
                        capture_output=True, text=True, encoding="utf-8", errors="replace", cwd=HERE)
    print("[11] 幂等 --check rc=", r3.returncode)
    print(r3.stdout[-900:])

    rep = {"sim_sha_before": sha_sim, "sim_sha_after": sha256(SIM), "orig_rows": n_old,
           "diff_cells": len(diff), "diff_cols": dict(Counter(d["col0"] for d in diff)),
           "new_rows": len(new_rows), "ledger_add": post_led - pre_led,
           "time_add": post_t - pre_t, "map_add": post_m - pre_m,
           "peer_col_violations": len(bad), "check_rc": r1.returncode, "apply_rc": r2.returncode,
           "idem_check_rc": r3.returncode}
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(rep, f, ensure_ascii=False, indent=2)
    print("[12] 彩排报告 →", OUT)


main()
