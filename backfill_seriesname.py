#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""媒体系列名回填：对缺失「车型名称」的行，用已记录媒体来源 URL 对应的车系名补空（只补空）。"""
import re
import time
import ssl
import glob
import json
import urllib.request
from collections import defaultdict

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/126.0 Safari/537.36")
CTX = ssl.create_default_context()
FILL_MEDIA = PatternFill("solid", fgColor="F8CBAD")


def get(url):
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept": "application/json"})
    with urllib.request.urlopen(req, timeout=25, context=CTX) as r:
        return json.loads(r.read().decode("utf-8", "ignore"))


def series_name(sid):
    u = (f"https://car-web-api.autohome.com.cn/car/spec/listSpec"
         f"?type=0x001f&from=1&pm=1&pluginversion=11.64.8&seriesid={sid}")
    return (get(u).get("result") or {}).get("seriesname")


def main():
    wb_path = sorted(glob.glob("NEV公告参数汇总表_合并版*.xlsx"))[-1]
    wb = load_workbook(wb_path)
    ws = wb["NEV公告参数汇总"]
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    hi = {h: i for i, h in enumerate(headers)}
    need = []  # (excel_row, seriesid)
    for row in ws.iter_rows(min_row=2):
        nm = row[hi["车型名称"]].value
        if nm is not None and str(nm).strip():
            continue
        src = row[hi["媒体校验来源"]].value
        if not src:
            continue
        m = re.search(r"series/(\d+)", str(src))
        if m:
            need.append((row[0].row, m.group(1)))
    print(f"rows missing 车型名称 with media source: {len(need)}")
    cache = {}
    filled = 0
    for excel_row, sid in need:
        if sid not in cache:
            try:
                cache[sid] = series_name(sid)
            except Exception as e:
                print(f"  fail {sid}: {e}")
                cache[sid] = None
            time.sleep(0.15)
        name = cache.get(sid)
        if not name:
            continue
        cell = ws.cell(row=excel_row, column=hi["车型名称"] + 1)
        cell.value = name
        cell.fill = FILL_MEDIA
        filled += 1
    wb.save(wb_path)
    print(f"filled 车型名称: {filled}")


if __name__ == "__main__":
    main()
