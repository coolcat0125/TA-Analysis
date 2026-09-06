# -*- coding: utf-8 -*-
"""v4.2.1 数据清洗：清空六字段物理范围越界的无效值（原值入变更记录，单元格浅红底纹）。

复用 audit_data.py 的记录枚举口径（跳过全空行，records 从第 2 行起算）。
保留例外：电机总功率 2220kW（仰望U9 Xtreme，4x555kW 真实值）。
"""
import math
import re
import sys
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

sys.stdout.reconfigure(encoding="utf-8")

SRC = "NEV公告参数汇总表_合并版（341~410批）.xlsx"
FIELDS = {
    "整备质量": ("整备质量(kg)", (300.0, 10000.0)),
    "纯电续航": ("纯电续航里程(km)", (5.0, 2000.0)),
    "电池容量": ("电池容量(kWh)", (1.0, 300.0)),
    "电池能量密度": ("电池能量密度(Wh/kg)", (30.0, 400.0)),
    "百公里电耗": ("百公里电耗(kWh/100km)", (3.0, 40.0)),
    "电机总功率": ("电机总功率(kW)", (1.0, 1500.0)),
}
KEEP = {("电机总功率", 3331)}  # 记录号口径：仰望U9 Xtreme，真实值
RED = PatternFill(fill_type="solid", start_color="FFC7CE", end_color="FFC7CE")


def first_number(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value) if math.isfinite(value) else None
    text = str(value).replace(",", "")
    m = re.search(r"-?\d+(?:\.\d+)?", text)
    return float(m.group()) if m else None


def is_missing(value):
    return value is None or (isinstance(value, str) and value.strip() in {"", "/", "-", "--", "N/A"})


wb = load_workbook(SRC)
ws = wb["NEV公告参数汇总"]
rows_iter = ws.iter_rows(values_only=False)
headers = [str(c.value).strip() if c.value is not None else "" for c in next(rows_iter)]
col_idx = {h: i for i, h in enumerate(headers)}

# 与 audit_data.py 相同的记录过滤，同时记录真实 Excel 行号
records = []
for cell_row in rows_iter:
    values = [c.value for c in cell_row]
    if any(not is_missing(v) for v in values):
        records.append(cell_row)
print(f"记录数 {len(records)}")

targets = []  # (field_key, record_no, excel_row, cell, original, batch, model, dtype, header)
for key, (header, (lo, hi)) in FIELDS.items():
    ci = col_idx.get(header)
    if ci is None:
        raise SystemExit(f"找不到列 {header}")
    for rec_no, cell_row in enumerate(records, start=2):
        cell = cell_row[ci]
        value = cell.value
        if is_missing(value):
            continue
        parsed = first_number(value)
        if parsed is None or lo <= parsed <= hi:
            continue
        if (key, rec_no) in KEEP:
            print(f"保留 KEEP {key} 记录{rec_no} 原值={value!r}")
            continue
        b = records[rec_no - 2][col_idx["批次"]].value
        m = records[rec_no - 2][col_idx["产品型号"]].value
        d = records[rec_no - 2][col_idx["动力类型"]].value
        targets.append((key, rec_no, cell.row, cell, value, b, m, d, header))

print(f"待清空 {len(targets)} 格")
if len(targets) != 68:
    ans = input(f"预期 68 格，实际 {len(targets)} 格，继续? [y/N] ")
    if ans.strip().lower() != "y":
        sys.exit(1)

ledger = wb["变更记录"]
last_row = ledger.max_row
seq = int(ledger.cell(row=last_row, column=1).value or 0)
print(f"台账续接序号 {seq + 1}")

for key, rec_no, excel_row, cell, value, batch, model, dtype, header in targets:
    orig = str(value)
    if isinstance(value, float):
        orig = repr(value)
    seq += 1
    ledger.append([seq, batch, model, dtype, header, "数据清洗更正", orig, "(空)",
                   "audit_data.py范围检查越界清洗 2026-09-06"])
    cell.value = None
    cell.fill = RED

wb.save(SRC)
print(f"完成：清空 {len(targets)} 格，台账新增 {len(targets)} 行（现最大序号 {seq}）")
