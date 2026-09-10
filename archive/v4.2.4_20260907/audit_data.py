#!/usr/bin/env python3
"""TA-Analysis data-quality and version audit.

Copyright © 2026 David YE
SPDX-License-Identifier: MIT
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import math
import re
import sys
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


COPYRIGHT = "Copyright © 2026 David YE"
DEFAULT_GATES = {
    "expected_min_records": 4494,
    "required_headers": ["批次", "产品型号", "动力类型"],
    "coverage_floors": {
        "百公里电耗": 0.70,
        "电池能量密度": 0.85,
        "电池容量": 0.90,
        "纯电续航": 0.70,
        "电机总功率": 0.65,
    },
}
RANGE_RULES = {
    "整备质量": (300.0, 10000.0),
    "纯电续航": (5.0, 2000.0),
    "电池容量": (1.0, 300.0),
    "电池能量密度": (30.0, 400.0),
    "百公里电耗": (3.0, 40.0),
    "电机总功率": (1.0, 1500.0),
    "发动机排量": (300.0, 8000.0),
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def first_number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value) if math.isfinite(value) else None
    text = str(value).replace(",", "")
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    return float(match.group()) if match else None


def is_missing(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() in {"", "/", "-", "--", "N/A"})


def find_column(headers: list[str], key: str) -> int | None:
    for index, header in enumerate(headers):
        if key in header:
            return index
    return None


def load_gates(path: Path | None) -> dict[str, Any]:
    gates = json.loads(json.dumps(DEFAULT_GATES))
    if not path:
        return gates
    with path.open(encoding="utf-8") as source:
        supplied = json.load(source)
    for name in ("expected_min_records", "required_headers", "coverage_floors"):
        if name in supplied:
            gates[name] = supplied[name]
    return gates


def discover_source(explicit: str | None) -> Path:
    if explicit:
        return Path(explicit)
    candidates = sorted(Path.cwd().glob("NEV公告参数汇总表_合并版*.xlsx"))
    if not candidates:
        raise FileNotFoundError("未找到 NEV公告参数汇总表_合并版*.xlsx；请用 --input 指定源表。")
    return candidates[-1]


def audit(input_path: Path, gates: dict[str, Any]) -> dict[str, Any]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    sheet = workbook["NEV公告参数汇总"] if "NEV公告参数汇总" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    records = [tuple(row) for row in rows if any(not is_missing(cell) for cell in row)]
    batch_col = find_column(headers, "批次")
    batches = Counter()
    for row in records:
        if batch_col is not None and batch_col < len(row) and not is_missing(row[batch_col]):
            batches[str(row[batch_col]).strip()] += 1

    coverage: dict[str, dict[str, Any]] = {}
    ranges: dict[str, dict[str, Any]] = {}
    for key, bounds in RANGE_RULES.items():
        column = find_column(headers, key)
        if column is None:
            continue
        populated = 0
        numeric = 0
        outliers = 0
        samples: list[dict[str, Any]] = []
        for row_index, row in enumerate(records, start=2):
            value = row[column] if column < len(row) else None
            if is_missing(value):
                continue
            populated += 1
            parsed = first_number(value)
            if parsed is None:
                continue
            numeric += 1
            if parsed < bounds[0] or parsed > bounds[1]:
                outliers += 1
                if len(samples) < 10:
                    samples.append({"excel_row": row_index, "value": value, "parsed": parsed})
        coverage[key] = {
            "header": headers[column],
            "populated": populated,
            "coverage": round(populated / len(records), 6) if records else 0,
            "missing": len(records) - populated,
        }
        ranges[key] = {
            "expected_range": list(bounds),
            "numeric_values": numeric,
            "outliers": outliers,
            "sample_outliers": samples,
        }

    required_missing = [item for item in gates["required_headers"] if find_column(headers, item) is None]
    failed_gates = []
    if len(records) < gates["expected_min_records"]:
        failed_gates.append(f"记录数 {len(records)} 小于最低门槛 {gates['expected_min_records']}")
    if required_missing:
        failed_gates.append("缺少必需列: " + ", ".join(required_missing))
    for key, minimum in gates["coverage_floors"].items():
        actual = coverage.get(key, {}).get("coverage")
        if actual is None:
            failed_gates.append(f"未找到覆盖率字段: {key}")
        elif actual < float(minimum):
            failed_gates.append(f"{key} 覆盖率 {actual:.1%} 低于门槛 {float(minimum):.1%}")

    return {
        "schema_version": "1.0.0",
        "copyright": COPYRIGHT,
        "generated_at_utc": dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat(),
        "input": {"filename": input_path.name, "sha256": sha256(input_path), "worksheet": sheet.title},
        "dataset": {"records": len(records), "columns": len(headers), "headers": headers, "batches": dict(sorted(batches.items()))},
        "coverage": coverage,
        "range_checks": ranges,
        "quality_gates": {"passed": not failed_gates, "failures": failed_gates, "configuration": gates},
    }


def markdown(report: dict[str, Any]) -> str:
    lines = [
        "# TA-Analysis 数据质量与版本审计报告",
        "",
        COPYRIGHT,
        "",
        f"- 生成时间（UTC）：{report['generated_at_utc']}",
        f"- 输入文件：`{report['input']['filename']}`",
        f"- SHA-256：`{report['input']['sha256']}`",
        f"- 工作表：`{report['input']['worksheet']}`",
        f"- 记录数：{report['dataset']['records']}；字段数：{report['dataset']['columns']}",
        "",
        "## 质量门禁",
        "",
        "**通过**" if report["quality_gates"]["passed"] else "**未通过**",
    ]
    if report["quality_gates"]["failures"]:
        lines.extend(f"- {failure}" for failure in report["quality_gates"]["failures"])
    lines.extend(["", "## 字段覆盖率", "", "| 字段 | 表头 | 已填 | 缺失 | 覆盖率 |", "|---|---|---:|---:|---:|"])
    for key, item in report["coverage"].items():
        lines.append(f"| {key} | {item['header']} | {item['populated']} | {item['missing']} | {item['coverage']:.1%} |")
    lines.extend(["", "## 物理范围检查", "", "| 字段 | 合理区间 | 数值样本 | 异常数 |", "|---|---|---:|---:|"])
    for key, item in report["range_checks"].items():
        low, high = item["expected_range"]
        lines.append(f"| {key} | {low}–{high} | {item['numeric_values']} | {item['outliers']} |")
    lines.extend(["", "## 批次分布", ""])
    for batch, count in report["dataset"]["batches"].items():
        lines.append(f"- {batch}: {count}")
    lines.extend(["", "---", "", "本报告由 `audit_data.py` 自动生成；请将 JSON 审计产物随每次数据批次更新一并归档。"])
    return "\n".join(lines) + "\n"


def main() -> int:
    parser = argparse.ArgumentParser(description="TA-Analysis 数据质量与版本审计")
    parser.add_argument("--input", help="源 Excel 文件；默认自动发现 341~410 批汇总表")
    parser.add_argument("--gates", help="质量门禁 JSON；默认使用脚本内置门槛")
    parser.add_argument("--output-dir", default="audit-output", help="报告输出目录")
    parser.add_argument("--strict", action="store_true", help="质量门禁未通过时返回非零状态码")
    args = parser.parse_args()
    input_path = discover_source(args.input)
    gates = load_gates(Path(args.gates) if args.gates else None)
    report = audit(input_path, gates)
    output = Path(args.output_dir)
    output.mkdir(parents=True, exist_ok=True)
    (output / "data_audit_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (output / "data_audit_report.md").write_text(markdown(report), encoding="utf-8")
    print(f"审计完成：{output / 'data_audit_report.json'}")
    return 1 if args.strict and not report["quality_gates"]["passed"] else 0


if __name__ == "__main__":
    sys.exit(main())
