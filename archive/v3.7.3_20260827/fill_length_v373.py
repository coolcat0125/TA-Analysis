# -*- coding: utf-8 -*-
"""
车长(mm) 补充 v3.7.3
====================
事实：仓库两份补充表（行业分析底表/数据补充比对表）均无车长实测数据。
口径：按用户指示「以轴距为基准」估算补充 —— 车长 = 轴距 × 车身形式系数，四舍五入到10mm。
     所有估算值以浅紫底纹标记，逐行记入变更记录（变更类型=轴距估算，数据来源注明非实测）；
     看板分析仍以轴距为基准（本脚本不改看板口径）。
系数依据（现代乘用车 车长/轴距 典型比值）：
     Car 1.72 / Lux Car 1.66 / SUV 1.69 / Lux SUV 1.66 / MPV 1.68 / Lux MPV 1.65 / Sports 1.64 / 其他 1.70
Copyright © 2026 David YE
"""
import openpyxl
import os
import sys
import time

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, 'NEV公告参数汇总表_合并版（341~410批）.xlsx')
SHEET = 'NEV公告参数汇总'
PURPLE = openpyxl.styles.PatternFill('solid', fgColor='D9D2E9')   # 轴距估算·浅紫底纹

RATIO = [
    ('Lux Car', 1.66), ('Lux SUV', 1.66), ('Lux MPV', 1.65), ('Sports', 1.64),
    ('Car', 1.72), ('SUV', 1.69), ('MPV', 1.68),
]


def seg_ratio(seg):
    s = str(seg or '').strip()
    for prefix, k in RATIO:
        if s.startswith(prefix):
            return k
    return 1.70


def main():
    print(f'[1/4] 加载工作簿: {os.path.basename(SRC)}')
    wb = openpyxl.load_workbook(SRC)
    ws = wb[SHEET]
    hdr = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(hdr)}
    c_batch, c_model = idx['批次'], idx['产品型号']
    c_seg, c_ab, c_lg = idx['细分市场'], idx['轴距(mm)'], idx['车长(mm)']
    n_rows = ws.max_row

    cr = wb['变更记录']
    last_seq = 0
    for r in range(cr.max_row, 1, -1):
        v = cr.cell(r, 1).value
        if isinstance(v, int):
            last_seq = v
            break

    # 颜色说明追加图例（幂等）
    cs = wb['颜色说明']
    has_legend = any(str(row[1] or '') == '轴距估算' for row in cs.iter_rows(values_only=True) if len(row) > 1)
    if not has_legend:
        cs.append(['浅紫底纹', '轴距估算',
                   '车长(mm)由轴距按车身形式系数估算补充（非公告实测值）；分析仍以轴距为基准，车长仅作参考维度'])

    print('[2/4] 以轴距为基准估算补充车长…')
    n_fill = 0
    n_skip = 0
    for ri in range(2, n_rows + 1):
        lg_cell = ws.cell(ri, c_lg + 1)
        if lg_cell.value not in (None, ''):
            continue
        ab = ws.cell(ri, c_ab + 1).value
        if not isinstance(ab, (int, float)) or not (1800 <= ab <= 4200):
            n_skip += 1
            continue
        seg = ws.cell(ri, c_seg + 1).value
        k = seg_ratio(seg)
        v = int(round(ab * k / 10.0) * 10)
        lg_cell.value = v
        lg_cell.fill = PURPLE
        last_seq += 1
        cr.append([last_seq, ws.cell(ri, c_batch + 1).value, ws.cell(ri, c_model + 1).value or '',
                   str(ws.cell(ri, idx['动力类型'] + 1).value or ''), '车长(mm)', '轴距估算',
                   '(空)', v, f'轴距{int(ab)}×系数{k}（v3.7.3估算，非实测）'])
        n_fill += 1

    filled = sum(1 for r in ws.iter_rows(min_row=2, min_col=c_lg + 1, max_col=c_lg + 1, values_only=True)
                 if r[0] not in (None, ''))
    print(f'      估算补充 {n_fill} 条 / 轴距缺失跳过 {n_skip} 条')
    print(f'      车长列填报率：{filled}/{n_rows-1}（{filled/(n_rows-1)*100:.1f}%）')

    print('[3/4] 保存…')
    saved = None
    try:
        wb.save(SRC)
        saved = SRC
        print('      已就地覆盖')
    except PermissionError:
        tmp = SRC.replace('（341~410批）', '（341~410批）v373tmp')
        wb.save(tmp)
        for i in range(6):
            time.sleep(20)
            try:
                os.replace(tmp, SRC)
                saved = SRC
                print(f'      重试成功（{i+1}）')
                break
            except PermissionError:
                pass
        if not saved:
            kept = SRC.replace('（341~410批）', '（341~410批·车长补充版）')
            os.replace(tmp, kept)
            saved = kept
            print(f'      原件被占用，保留为 {os.path.basename(kept)}')

    print('[4/4] 复检…')
    wb2 = openpyxl.load_workbook(saved, read_only=True)
    ws2 = wb2[SHEET]
    n_lg = 0
    for r in ws2.iter_rows(min_row=2, min_col=c_lg + 1, max_col=c_lg + 1, values_only=True):
        if r[0] not in (None, ''):
            n_lg += 1
    print(f'      车长(mm) 非空 {n_lg} / {n_rows-1}（{n_lg/(n_rows-1)*100:.1f}%）· 变更记录 {wb2["变更记录"].max_row-1} 行')
    wb2.close()
    print('✓ 完成')


if __name__ == '__main__':
    sys.exit(main())
