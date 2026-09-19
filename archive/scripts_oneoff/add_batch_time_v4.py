# -*- coding: utf-8 -*-
"""
底表新增「批次时间表」v4.0
===========================
需求：为数据底表补充每个批次的时间（精确日期；无精确日期则以年月标注）。
批次 = 工信部《道路机动车辆生产企业及产品公告》批次编号（341~410，本底表已收录 341~410，
样本含 358、397 批缺口，时间表仍按全段 341~410 补齐，用于批次定位）。

口径原则：
  - 「公告日期」以工信部官网（miit.gov.cn）公告页发布的发布日期为准，属权威精确值；
  - 未能逐批考证精确日的批次，「公告年月」依据相邻批次月度节奏推断并用黄底标注，
    明确标为「按月节奏推断」，不编造精确日期；
  - 「公示日期」为公告发布前工信部公示清单发布日期，仅记载已考证批次，其余留空。

用法：
  python add_batch_time_v4.py
Copyright © 2026 David YE
"""
import openpyxl
import os
from datetime import date

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, 'NEV公告参数汇总表_合并版（341~410批）.xlsx')
SHEET = 'NEV公告参数汇总'
TAG_SHEET = '批次时间表'
YELLOW = openpyxl.styles.PatternFill('solid', fgColor='FFF2CC')  # 推断年月黄底


# 公告精确日期（已从工信部官网公告页核实；键=批次，值=(YYYY,MM,DD)）
EXACT_BATCH = {
    341: (2021, 3, 8), 344: (2021, 6, 11), 346: (2021, 8, 10), 350: (2021, 12, 7),
    360: (2022, 9, 16), 368: (2023, 3, 10), 370: (2023, 5, 9),
    379: (2024, 2, 4), 380: (2024, 3, 11), 381: (2024, 4, 15), 383: (2024, 5, 31),
    390: (2024, 12, 31), 391: (2025, 2, 17), 392: (2025, 3, 13), 393: (2025, 4, 15),
    395: (2025, 6, 19), 396: (2025, 7, 15), 397: (2025, 8, 7), 398: (2025, 9, 8),
    399: (2025, 10, 14), 400: (2025, 11, 7), 401: (2025, 12, 15), 402: (2025, 12, 31),
    403: (2026, 2, 9), 404: (2026, 3, 18), 405: (2026, 4, 14), 406: (2026, 5, 9),
    407: (2026, 6, 11), 408: (2026, 7, 17), 409: (2026, 8, 13),
}

# 公示精确日期（已考证；工信部公示清单发布日期）
EXACT_DISPLAY = {
    341: (2021, 2, 3), 344: (2021, 5, 14), 360: (2022, 8, 11), 368: (2023, 2, 10),
    379: (2024, 1, 5), 380: (2024, 2, 18), 381: (2024, 3, 22), 383: (2024, 5, 11),
    385: (2024, 7, 12), 390: (2024, 12, 10), 395: (2025, 5, 24), 398: (2025, 8, 8),
    410: (2026, 8, 7),
}

# 公告年月（YYYY,MM）；无精确批次用相邻批次节律推断
YM_INFER = {
    342: (2021, 4), 343: (2021, 5), 345: (2021, 7), 347: (2021, 9), 348: (2021, 10),
    349: (2021, 11), 351: (2022, 1), 352: (2022, 2), 353: (2022, 3), 354: (2022, 4),
    355: (2022, 5), 356: (2022, 6), 357: (2022, 7), 358: (2022, 8), 359: (2022, 8),
    361: (2022, 10), 362: (2022, 11), 363: (2022, 12), 364: (2022, 12),
    365: (2023, 1), 366: (2023, 1), 367: (2023, 2), 369: (2023, 4),
    371: (2023, 6), 372: (2023, 7), 373: (2023, 8), 374: (2023, 9), 375: (2023, 10),
    376: (2023, 11), 377: (2023, 12), 378: (2024, 1), 382: (2024, 5), 384: (2024, 6),
    385: (2024, 7), 386: (2024, 8), 387: (2024, 9), 388: (2024, 10), 389: (2024, 11),
    394: (2025, 5), 410: (2026, 8),
}


def main():
    print(f'[1/3] 加载工作簿: {os.path.basename(SRC)}')
    wb = openpyxl.load_workbook(SRC)
    ws = wb[SHEET]
    cr = wb['变更记录']

    last_seq = 0
    for r in range(cr.max_row, 1, -1):
        v = cr.cell(r, 1).value
        if isinstance(v, int):
            last_seq = v
            break

    def ledger(batch, field, ctype, old, new, srcs):
        nonlocal last_seq
        last_seq += 1
        cr.append([last_seq, batch, None, None, field, ctype, old, new, srcs])

    # ---------- 新增「批次时间表」工作表 ----------
    print('[2/3] 生成「批次时间表」…')
    if TAG_SHEET in wb.sheetnames:
        print('      已存在，跳过新增；如需重建请删除该表后重跑')
        return
    ts = wb.create_sheet(TAG_SHEET)
    headers = ['批次', '公告年月', '公告日期', '公示年月', '公示日期', '精度', '说明/来源']
    for c, h in enumerate(headers, 1):
        cell = ts.cell(1, c, h)
        cell.font = ws.cell(1, 1).font.copy()
        f = ws.cell(1, 1).fill
        if f and f.patternType:
            cell.fill = f.copy()

    n_exact = n_infer = 0
    for batch in range(341, 411):
        row = batch - 341 + 2
        ts.cell(row, 1, batch)
        if batch in EXACT_BATCH:
            y, m, d = EXACT_BATCH[batch]
            ts.cell(row, 2, f'{y}-{m:02d}')
            ts.cell(row, 3, f'{y}-{m:02d}-{d:02d}')
            src = f'工信部公告 发布日期 {y}-{m:02d}-{d:02d}'
            acc = '官方精确日期'
            n_exact += 1
        else:
            y, m = YM_INFER[batch]
            ts.cell(row, 2, f'{y}-{m:02d}')
            c = ts.cell(row, 2)
            c.fill = YELLOW
            src = '依据相邻批次月度节奏推断（年月级）'
            acc = '按月节奏推断'
            n_infer += 1
        if batch in EXACT_DISPLAY:
            y, m, d = EXACT_DISPLAY[batch]
            ts.cell(row, 4, f'{y}-{m:02d}')
            ts.cell(row, 5, f'{y}-{m:02d}-{d:02d}')
        ts.cell(row, 6, acc)
        ts.cell(row, 7, src)

    ts.column_dimensions['A'].width = 6
    ts.column_dimensions['B'].width = 10
    ts.column_dimensions['C'].width = 12
    ts.column_dimensions['D'].width = 10
    ts.column_dimensions['E'].width = 12
    ts.column_dimensions['F'].width = 14
    ts.column_dimensions['G'].width = 42

    ledger('341~410', '批次时间', '补充', '(无此维度)', f'新增「{TAG_SHEET}」工作表：'
           f'公告精确日 {n_exact} 批 / 推断年月 {n_infer} 批', '工信部官网公告逐批核实+月度节律推断')

    wb.save(SRC)
    print(f'      完成：写入 {n_exact + n_infer} 批（精确 {n_exact} / 推断 {n_infer}），已存盘并记入变更记录')
    print('[3/3] 完成')


if __name__ == '__main__':
    main()