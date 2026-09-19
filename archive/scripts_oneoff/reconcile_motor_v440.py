# -*- coding: utf-8 -*-
"""v4.4.0 电机列对账（v2 修正版）：仅当 后电机 存在时强制 总=前+后。
v1 教训：后电机为空且 前≠总 时把总压平成前 —— 会把四驱车的系统总功率
错误压成单电机款共识前值（65 行损坏，已由 fix_motor_logic_v441.py 回滚）。
新规则：后空 → 不动总（维持"仅总未拆"或单电机本来就相等，不再改写）。
"""
import sys

import openpyxl
from openpyxl.styles import PatternFill

sys.stdout.reconfigure(encoding="utf-8")
SRC = "NEV公告参数汇总表_合并版（341~410批）.xlsx"
RED = PatternFill("solid", fgColor="FFC7CE")


def parse_pt(v):
    if v is None:
        return None, None, ""
    s = str(v).strip()
    if not s:
        return None, None, ""
    parts = s.split("/")
    try:
        p = float(parts[0]) if parts[0] else None
    except ValueError:
        p = None
    try:
        t = float(parts[1]) if len(parts) > 1 and parts[1] else None
    except ValueError:
        t = None
    return p, t, s


def pt_cell(p, t):
    if p is None:
        return ""
    a = str(int(p)) if float(p) == int(p) else str(p)
    if t is None:
        return a
    b = str(int(t)) if float(t) == int(t) else str(t)
    return f"{a}/{b}"


def main():
    wb = openpyxl.load_workbook(SRC)
    ws = wb["NEV公告参数汇总"]
    hdr = [str(c.value).strip() if c.value else "" for c in ws[1]]
    fi = {h: i for i, h in enumerate(hdr)}
    iF, iT, iR = fi["前电机功率/扭矩"], fi["电机总功率/扭矩"], fi["后电机功率/扭矩"]
    cr = wb["变更记录"]
    seq = 0
    for rr in range(cr.max_row, 1, -1):
        v = cr.cell(rr, 1).value
        if isinstance(v, int):
            seq = v
            break
    fixed = kept = 0
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value is None:
            continue
        fp, ft, _ = parse_pt(ws.cell(r, iF + 1).value)
        rp, rt, _ = parse_pt(ws.cell(r, iR + 1).value)
        tp, tt, ts = parse_pt(ws.cell(r, iT + 1).value)
        # v2：仅当 前后都存在 才强制 总=前+后；后空绝不改写总列（防压平四驱系统值）
        if fp is None or rp is None:
            continue
        exp_p = fp + rp
        exp_t = ft + rt if (ft is not None and rt is not None) else None
        if tp is not None and abs(tp - exp_p) <= 2 and (
                (exp_t is None) or (tt is not None and abs(tt - exp_t) <= 2)):
            kept += 1
            continue
        val = pt_cell(exp_p, exp_t)
        cell = ws.cell(r, iT + 1)
        old = cell.value
        cell.value = val
        cell.fill = RED
        fixed += 1
        seq += 1
        cr.append([seq, ws.cell(r, 1).value, ws.cell(r, fi["产品型号"] + 1).value,
                   ws.cell(r, fi["动力类型"] + 1).value, "电机总功率/扭矩", "数据清洗更正",
                   "" if old is None else str(old)[:40], val, "对账修正(v2)：前后齐全，总=前+后"])
    wb.save(SRC)
    print(f"对账完成(v2)：一致 {kept} 行，重写总列 {fixed} 行（台账末序号 {seq}）")


if __name__ == "__main__":
    main()
