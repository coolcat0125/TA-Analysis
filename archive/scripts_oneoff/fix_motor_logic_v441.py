# -*- coding: utf-8 -*-
"""v4.4.1 电机列逻辑错误修复（用户核实反馈）
================================================================
A. 小米SU7 Ultra（批403 XMA7001MBEV*，7行）：源数据将系统总功率 1138/1770 误作单电机
   塞入前电机列 → 按汽车之家参数（用户附件）更正：前 288/500、后 850/1270、总 1138/1770；
   只补空顺带 电池类型=三元锂电池 / 电芯供应商=宁德时代 / 电池容量=93.7（如空）
B. 压平回滚：对账(reconcile)曾把"前=跨款共识小值、总=真实系统值"的行压平成 总=前，
   从台账「对账修正」原值恢复大值总功率，并清除错配的前电机值（该值来自单电机款共识）
C. 全表同类扫描：后空 & 前存在 & 总>前×1.05 → 清除前电机值（跨款共识错配/系统值误入），保留总
D. 顺带扫描 前>总×1.05（后空）清单输出不自动改（留人工/证据处置）
全部台账（类型=数据清洗更正/证据核实更正，媒体橙 F8CBAD / 浅红 FFC7CE）
"""
import re
import sys

import openpyxl
from openpyxl.styles import PatternFill

sys.stdout.reconfigure(encoding="utf-8")
SRC = "NEV公告参数汇总表_合并版（341~410批）.xlsx"
RED = PatternFill("solid", fgColor="FFC7CE")
MEDIA = PatternFill("solid", fgColor="F8CBAD")


def parse_pt(v):
    if v is None:
        return None, None, ""
    s = str(v).strip()
    if not s:
        return None, None, ""
    parts = s.split("/")
    try:
        p = float(parts[0]) if parts[0] else None
    except ValueError:
        p = None
    try:
        t = float(parts[1]) if len(parts) > 1 and parts[1] else None
    except ValueError:
        t = None
    return p, t, s


def pt_cell(p, t):
    if p is None:
        return ""
    a = str(int(p)) if float(p) == int(p) else str(p)
    if t is None:
        return a
    b = str(int(t)) if float(t) == int(t) else str(t)
    return f"{a}/{b}"


def main():
    wb = openpyxl.load_workbook(SRC)
    ws = wb["NEV公告参数汇总"]
    hdr = [str(c.value).strip() if c.value else "" for c in ws[1]]
    fi = {h: i for i, h in enumerate(hdr)}
    iF, iT, iR = fi["前电机功率/扭矩"], fi["电机总功率/扭矩"], fi["后电机功率/扭矩"]
    rowmap = {}
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value is None:
            continue
        rowmap.setdefault((str(ws.cell(r, 1).value).strip(),
                           str(ws.cell(r, fi["产品型号"] + 1).value).strip()), r)

    cr = wb["变更记录"]
    seq = 0
    for rr in range(cr.max_row, 1, -1):
        v = cr.cell(rr, 1).value
        if isinstance(v, int):
            seq = v
            break
    led_rows = []
    stats = {"ultra": 0, "restore": 0, "clearfront": 0, "batt": 0, "flagged": 0}

    def led(row_vals):
        nonlocal seq
        seq += 1
        led_rows.append(row_vals)
        cr.append([seq] + row_vals)

    # ---- B) 压平回滚（先于 C，依据台账原值）----
    for rr in range(2, cr.max_row + 1):
        vals = [cr.cell(rr, c).value for c in range(1, 10)]
        if len(vals) < 9 or vals[5] != "数据清洗更正":
            continue
        if not str(vals[8] or "").startswith("对账修正") or vals[4] != "电机总功率/扭矩":
            continue

        def pn(s):
            try:
                return float(str(s).split("/")[0])
            except (ValueError, TypeError, IndexError):
                return None
        po, pw = pn(vals[6]), pn(vals[7])
        if po is None or pw is None or po <= pw * 1.05:
            continue
        key = (str(vals[1]).strip(), str(vals[2] or "").strip())
        r = rowmap.get(key)
        if r is None:
            continue
        fp, ft, _ = parse_pt(ws.cell(r, iF + 1).value)
        rp, rt, _ = parse_pt(ws.cell(r, iR + 1).value)
        # 该行当前 rear 为空且 front == 被压平的小值 → 确属错配：恢复大值总 + 清前
        if rp is None and fp is not None and abs(fp - pw) <= 2:
            op, ot, _ = parse_pt(po)
            ws.cell(r, iT + 1).value = pt_cell(op, ot)
            ws.cell(r, iT + 1).fill = RED
            led([vals[1], vals[2], vals[3], "电机总功率/扭矩", "数据清洗更正",
                 pw, pt_cell(op, ot), "压平回滚：恢复共识覆盖前的真实系统总功率"])
            ws.cell(r, iF + 1).value = None
            ws.cell(r, iF + 1).fill = RED
            led([vals[1], vals[2], vals[3], "前电机功率/扭矩", "数据清洗更正",
                 pt_cell(fp, ft), "(空)", "跨款共识错配：单电机款前电机值误套四驱款，清除（总功率已恢复为系统值）"])
            stats["restore"] += 1

    # ---- C) 全表同类扫描（后空 & 前 & 总 差异>5%）----
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value is None:
            continue
        fp, ft, fs = parse_pt(ws.cell(r, iF + 1).value)
        rp, rt, _ = parse_pt(ws.cell(r, iR + 1).value)
        tp, tt, ts = parse_pt(ws.cell(r, iT + 1).value)
        if rp is not None or fp is None or tp is None:
            continue
        if tp > fp * 1.05 and tp - fp > 5:
            ws.cell(r, iF + 1).value = None
            ws.cell(r, iF + 1).fill = RED
            stats["clearfront"] += 1
            led([ws.cell(r, 1).value, ws.cell(r, fi["产品型号"] + 1).value,
                 ws.cell(r, fi["动力类型"] + 1).value, "前电机功率/扭矩", "数据清洗更正",
                 fs, "(空)", f"疑似系统值误作前电机/跨款错配（总{ts}>前{fs}），清除前值保留系统总功率"])

    # ---- A) SU7 Ultra 按附件更正 ----
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value is None:
            continue
        model = str(ws.cell(r, fi["产品型号"] + 1).value or "").strip()
        if not (str(ws.cell(r, 1).value).strip() == "403" and re.match(r"^XMA7001MBEV", model)):
            continue
        ws.cell(r, iF + 1).value = "288/500"
        ws.cell(r, iF + 1).fill = MEDIA
        ws.cell(r, iR + 1).value = "850/1270"
        ws.cell(r, iR + 1).fill = MEDIA
        ws.cell(r, iT + 1).value = "1138/1770"
        ws.cell(r, iT + 1).fill = MEDIA
        stats["ultra"] += 1
        for fld, col, val in (("前电机功率/扭矩", iF, "288/500"), ("后电机功率/扭矩", iR, "850/1270"),
                              ("电机总功率/扭矩", iT, "1138/1770")):
            led([403, model, ws.cell(r, fi["动力类型"] + 1).value, fld, "证据核实更正",
                 "1138/1770(前电机列误存系统值)" if fld != "电机总功率/扭矩" else "1138/1770", val,
                 "汽车之家SU7 Ultra 2025款标准版参数（用户附件）：三电机 前288kW/500Nm+后850kW/1270Nm=1138kW/1770Nm"])
        # 只补空：电池类型/电芯供应商/电池容量
        for name, val in (("电池类型", "三元锂电池"), ("电芯供应商", "宁德时代"), ("电池容量(kWh)", 93.7)):
            c = ws.cell(r, fi[name] + 1)
            if c.value in (None, ""):
                c.value = val
                c.fill = MEDIA
                led([403, model, ws.cell(r, fi["动力类型"] + 1).value, name, "证据核实更正",
                     "(空)", val, "汽车之家SU7 Ultra参数（用户附件）：麒麟电池三元锂/宁德时代/93.7kWh"])
                stats["batt"] += 1

    # ---- D) 前>总 清单（不自动改）----
    print("== 前>总（后空）留人工核实清单 ==")
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value is None:
            continue
        fp, _, fs = parse_pt(ws.cell(r, iF + 1).value)
        rp, _, _ = parse_pt(ws.cell(r, iR + 1).value)
        tp, _, ts = parse_pt(ws.cell(r, iT + 1).value)
        if rp is None and fp is not None and tp is not None and fp > tp * 1.05:
            stats["flagged"] += 1
            if stats["flagged"] <= 10:
                print(f"  批{ws.cell(r,1).value} {ws.cell(r, fi['产品型号']+1).value} 前{fs} 总{ts}")

    for row in led_rows:
        pass
    # 保存（文件被 Excel/WPS 占用时写临时文件并重试替换）
    import os
    import time
    saved = False
    try:
        wb.save(SRC)
        saved = True
    except PermissionError:
        tmp = SRC.replace("（341~410批）", "（341~410批）tmp")
        wb.save(tmp)
        for i in range(12):
            time.sleep(10)
            try:
                os.replace(tmp, SRC)
                saved = True
                print(f"重试替换成功（第 {i+1} 次）")
                break
            except PermissionError:
                continue
        if not saved:
            kept = SRC.replace("（341~410批）", "（341~410批）fixed")
            os.replace(tmp, kept)
            print(f"持续被占用，结果保留为 {kept}（请手工替换主文件）")
    print(f"台账追加 {len(led_rows)} 条（末序号 {seq}）")
    print("统计:", stats)


if __name__ == "__main__":
    main()
